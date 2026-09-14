# -*- coding: utf-8 -*-
"""Build the brief's five reference decks and render every slide to PNG.

    venv313/Scripts/python.exe tests/preview_decks.py [subject-id ...]

There is no assertion here on purpose. This is the tool for looking at
the result the way a designer would — the checks live in
test_presentation_templates.py. Rendering goes through LibreOffice, which
is also what renders the in-app preview (app/pptx_pdf.py), so what comes
out of here is what a teacher sees on their phone.

Output: tests/_decks/png/<id>-NN.png, one per slide.
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from pptx import Presentation                                     # noqa: E402
from app.export_builder import build_presentation_pptx            # noqa: E402
from app import pptx_pdf                                          # noqa: E402

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_decks")
PNG = os.path.join(OUT, "png")


# The five the brief names, with content shaped the way the model really
# writes it: a kind on every slide, short bullet fragments, a paragraph.
DECKS = {
    "biology": {
        "title": "Строение клетки", "subject": "Биология", "grade": "8 класс",
        "description": "Из чего состоит клетка и как работает каждая её часть.",
        "slides": [
            {"title": "Что такое клетка", "kind": "intro",
             "bullet_points": ["Наименьшая единица живого"],
             "body": "Клетка — наименьшая единица живого организма, способная к самостоятельному обмену веществ, росту и делению. Всё, что происходит в организме, начинается здесь."},
            {"title": "Из чего состоит клетка", "kind": "concepts",
             "bullet_points": ["Мембрана", "Ядро", "Митохондрии", "Рибосомы", "Цитоплазма"]},
            {"title": "Как клетка получает энергию", "kind": "example",
             "bullet_points": ["Глюкоза поступает внутрь", "Митохондрия расщепляет её", "Выделяется АТФ"],
             "body": "Дыхание клетки — это цепочка, а не одно событие."},
            {"title": "Растительная и животная клетка", "kind": "explanation",
             "bullet_points": ["Есть клеточная стенка", "Есть хлоропласты", "Стенки нет", "Хлоропластов нет"]},
            {"title": "Что запомнить", "kind": "summary",
             "bullet_points": ["Клетка — единица жизни", "У каждой части своя работа", "Деление даёт рост"]},
        ],
    },
    "mathematics": {
        "title": "Квадратные уравнения", "subject": "Математика", "grade": "8 класс",
        "description": "Как узнать квадратное уравнение и как его решить.",
        "slides": [
            {"title": "Что это такое", "kind": "intro",
             "bullet_points": ["Уравнение со степенью два"],
             "body": "Квадратным называют уравнение вида ax² + bx + c = 0, где a не равно нулю. Степень два меняет всё: у такого уравнения может быть два корня, один или ни одного."},
            {"title": "Порядок решения", "kind": "example",
             "bullet_points": ["Записать коэффициенты", "Найти дискриминант", "Подставить в формулу", "Проверить корни"],
             "body": "Каждый шаг проверяем отдельно — так ошибка находится сразу."},
            {"title": "Сколько корней", "kind": "concepts",
             "bullet_points": ["D > 0 — два корня", "D = 0 — один", "D < 0 — корней нет"]},
            {"title": "Формула корней", "kind": "formula",
             "bullet_points": ["$x = \\frac{-b \\pm \\sqrt{D}}{2a}$", "D — это дискриминант"]},
            {"title": "Итог урока", "kind": "summary",
             "bullet_points": ["Узнаём по степени", "Считаем дискриминант", "Находим корни"]},
        ],
    },
    "geography": {
        "title": "Климатические пояса Земли", "subject": "География", "grade": "7 класс",
        "description": "Почему на Земле разный климат и как он распределён.",
        "slides": [
            {"title": "Откуда берётся климат", "kind": "intro",
             "bullet_points": ["Всё решает угол падения солнечных лучей"],
             "body": "Земля шарообразна, поэтому солнечные лучи падают на разные широты под разным углом. У экватора они почти отвесны, у полюсов скользят по поверхности — отсюда и пояса."},
            {"title": "Основные пояса", "kind": "concepts",
             "bullet_points": ["Экваториальный", "Тропический", "Умеренный", "Арктический"]},
            {"title": "Экваториальный и арктический", "kind": "explanation",
             "bullet_points": ["Жарко весь год", "Осадков очень много", "Холодно весь год", "Осадков почти нет"]},
            {"title": "Сколько поясов", "kind": "example",
             "bullet_points": ["13 климатических поясов", "7 основных и 6 переходных", "Пояса симметричны экватору"]},
            {"title": "Главное", "kind": "summary",
             "bullet_points": ["Климат зависит от широты", "Пояса идут полосами", "Переходные лежат между основными"]},
        ],
    },
    "history": {
        "title": "Древний Египет", "subject": "История Таджикистана", "grade": "5 класс",
        "description": "Государство на Ниле: как оно возникло и чем запомнилось.",
        "slides": [
            {"title": "Страна на Ниле", "kind": "intro",
             "bullet_points": ["Всё держалось на реке"],
             "body": "Египет вытянулся узкой полосой вдоль Нила. Река разливалась каждый год и оставляла плодородный ил, и только поэтому в пустыне смогло вырасти крупное государство."},
            {"title": "Как шла история", "kind": "summary",
             "bullet_points": ["Объединение страны", "Строительство пирамид", "Расцвет Нового царства", "Завоевание Египта"]},
            {"title": "Кто управлял страной", "kind": "concepts",
             "bullet_points": ["Фараон", "Жрецы", "Писцы", "Земледельцы"]},
            {"title": "Пирамиды", "kind": "explanation",
             "bullet_points": ["Гробницы фараонов"],
             "body": "Пирамида была не просто могилой, а целым комплексом: храм, дорога, склады. Строили её десятилетиями, и работала на этом не армия рабов, а наёмные общинники в свободное от полевых работ время."},
            {"title": "Что запомнить", "kind": "summary",
             "bullet_points": ["Нил кормил страну", "Власть у фараона", "Письменность — иероглифы"]},
        ],
    },
    "physics": {
        "title": "Закон Ома", "subject": "Физика", "grade": "8 класс",
        "description": "Как связаны сила тока, напряжение и сопротивление.",
        "slides": [
            {"title": "Три величины", "kind": "intro",
             "bullet_points": ["Ток, напряжение, сопротивление"],
             "body": "В любой цепи работают три величины, и они связаны жёстко: зная две, всегда можно найти третью. Эту связь и описывает закон Ома."},
            {"title": "Что означает каждая", "kind": "concepts",
             "bullet_points": ["Сила тока — сколько заряда проходит", "Напряжение — что толкает заряд", "Сопротивление — что мешает"]},
            {"title": "Формула", "kind": "formula",
             "bullet_points": ["$I = \\frac{U}{R}$", "Ток прямо пропорционален напряжению"]},
            {"title": "Как решать задачу", "kind": "example",
             "bullet_points": ["Начертить схему", "Выписать известное", "Подставить в формулу"],
             "body": "Схема нужна всегда — она показывает, последовательное соединение или параллельное."},
            {"title": "Итог", "kind": "summary",
             "bullet_points": ["I = U / R", "Больше R — меньше I", "Больше U — больше I"]},
        ],
    },
}

for d in DECKS.values():
    d["language"] = "Русский"


def render(deck_id, deck):
    os.makedirs(PNG, exist_ok=True)
    data = build_presentation_pptx(deck).getvalue()
    src = os.path.join(OUT, f"{deck_id}.pptx")
    with open(src, "wb") as fh:
        fh.write(data)

    soffice = pptx_pdf.find_binary()
    if not soffice:
        print("LibreOffice not found — wrote the .pptx only")
        return

    n = len(Presentation(io.BytesIO(data)).slides)
    tmp = tempfile.mkdtemp()
    try:
        for idx in range(n):
            prs = Presentation(io.BytesIO(data))
            lst = prs.slides._sldIdLst
            for j, sld in enumerate(list(lst)):
                if j != idx:
                    lst.remove(sld)
            one = os.path.join(tmp, f"{deck_id}-{idx:02d}.pptx")
            prs.save(one)
        subprocess.run([soffice, "--headless", "--convert-to",
                        "png:impress_png_Export", "--outdir", PNG]
                       + [os.path.join(tmp, f) for f in sorted(os.listdir(tmp))],
                       capture_output=True, timeout=300)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"{deck_id}: {n} слайдов -> {PNG}")


if __name__ == "__main__":
    wanted = sys.argv[1:] or list(DECKS)
    for deck_id in wanted:
        if deck_id not in DECKS:
            print(f"нет такой колоды: {deck_id}")
            continue
        render(deck_id, DECKS[deck_id])
