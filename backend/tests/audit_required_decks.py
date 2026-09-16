import asyncio
import io
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from pptx import Presentation
from pptx.util import Emu

from app.ai_service import _render_slide_images
from app.export_builder import build_presentation_pptx

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_decks")
SLIDE_W_EMU = Emu(int(13.333 * 914400))
SLIDE_H_EMU = Emu(int(7.5 * 914400))
MIN_READABLE_PT = 9

DECKS = {
    "python": {
        "title": "Python: циклы и условия", "subject": "Информатика", "grade": "10 класс",
        "description": "Как Python управляет повторением и выбором действий.",
        "slides": [
            {"title": "Что такое Python", "kind": "intro",
             "bullet_points": ["Понятный синтаксис", "Используется везде"],
             "body": "Python — язык программирования с простым синтаксисом, который используют в вебе, анализе данных и автоматизации.",
             "image_query": "Python programming language logo"},
            {"title": "Цикл for", "kind": "example",
             "code": {"language": "python", "snippet": "for i in range(5):\n    print(i)",
                      "explanation": "Печатает числа от 0 до 4"},
             "bullet_points": ["range задаёт диапазон", "Тело цикла — с отступом"]},
            {"title": "Цикл while", "kind": "explanation",
             "code": {"language": "python", "snippet": "n = 5\nwhile n > 0:\n    print(n)\n    n -= 1"},
             "bullet_points": ["Условие проверяется перед каждым шагом"]},
            {"title": "Условие if / else", "kind": "concepts",
             "code": {"language": "python", "snippet": "if x > 0:\n    print('positive')\nelse:\n    print('not positive')"},
             "bullet_points": ["Одна из двух веток выполнится"]},
            {"title": "For или while", "kind": "comparison",
             "bullet_points": ["Известно количество шагов", "Известно условие остановки",
                               "Неизвестно число шагов", "Проверка перед каждым шагом"]},
            {"title": "Что запомнить", "kind": "summary",
             "bullet_points": ["for — по последовательности", "while — по условию", "Оба используют отступы"]},
        ],
    },
    "biology": {
        "title": "Строение клетки", "subject": "Биология", "grade": "8 класс",
        "description": "Из чего состоит клетка и как работает каждая её часть.",
        "slides": [
            {"title": "Что такое клетка", "kind": "intro",
             "bullet_points": ["Наименьшая единица живого"],
             "body": "Клетка — наименьшая единица живого организма, способная к самостоятельному обмену веществ, росту и делению.",
             "image_query": "plant cell structure diagram"},
            {"title": "Из чего состоит клетка", "kind": "concepts",
             "bullet_points": ["Мембрана", "Ядро", "Митохондрии", "Рибосомы", "Цитоплазма"]},
            {"title": "Как клетка получает энергию", "kind": "example",
             "bullet_points": ["Глюкоза поступает внутрь", "Митохондрия расщепляет её", "Выделяется АТФ"],
             "image_query": "mitochondrion structure diagram"},
            {"title": "Растительная и животная клетка", "kind": "explanation",
             "bullet_points": ["Есть клеточная стенка", "Есть хлоропласты", "Стенки нет", "Хлоропластов нет"],
             "image_query": "plant animal cell comparison diagram"},
            {"title": "Ядро клетки", "kind": "explanation",
             "bullet_points": ["Хранит ДНК", "Управляет клеткой"],
             "image_query": "cell nucleus diagram"},
            {"title": "Что запомнить", "kind": "summary",
             "bullet_points": ["Клетка — единица жизни", "У каждой части своя работа", "Деление даёт рост"]},
        ],
    },
    "mathematics": {
        "title": "Квадратные уравнения", "subject": "Математика", "grade": "8 класс",
        "description": "Как узнать квадратное уравнение и как его решить.",
        "slides": [
            {"title": "Что это такое", "kind": "intro",
             "bullet_points": ["Уравнение со степенью два"],
             "body": "Квадратным называют уравнение вида ax² + bx + c = 0, где a не равно нулю."},
            {"title": "Порядок решения", "kind": "example",
             "bullet_points": ["Записать коэффициенты", "Найти дискриминант", "Подставить в формулу", "Проверить корни"],
             "body": "Каждый шаг проверяем отдельно."},
            {"title": "Сколько корней", "kind": "concepts",
             "bullet_points": ["D > 0 — два корня", "D = 0 — один", "D < 0 — корней нет"]},
            {"title": "Формула корней", "kind": "formula",
             "bullet_points": ["$x = \\frac{-b \\pm \\sqrt{D}}{2a}$", "D — это дискриминант"]},
            {"title": "Уравнение и график", "kind": "explanation",
             "bullet_points": ["Корни — это точки пересечения с осью X"],
             "image_query": "parabola graph quadratic function"},
            {"title": "Итог урока", "kind": "summary",
             "bullet_points": ["Узнаём по степени", "Считаем дискриминант", "Находим корни"]},
        ],
    },
    "geography": {
        "title": "Климатические пояса Земли", "subject": "География", "grade": "7 класс",
        "description": "Почему на Земле разный климат и как он распределён.",
        "slides": [
            {"title": "Откуда берётся климат", "kind": "intro",
             "bullet_points": ["Угол падения солнечных лучей"],
             "body": "Земля шарообразна, поэтому солнечные лучи падают на разные широты под разным углом.",
             "image_query": "sun angle latitude diagram"},
            {"title": "Основные пояса", "kind": "concepts",
             "bullet_points": ["Экваториальный", "Тропический", "Умеренный", "Арктический"],
             "image_query": "world climate zones map"},
            {"title": "Экваториальный и арктический", "kind": "explanation",
             "bullet_points": ["Жарко весь год", "Осадков очень много", "Холодно весь год", "Осадков почти нет"]},
            {"title": "Сколько поясов", "kind": "example",
             "bullet_points": ["13 климатических поясов", "7 основных и 6 переходных", "Пояса симметричны экватору"]},
            {"title": "Пояса на карте", "kind": "explanation",
             "bullet_points": ["Каждый пояс — своя полоса"],
             "image_query": "climate zones world map illustrated"},
            {"title": "Главное", "kind": "summary",
             "bullet_points": ["Климат зависит от широты", "Пояса идут полосами", "Переходные лежат между основными"]},
        ],
    },
    "history": {
        "title": "Древний Египет", "subject": "История Таджикистана", "grade": "6 класс",
        "description": "Государство на Ниле: как оно возникло и чем запомнилось.",
        "slides": [
            {"title": "Страна на Ниле", "kind": "intro",
             "bullet_points": ["Всё держалось на реке"],
             "body": "Египет вытянулся узкой полосой вдоль Нила. Река разливалась каждый год и оставляла плодородный ил.",
             "image_query": "Nile river ancient Egypt map"},
            {"title": "Как шла история", "kind": "summary",
             "bullet_points": ["Объединение страны", "Строительство пирамид", "Расцвет Нового царства", "Завоевание Египта"]},
            {"title": "Кто управлял страной", "kind": "concepts",
             "bullet_points": ["Фараон", "Жрецы", "Писцы", "Земледельцы"]},
            {"title": "Пирамиды", "kind": "explanation",
             "bullet_points": ["Гробницы фараонов"],
             "body": "Пирамида была не просто могилой, а целым комплексом: храм, дорога, склады.",
             "image_query": "Great Pyramids Giza ancient Egypt photograph"},
            {"title": "Иероглифы", "kind": "example",
             "bullet_points": ["Письмо рисунками", "Читали жрецы и писцы"],
             "image_query": "Egyptian hieroglyphs ancient inscription"},
            {"title": "Что запомнить", "kind": "summary",
             "bullet_points": ["Нил кормил страну", "Власть у фараона", "Письменность — иероглифы"]},
        ],
    },
}
for d in DECKS.values():
    d["language"] = "Русский"


def audit_pptx(data: bytes) -> list[str]:
    problems = []
    prs = Presentation(io.BytesIO(data))
    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape.left is None or shape.top is None:
                continue
            if shape.left < -Emu(1000) or shape.top < -Emu(1000):
                problems.append(f"s{n}: текст за левым/верхним краем: {text[:30]!r}")
            right = shape.left + (shape.width or 0)
            bottom = shape.top + (shape.height or 0)
            if right > SLIDE_W_EMU + Emu(152400):
                problems.append(f"s{n}: текст выходит вправо: {text[:30]!r}")
            if bottom > SLIDE_H_EMU + Emu(152400):
                problems.append(f"s{n}: текст выходит вниз: {text[:30]!r}")
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size and run.font.size.pt < MIN_READABLE_PT:
                        problems.append(f"s{n}: {run.font.size.pt}pt мельче порога: {text[:24]!r}")
    return problems


async def run_one(deck_id, deck):
    t0 = time.perf_counter()
    await _render_slide_images(deck, deck["title"])
    img_time = time.perf_counter() - t0

    slides = deck["slides"]
    images = [s["image"]["path"] for s in slides if isinstance(s.get("image"), dict) and s["image"].get("path")]
    unique = len(set(images))
    dup = len(images) - unique

    data = build_presentation_pptx(deck).getvalue()
    with open(os.path.join(OUT, f"required_{deck_id}.pptx"), "wb") as fh:
        fh.write(data)

    prs = Presentation(io.BytesIO(data))
    layouts = [s._element.cSld.get("name") or "-" for s in prs.slides][1:]
    problems = audit_pptx(data)

    print(f"\n=== {deck_id} — «{deck['title']}» ({deck['subject']}, {deck['grade']}) ===")
    print(f"  время подбора картинок: {img_time:.1f}с")
    print(f"  слайдов с картинкой: {len(images)}/{len(slides)}   уникальных: {unique}   дублей: {dup}")
    print(f"  layouts: {' -> '.join(layouts)}")
    print(f"  разных layouts: {len(set(layouts))} из {len(layouts)}")
    print(f"  размер файла: {len(data) // 1024} КБ")
    if problems:
        print(f"  ПРОБЛЕМЫ ВЁРСТКИ ({len(problems)}):")
        for p in problems[:5]:
            print(f"    - {p}")
    else:
        print("  вёрстка: без проблем")
    return {
        "deck": deck_id, "images": len(images), "unique": unique, "dup": dup,
        "layouts": len(set(layouts)), "total_slides": len(layouts), "problems": len(problems),
    }


async def main():
    results = []
    for deck_id, deck in DECKS.items():
        results.append(await run_one(deck_id, deck))

    print("\n" + "=" * 60)
    print("ИТОГО")
    print("=" * 60)
    total_images = sum(r["images"] for r in results)
    total_unique = sum(r["unique"] for r in results)
    total_dup = sum(r["dup"] for r in results)
    total_problems = sum(r["problems"] for r in results)
    for r in results:
        print(f"  {r['deck']:<12} картинок={r['images']}  уникальных={r['unique']}  "
              f"дублей={r['dup']}  layouts={r['layouts']}/{r['total_slides']}  "
              f"проблем={r['problems']}")
    print(f"\n  всего картинок: {total_images}, уникальных: {total_unique}, дублей: {total_dup}")
    print(f"  всего проблем вёрстки: {total_problems}")
    print(f"\nФайлы: {OUT}/required_*.pptx")


asyncio.run(main())
