import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from pptx import Presentation
from pptx.util import Emu

from app import subject_templates
from app.export_builder import build_presentation_pptx, resolve_deck_theme
from app.image_query import build_queries

FAIL = []
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_decks")


def check(name, condition, detail=""):
    print(("  PASS  " if condition else "  FAIL  ") + name + (f"   [{detail}]" if detail else ""))
    if not condition:
        FAIL.append(name)


REQUIRED_SUBJECTS = [
    "Математика", "Физика", "Химия", "Биология", "География",
    "История Таджикистана", "Информатика", "Русский язык",
    "Таджикская литература", "Английский язык", "Таджикский язык",
    "Обществознание", "Экология", "Начальные классы",
]

SLIDE_W_EMU = Emu(int(13.333 * 914400))
SLIDE_H_EMU = Emu(int(7.5 * 914400))

MIN_READABLE_PT = 9


def deck_fixture(subject, language="Русский", grade="8 класс"):
    return {
        "title": f"Открытый урок по предмету «{subject}»",
        "description": "Пробный урок, собранный тестом для проверки вёрстки шаблона.",
        "subject": subject,
        "grade": grade,
        "language": language,
        "slides": [
            {"title": "Введение в тему", "kind": "intro",
             "bullet_points": ["Что изучаем", "Зачем это нужно"],
             "body": "Этот урок вводит тему и объясняет, почему она важна для дальнейшей работы на уроках.",
             "speaker_notes": "Заметки для учителя."},
            {"title": "Ключевые понятия", "kind": "concepts",
             "bullet_points": ["Первое понятие", "Второе понятие", "Третье понятие",
                               "Четвёртое понятие", "Пятое понятие", "Шестое"],
             "body": "Короткое пояснение к списку понятий."},
            {"title": "Разбор основного механизма подробно", "kind": "explanation",
             "bullet_points": ["Достаточно длинный пункт, который обязательно перенесётся на вторую строку",
                               "Короткий"],
             "body": "Развёрнутое объяснение в два предложения. Оно должно поместиться под карточками."},
            {"title": "Формула", "kind": "formula",
             "bullet_points": ["$a^{2} + b^{2} = c^{2}$", "Чтение формулы"]},
            {"title": "Пример", "kind": "example",
             "bullet_points": ["Условие задачи", "Первый шаг решения", "Второй шаг", "Ответ"],
             "body": "Разбор примера целиком."},
            {"title": "Задание для класса", "kind": "task",
             "bullet_points": ["Выполните первое", "Выполните второе", "Проверьте соседа"]},
            {"title": "Итоги урока", "kind": "summary",
             "bullet_points": ["Вывод один", "Вывод два", "Вывод три"],
             "body": "Краткое резюме урока."},
        ],
    }



print("\n1. Разрешение предмета в шаблон")
for subject in REQUIRED_SUBJECTS:
    tpl = subject_templates.resolve(subject)
    check(f"{subject} -> {tpl.id}", tpl.id != "general", tpl.id)

check("незнакомый предмет падает в нейтральный шаблон",
      subject_templates.resolve("Астрономия XYZ").id == "general")
check("1-4 класс без предмета -> начальные классы",
      subject_templates.resolve("", "3 класс").id == "primary_school")
check("предмет важнее класса (п.8 ТЗ)",
      subject_templates.resolve("Биология", "3 класс").id == "biology",
      subject_templates.resolve("Биология", "3 класс").id)



print("\n2. Шаблоны различаются структурно, а не только цветом")
signatures = {}
for subject in REQUIRED_SUBJECTS:
    tpl = subject_templates.resolve(subject)
    signatures[subject] = (tpl.header, tpl.card, tpl.marker, tpl.decor, tpl.cover)

unique = len(set(signatures.values()))
check(f"{unique} уникальных структурных сигнатур из {len(REQUIRED_SUBJECTS)}",
      unique == len(REQUIRED_SUBJECTS),
      "; ".join(f"{s}={sig}" for s, sig in signatures.items()) if unique != len(REQUIRED_SUBJECTS) else "")

accents = {subject_templates.resolve(s).accent for s in REQUIRED_SUBJECTS}
check("у каждого шаблона своя палитра", len(accents) == len(REQUIRED_SUBJECTS),
      f"{len(accents)} цветов")



print("\n3. Сборка PPTX и проверка вёрстки")
os.makedirs(OUT_DIR, exist_ok=True)


def audit(buf, label):
    problems = []
    prs = Presentation(buf)
    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape.left is None or shape.top is None:
                continue
            if shape.left < -Emu(1000) or shape.top < -Emu(1000):
                problems.append(f"s{n}: текст начинается за пределами слайда: {text[:30]!r}")
            right = shape.left + (shape.width or 0)
            bottom = shape.top + (shape.height or 0)
            if right > SLIDE_W_EMU + Emu(152400):
                problems.append(f"s{n}: текст выходит вправо: {text[:30]!r}")
            if bottom > SLIDE_H_EMU + Emu(152400):
                problems.append(f"s{n}: текст выходит вниз: {text[:30]!r}")
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size and run.font.size.pt < MIN_READABLE_PT:
                        problems.append(f"s{n}: {run.font.size.pt}pt — мельче порога: {text[:24]!r}")
    return problems


built = 0
for subject in REQUIRED_SUBJECTS:
    try:
        buf = build_presentation_pptx(deck_fixture(subject))
    except Exception as e:
        check(f"сборка «{subject}»", False, f"{type(e).__name__}: {e}")
        continue
    built += 1
    data = buf.getvalue()
    path = os.path.join(OUT_DIR, f"{subject_templates.resolve(subject).id}.pptx")
    with open(path, "wb") as fh:
        fh.write(data)
    problems = audit(io.BytesIO(data), subject)
    check(f"вёрстка «{subject}»", not problems,
          problems[0] if problems else f"{len(data) // 1024} КБ")

check(f"собрано {built}/{len(REQUIRED_SUBJECTS)} шаблонов", built == len(REQUIRED_SUBJECTS))

for legacy_id in ("playful", "google", "zamonaviy", "klassik", "rangli", "minimal"):
    fixture = deck_fixture("Биология")
    fixture["template"] = legacy_id
    try:
        old_deck = build_presentation_pptx(fixture)
    except Exception as e:
        check(f"старая колода «{legacy_id}»", False, f"{type(e).__name__}: {e}")
        continue
    problems = audit(io.BytesIO(old_deck.getvalue()), legacy_id)
    check(f"старая колода «{legacy_id}» собирается как раньше", not problems,
          problems[0] if problems else "ок")



print("\n3b. Внутри колоды слайды устроены по-разному")
for subject in ("Биология", "Математика", "История Таджикистана", "Физика"):
    data = build_presentation_pptx(deck_fixture(subject)).getvalue()
    prs = Presentation(io.BytesIO(data))
    used = [s._element.cSld.get("name") or "-" for s in prs.slides][1:]
    distinct = len(set(used))
    check(f"«{subject}»: {distinct} разных композиций на {len(used)} слайдов",
          distinct >= 4, ", ".join(used))
    runs = [a for a, b in zip(used, used[1:]) if a == b]
    check(f"«{subject}»: нет двух одинаковых подряд", not runs,
          ", ".join(runs) if runs else "")

_by_subject = {}
for subject in ("Биология", "Математика", "История Таджикистана"):
    prs = Presentation(io.BytesIO(build_presentation_pptx(deck_fixture(subject)).getvalue()))
    _by_subject[subject] = tuple(s._element.cSld.get("name") or "-" for s in prs.slides)
check("разные предметы дают разные последовательности композиций",
      len(set(_by_subject.values())) == len(_by_subject),
      " | ".join(f"{k}: {'>'.join(v[1:])}" for k, v in _by_subject.items()))



print("\n4. Таджикский язык")
TJ_LETTERS = "ӣҷҳқӯғ"
tj = deck_fixture("Биология", language="Таджикский")
tj["title"] = "Сохти ҳуҷайра — ғояи асосӣ"
tj["slides"][0]["bullet_points"] = ["Мафҳумҳои асосӣ", "Ғизо ва қувва", "Дарозии ӯ"]
tj["slides"][0]["body"] = ("Ҳуҷайра воҳиди асосии организмҳои зинда аст: ӯ ғизо мегирад, "
                           "қувва ҳосил мекунад ва тақсим мешавад.")
buf = build_presentation_pptx(tj)
data = buf.getvalue()
with open(os.path.join(OUT_DIR, "tajik_biology.pptx"), "wb") as fh:
    fh.write(data)

prs = Presentation(io.BytesIO(data))
all_text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                     if sh.has_text_frame)
check("таджикские буквы дошли до PPTX без потерь",
      all(ch in all_text for ch in TJ_LETTERS),
      "".join(ch for ch in TJ_LETTERS if ch not in all_text) or "все на месте")

fonts = {run.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame
         for p in sh.text_frame.paragraphs for run in p.runs if run.font.name}
unsafe = fonts - subject_templates._TAJIK_SAFE - {None}
check("все шрифты таджикской колоды умеют рисовать ӣҷҳқӯғ",
      not unsafe, ", ".join(sorted(unsafe)) if unsafe else ", ".join(sorted(fonts)))

prim_ru = resolve_deck_theme({"subject": "Начальные классы", "language": "Русский"})
prim_tj = resolve_deck_theme({"subject": "Начальные классы", "language": "Таджикский"})
check("декоративный шрифт сохраняется для русского",
      prim_ru["title_font"] == "Comic Sans MS", prim_ru["title_font"])
check("…и заменяется для таджикского",
      prim_tj["title_font"] in subject_templates._TAJIK_SAFE, prim_tj["title_font"])



print("\n5. Явный выбор шаблона не подменяется (п.14 ТЗ)")
legacy = resolve_deck_theme({"subject": "Биология", "template": "google"})
check("выбранный в мастере «google» реально применяется",
      legacy.get("subject_template") == "" and legacy["header"] == "underline",
      str(legacy.get("subject_template")))
auto = resolve_deck_theme({"subject": "Биология", "template": ""})
check("без явного выбора применяется предметный шаблон",
      auto.get("subject_template") == "biology", str(auto.get("subject_template")))
by_id = resolve_deck_theme({"subject": "Биология", "template": "chemistry"})
check("предметный шаблон можно выбрать вручную",
      by_id.get("subject_template") == "chemistry", str(by_id.get("subject_template")))



print("\n6. Подбор изображений")
slide = {"title": "Строение клетки", "image_query": "plant cell",
         "bullet_points": ["Мембрана", "Ядро", "Митохондрии"]}
qs = build_queries(slide, "Строение клетки", "Биология", "8 класс")
check("запрос уточняется предметом", any("cell structure" in q or "labeled" in q for q in qs),
      qs[0])
check("исходная фраза модели остаётся последней", qs[-1] == "plant cell", qs[-1])

generic = {"title": "Введение", "image_query": "biology"}
qg = build_queries(generic, "Закон Ома", "Физика", "8 класс")
check("слишком общий запрос перестраивается вокруг темы",
      "biology" not in qg[0], qg[0])

ohm = {"title": "Закон Ома", "image_query": "electric circuit"}
qo = build_queries(ohm, "Закон Ома", "Физика", "8 класс")
check("ветка предмета добавляет свои уточнения",
      any("circuit diagram" in q for q in qo), qo[0])

egypt = {"title": "Пирамиды", "image_query": "egyptian pyramids"}
qe = build_queries(egypt, "Древний Египет", "Всемирная история", "5 класс")
check("история просит исторический материал",
      any("historical" in q or "museum" in q for q in qe), qe[0])



print("\n7. Фронтенд повторяет бэкенд один в один")
MIRROR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "frontend", "src", "lib", "subject-templates.ts")

if not os.path.exists(MIRROR):
    check("frontend/src/lib/subject-templates.ts на месте", False, MIRROR)
else:
    import re

    ts = io.open(MIRROR, encoding="utf-8").read()
    ts_ids = set(re.findall(r'id: "([a-z_]+)"', ts))
    py_ids = set(subject_templates.TEMPLATES)
    check("тот же набор id", ts_ids == py_ids,
          f"только в TS: {sorted(ts_ids - py_ids)}; только в Python: {sorted(py_ids - ts_ids)}")

    pattern = re.compile(
        r'id: "([a-z_]+)".*?accent: "(#[0-9A-Fa-f]{6})".*?'
        r'header: "([a-z]+)", card: "([a-z]+)", marker: "([a-z]+)", decor: "([a-z]+)",'
        r'\s*cover: "([a-z]+)"',
        re.S)
    drift = []
    compared = 0
    for tid, accent, header, card, marker, decor, cover in pattern.findall(ts):
        tpl = subject_templates.TEMPLATES.get(tid)
        if tpl is None:
            continue
        compared += 1
        for field, ts_val, py_val in (("accent", accent.upper(), tpl.accent.upper()),
                                      ("header", header, tpl.header),
                                      ("card", card, tpl.card),
                                      ("marker", marker, tpl.marker),
                                      ("decor", decor, tpl.decor),
                                      ("cover", cover, tpl.cover)):
            if ts_val != py_val:
                drift.append(f"{tid}.{field}: TS={ts_val} Python={py_val}")
    check(f"совпадают палитра и структура (сверено {compared})",
          bool(compared) and not drift, "; ".join(drift[:3]) or f"{compared} шаблонов")


print(f"\nФайлы для просмотра: {OUT_DIR}")
if FAIL:
    print(f"\n{len(FAIL)} ПРОВАЛЕНО:")
    for name in FAIL:
        print("   - " + name)
    sys.exit(1)
print("\nВсе проверки пройдены.")
