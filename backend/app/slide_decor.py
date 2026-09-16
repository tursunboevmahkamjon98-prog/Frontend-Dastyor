
from __future__ import annotations

import io
import math

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt, Emu

from app import pptx_shapes


_SLIDE_W_IN = 13.333
_SLIDE_H_IN = 7.5



_PATTERN_PPI = 100
_pattern_cache: dict[tuple, bytes] = {}


def _pattern_png(motif: str, tint: RGBColor, bg: tuple[int, int, int]) -> bytes | None:
    key = (motif, int(tint[0]), int(tint[1]), int(tint[2]), bg)
    hit = _pattern_cache.get(key)
    if hit is not None:
        return hit

    w = int(_SLIDE_W_IN * _PATTERN_PPI)
    h = int(_SLIDE_H_IN * _PATTERN_PPI)
    ink = (int(tint[0]), int(tint[1]), int(tint[2]))
    img = Image.new("RGB", (w, h), tuple(bg))
    d = ImageDraw.Draw(img)

    if motif == "grid":
        step = int(0.42 * _PATTERN_PPI)
        for x in range(step, w, step):
            d.line([(x, 0), (x, h)], fill=ink, width=1)
        for y in range(step, h, step):
            d.line([(0, y), (w, y)], fill=ink, width=1)
    elif motif == "dots":
        step = int(0.5 * _PATTERN_PPI)
        r = 1
        for y in range(step, h, step):
            for x in range(step, w, step):
                d.ellipse([x - r, y - r, x + r, y + r], fill=ink)
    else:
        return None

    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    data = buf.getvalue()
    _pattern_cache[key] = data
    return data


def _add_pattern(slide, motif: str, tint: RGBColor, bg: tuple[int, int, int]) -> bool:
    data = _pattern_png(motif, tint, bg)
    if not data:
        return False
    slide.shapes.add_picture(io.BytesIO(data), 0, 0,
                             Inches(_SLIDE_W_IN), Inches(_SLIDE_H_IN))
    return True




def _mix(color: RGBColor, bg: tuple[int, int, int], amount: float) -> RGBColor:
    amount = max(0.0, min(1.0, amount))
    return RGBColor(
        int(round(color[0] + (bg[0] - color[0]) * amount)),
        int(round(color[1] + (bg[1] - color[1]) * amount)),
        int(round(color[2] + (bg[2] - color[2]) * amount)),
    )


def _shape(slide, shape_type, x_in, y_in, w_in, h_in, fill=None,
           line=None, line_pt=1.0, rotation=None):
    shp = slide.shapes.add_shape(
        shape_type, Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.fill.solid()
        shp.line.fill.fore_color.rgb = line
        shp.line.width = Pt(line_pt)
    if rotation is not None:
        shp.rotation = rotation
    return pptx_shapes.flatten(shp)


def _glyph(slide, text, x_in, y_in, size_pt, color, font="Arial", bold=False):
    box = slide.shapes.add_textbox(
        Inches(x_in), Inches(y_in), Inches(size_pt / 50.0), Inches(size_pt / 50.0))
    tf = box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size_pt)
    p.font.bold = bold
    p.font.color.rgb = color
    return box




def _grid(slide, tint, strong, i, bg):
    step = 0.42
    n_cols = int(_SLIDE_W_IN / step)
    n_rows = int(_SLIDE_H_IN / step)
    for c in range(1, n_cols):
        _shape(slide, MSO_SHAPE.RECTANGLE, c * step, 0, 0.006, _SLIDE_H_IN, fill=tint)
    for r in range(1, n_rows):
        _shape(slide, MSO_SHAPE.RECTANGLE, 0, r * step, _SLIDE_W_IN, 0.006, fill=tint)


def _symbols(slide, tint, strong, i, bg):
    marks = ("+", "−", "×", "÷", "=", "√", "ƒ", "∑")
    spots = ((0.15, 1.9), (0.22, 3.4), (0.12, 5.1), (0.3, 6.4),
             (12.7, 2.1), (12.55, 3.7), (12.8, 5.3), (12.6, 6.5))
    for n, (x, y) in enumerate(spots):
        _glyph(slide, marks[(n + i) % len(marks)], x, y, 30, tint, bold=True)


def _construction(slide, tint, strong, i, bg):
    for n, r in enumerate((1.5, 2.35, 3.2)):
        _shape(slide, MSO_SHAPE.ARC, -r + 0.35, -r + 0.2, r * 2, r * 2,
               fill=None, line=tint if n else strong, line_pt=1.0)
    _shape(slide, MSO_SHAPE.RECTANGLE, 12.45, 0.0, 0.01, _SLIDE_H_IN, fill=tint)
    _shape(slide, MSO_SHAPE.RIGHT_TRIANGLE, 12.5, 6.15, 0.75, 0.75,
           fill=None, line=tint, line_pt=1.0)


def _wave(slide, tint, strong, i, bg):
    y = 7.12
    span, amp = 0.62, 0.17
    x = 0.9
    up = True
    while x < 12.2:
        _shape(slide, MSO_SHAPE.ARC, x, y - (amp if up else 0), span, amp * 2,
               fill=None, line=tint, line_pt=1.25, rotation=0 if up else 180)
        x += span
        up = not up
    _shape(slide, MSO_SHAPE.RIGHT_ARROW, 12.5, 3.3, 0.72, 0.2, fill=tint)


def _hexlattice(slide, tint, strong, i, bg):
    d = 0.52
    x0 = 12.55
    for n in range(6):
        offset = 0.0 if n % 2 == 0 else 0.24
        _shape(slide, MSO_SHAPE.HEXAGON, x0 - offset, 0.75 + n * (d * 0.86), d, d,
               fill=None, line=tint if n % 2 else strong, line_pt=1.1)
    _shape(slide, MSO_SHAPE.HEXAGON, 0.12, 6.45, 0.44, 0.44, fill=None,
           line=tint, line_pt=1.1)


def _organic(slide, tint, strong, i, bg):
    _shape(slide, MSO_SHAPE.OVAL, -1.5, -1.1, 3.4, 2.6, fill=tint)
    _shape(slide, MSO_SHAPE.OVAL, -0.8, -0.6, 1.9, 1.5, fill=strong)
    _shape(slide, MSO_SHAPE.OVAL, 12.15, 5.85, 2.6, 2.4, fill=tint)
    _shape(slide, MSO_SHAPE.OVAL, 12.55, 6.5, 1.2, 1.1, fill=strong)


def _contour(slide, tint, strong, i, bg):
    for n in range(5):
        r = 0.75 + n * 0.52
        _shape(slide, MSO_SHAPE.OVAL, -r * 0.55, 7.35 - r, r * 1.9, r * 1.35,
               fill=None, line=tint if n % 2 else strong, line_pt=1.0)


def _timeline(slide, tint, strong, i, bg):
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.9, 7.2, 11.5, 0.012, fill=strong)
    for n in range(13):
        x = 0.9 + n * (11.5 / 12)
        tall = n % 3 == 0
        _shape(slide, MSO_SHAPE.RECTANGLE, x, 7.2 - (0.11 if tall else 0.06),
               0.012, 0.11 if tall else 0.06, fill=strong if tall else tint)
    _shape(slide, MSO_SHAPE.OVAL, 0.9 + (i % 12) * (11.5 / 12) - 0.045, 7.155,
           0.1, 0.1, fill=strong)


def _dots(slide, tint, strong, i, bg):
    step = 0.5
    for r in range(1, int(_SLIDE_H_IN / step)):
        for c in range(1, int(_SLIDE_W_IN / step)):
            _shape(slide, MSO_SHAPE.OVAL, c * step, r * step, 0.028, 0.028, fill=tint)


def _ruled(slide, tint, strong, i, bg):
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.62, 0.0, 0.014, _SLIDE_H_IN, fill=strong)
    y = 1.55
    while y < 7.1:
        _shape(slide, MSO_SHAPE.RECTANGLE, 0.68, y, 0.16, 0.008, fill=tint)
        y += 0.3


def _ornament(slide, tint, strong, i, bg):
    def _chain(y, count, x0, gap):
        for n in range(count):
            _shape(slide, MSO_SHAPE.DIAMOND, x0 + n * gap, y, 0.1, 0.1,
                   fill=strong if n % 2 == 0 else tint)
    _chain(7.22, 9, 5.85, 0.19)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.12, 1.6, 0.012, 5.2, fill=tint)
    _shape(slide, MSO_SHAPE.DIAMOND, 0.07, 4.1, 0.12, 0.12, fill=strong)


def _letters(slide, tint, strong, i, bg):
    pairs = (("Aa", 0.05, 5.55), ("Bb", 12.35, 0.35), ("Cc", 12.5, 5.9))
    for n, (txt, x, y) in enumerate(pairs):
        _glyph(slide, txt, x, y, 54 if n == 0 else 40, tint, bold=True)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 12.45, 3.15, 0.72, 0.5,
           fill=None, line=tint, line_pt=1.1)


def _girih(slide, tint, strong, i, bg):
    y = 0.55
    n = 0
    while y < 7.0:
        _shape(slide, MSO_SHAPE.RECTANGLE, 12.62, y, 0.38, 0.38,
               fill=None, line=tint, line_pt=1.0, rotation=45)
        if n % 2 == 0:
            _shape(slide, MSO_SHAPE.DIAMOND, 12.71, y + 0.09, 0.2, 0.2, fill=strong)
        y += 0.56
        n += 1
    _shape(slide, MSO_SHAPE.RECTANGLE, 12.55, 0.0, 0.012, _SLIDE_H_IN, fill=tint)


def _network(slide, tint, strong, i, bg):
    nodes = ((0.28, 1.85), (0.62, 2.75), (0.2, 3.7), (0.66, 4.65), (0.3, 5.5))
    for n in range(len(nodes) - 1):
        (x1, y1), (x2, y2) = nodes[n], nodes[n + 1]
        length = math.hypot(x2 - x1, y2 - y1)
        angle = math.degrees(math.atan2(y2 - y1, x2 - x1))
        _shape(slide, MSO_SHAPE.RECTANGLE, (x1 + x2) / 2 - length / 2,
               (y1 + y2) / 2, length, 0.01, fill=tint, rotation=angle)
    for n, (x, y) in enumerate(nodes):
        d = 0.19 if n % 2 else 0.13
        _shape(slide, MSO_SHAPE.OVAL, x - d / 2, y - d / 2, d, d,
               fill=strong if n % 2 else tint)


def _cycle(slide, tint, strong, i, bg):
    _shape(slide, MSO_SHAPE.DONUT, 12.3, 2.85, 1.0, 1.0, fill=None,
           line=tint, line_pt=1.2)
    for n in range(3):
        _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, 12.74, 2.78, 0.16, 0.16,
               fill=strong, rotation=n * 120)
    _shape(slide, MSO_SHAPE.OVAL, -0.9, 6.6, 2.4, 1.7, fill=tint)


_MOTIFS = {
    "grid": _grid,
    "symbols": _symbols,
    "construction": _construction,
    "wave": _wave,
    "hexlattice": _hexlattice,
    "organic": _organic,
    "contour": _contour,
    "timeline": _timeline,
    "dots": _dots,
    "ruled": _ruled,
    "ornament": _ornament,
    "letters": _letters,
    "girih": _girih,
    "network": _network,
    "cycle": _cycle,
}




def _cover_axis(slide, accent, tint, strong, bg):
    ox, oy = 2.35, 6.05
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.35, oy, 4.3, 0.014, fill=strong)
    _shape(slide, MSO_SHAPE.RECTANGLE, ox, 4.0, 0.014, 3.3, fill=strong)
    for n in range(-3, 5):
        if n:
            _shape(slide, MSO_SHAPE.RECTANGLE, ox + n * 0.5, oy - 0.05, 0.012, 0.1, fill=tint)
            _shape(slide, MSO_SHAPE.RECTANGLE, ox - 0.05, oy + n * 0.38, 0.1, 0.012, fill=tint)
    for n in range(-8, 9):
        x = n * 0.24
        y = -(x * x) * 0.26
        _shape(slide, MSO_SHAPE.OVAL, ox + x, oy + y, 0.05, 0.05, fill=accent)


def _cover_construction(slide, accent, tint, strong, bg):
    for n, r in enumerate((1.1, 1.75, 2.4)):
        _shape(slide, MSO_SHAPE.ARC, 1.9 - r, 6.2 - r, r * 2, r * 2,
               fill=None, line=strong if n == 1 else tint, line_pt=1.2)
    _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, 1.05, 4.95, 1.85, 1.35,
           fill=None, line=accent, line_pt=1.5)
    _shape(slide, MSO_SHAPE.OVAL, 1.86, 6.16, 0.09, 0.09, fill=accent)


def _cover_trajectory(slide, accent, tint, strong, bg):
    for n in range(16):
        t = n / 15.0
        x = 0.45 + t * 4.3
        y = 6.95 - (2.6 * t - 2.35 * t * t) * 1.15
        d = 0.07 + 0.05 * (1 - abs(0.5 - t) * 2)
        _shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=strong if n % 3 else accent)
    _shape(slide, MSO_SHAPE.RIGHT_ARROW, 0.5, 6.55, 0.95, 0.16, fill=accent, rotation=-38)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.35, 7.02, 4.6, 0.014, fill=tint)


def _cover_molecule(slide, accent, tint, strong, bg):
    d = 1.05
    for n in range(3):
        _shape(slide, MSO_SHAPE.HEXAGON, 0.55 + n * (d * 0.78), 5.35 + (n % 2) * 0.45,
               d, d, fill=None, line=accent if n == 1 else strong, line_pt=1.4)
    for n, (x, y) in enumerate(((0.4, 5.25), (3.05, 6.55))):
        _shape(slide, MSO_SHAPE.OVAL, x, y, 0.22, 0.22, fill=accent if n else tint)
    _shape(slide, MSO_SHAPE.RECTANGLE, 3.2, 5.6, 0.75, 0.012, fill=tint, rotation=-28)


def _cover_organic(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.OVAL, -1.1, 4.45, 4.6, 3.4, fill=tint)
    _shape(slide, MSO_SHAPE.OVAL, 0.15, 5.15, 2.5, 2.0, fill=strong)
    _shape(slide, MSO_SHAPE.OVAL, 0.95, 5.75, 0.8, 0.72, fill=accent)
    _shape(slide, MSO_SHAPE.OVAL, 2.75, 6.45, 0.95, 0.8, fill=tint)


def _cover_globe(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.OVAL, 0.55, 4.5, 3.4, 3.4, fill=None, line=strong, line_pt=1.4)
    for w in (0.9, 1.9, 2.8):
        _shape(slide, MSO_SHAPE.OVAL, 0.55 + (3.4 - w) / 2, 4.5, w, 3.4,
               fill=None, line=tint, line_pt=1.0)
    for n in range(1, 4):
        _shape(slide, MSO_SHAPE.RECTANGLE, 0.62, 4.5 + n * 0.85, 3.26, 0.012, fill=tint)
    _shape(slide, MSO_SHAPE.OVAL, 2.05, 5.85, 0.16, 0.16, fill=accent)


def _cover_banner(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.95, 6.15, 4.9, 0.03, fill=accent)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.95, 6.26, 4.9, 0.012, fill=tint)
    for n in range(7):
        x = 0.95 + n * 0.78
        _shape(slide, MSO_SHAPE.RECTANGLE, x, 6.36, 0.014, 0.16 if n % 2 == 0 else 0.09,
               fill=strong if n % 2 == 0 else tint)
    _shape(slide, MSO_SHAPE.DIAMOND, 3.3, 5.82, 0.22, 0.22, fill=accent)


def _cover_terminal(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 4.75, 4.5, 2.45,
           fill=None, line=strong, line_pt=1.3)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.55, 5.2, 4.5, 0.012, fill=tint)
    for n, c in enumerate((accent, strong, tint)):
        _shape(slide, MSO_SHAPE.OVAL, 0.78 + n * 0.26, 4.92, 0.15, 0.15, fill=c)
    for n, w in enumerate((2.6, 3.4, 1.9, 3.0)):
        _shape(slide, MSO_SHAPE.RECTANGLE, 0.85 + (0.35 if n in (1, 3) else 0),
               5.5 + n * 0.36, w, 0.06, fill=tint)


def _cover_page(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.6, 4.55, 3.6, 2.75, fill=None, line=strong, line_pt=1.2)
    _shape(slide, MSO_SHAPE.RECTANGLE, 1.15, 4.55, 0.016, 2.75, fill=accent)
    y = 4.85
    n = 0
    while y < 7.15:
        _shape(slide, MSO_SHAPE.RECTANGLE, 1.3, y, 2.6 if n % 3 else 1.7, 0.012, fill=tint)
        y += 0.33
        n += 1


def _cover_frontispiece(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.65, 4.6, 3.9, 2.6, fill=None, line=strong, line_pt=1.2)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.85, 4.8, 3.5, 2.2, fill=None, line=tint, line_pt=0.9)
    for x, y in ((0.58, 4.53), (4.42, 4.53), (0.58, 7.13), (4.42, 7.13)):
        _shape(slide, MSO_SHAPE.DIAMOND, x, y, 0.16, 0.16, fill=accent)
    for n in range(5):
        _shape(slide, MSO_SHAPE.DIAMOND, 1.95 + n * 0.24, 6.75, 0.1, 0.1,
               fill=accent if n == 2 else tint)


def _cover_bubble(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 4.75, 2.7, 1.5, fill=tint)
    _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, 0.95, 6.2, 0.4, 0.4, fill=tint, rotation=180)
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 2.35, 5.95, 2.3, 1.25,
           fill=None, line=accent, line_pt=1.4)
    for n in range(3):
        _shape(slide, MSO_SHAPE.OVAL, 1.0 + n * 0.42, 5.35, 0.2, 0.2, fill=strong)


def _cover_ornament(slide, accent, tint, strong, bg):
    cx, cy, s = 2.0, 5.9, 1.75
    for rot in (0, 30, 60):
        _shape(slide, MSO_SHAPE.RECTANGLE, cx - s / 2, cy - s / 2, s, s,
               fill=None, line=strong if rot == 30 else tint, line_pt=1.2, rotation=rot)
    _shape(slide, MSO_SHAPE.DIAMOND, cx - 0.28, cy - 0.28, 0.56, 0.56, fill=accent)
    for n in range(6):
        _shape(slide, MSO_SHAPE.DIAMOND, 3.3 + n * 0.32, cy - 0.09, 0.18, 0.18,
               fill=accent if n % 2 == 0 else tint)


def _cover_nodes(slide, accent, tint, strong, bg):
    pts = ((0.85, 6.6), (1.75, 5.35), (2.85, 6.15), (3.9, 5.15), (3.35, 7.0))
    for n in range(len(pts)):
        for m in range(n + 1, len(pts)):
            if (n + m) % 2:
                continue
            (x1, y1), (x2, y2) = pts[n], pts[m]
            length = math.hypot(x2 - x1, y2 - y1)
            angle = math.degrees(math.atan2(y2 - y1, x2 - x1))
            _shape(slide, MSO_SHAPE.RECTANGLE, (x1 + x2) / 2 - length / 2,
                   (y1 + y2) / 2, length, 0.012, fill=tint, rotation=angle)
    for n, (x, y) in enumerate(pts):
        d = (0.46, 0.3, 0.38, 0.26, 0.34)[n]
        _shape(slide, MSO_SHAPE.OVAL, x - d / 2, y - d / 2, d, d,
               fill=accent if n == 0 else strong)


def _cover_cycle(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.DONUT, 0.85, 4.7, 2.7, 2.7, fill=None, line=strong, line_pt=1.5)
    for n in range(3):
        _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, 2.06, 4.55, 0.3, 0.3,
               fill=accent, rotation=n * 120)
    _shape(slide, MSO_SHAPE.OVAL, 1.72, 5.65, 0.95, 0.8, fill=tint, rotation=-25)
    _shape(slide, MSO_SHAPE.RECTANGLE, 3.85, 5.2, 0.014, 2.0, fill=tint)


def _cover_notebook(slide, accent, tint, strong, bg):
    _shape(slide, MSO_SHAPE.SUN, 0.35, 5.35, 0.85, 0.85, fill=strong)
    _shape(slide, MSO_SHAPE.HEART, 3.55, 6.35, 0.5, 0.45, fill=accent, rotation=-12)
    _shape(slide, MSO_SHAPE.STAR_5_POINT, 1.6, 6.6, 0.45, 0.45, fill=tint)
    _shape(slide, MSO_SHAPE.RECTANGLE, 1.1, 6.15, 2.2, 0.045, fill=accent, rotation=-2)


_COVERS = {
    "axis": _cover_axis,
    "construction": _cover_construction,
    "trajectory": _cover_trajectory,
    "molecule": _cover_molecule,
    "organic": _cover_organic,
    "globe": _cover_globe,
    "banner": _cover_banner,
    "terminal": _cover_terminal,
    "page": _cover_page,
    "frontispiece": _cover_frontispiece,
    "bubble": _cover_bubble,
    "ornament": _cover_ornament,
    "nodes": _cover_nodes,
    "cycle": _cover_cycle,
    "notebook": _cover_notebook,
}


def draw_cover(slide, cover: str, accent: RGBColor, support: RGBColor,
               bg: tuple[int, int, int]) -> None:
    fn = _COVERS.get(str(cover or "").strip().lower())
    if fn is None:
        return
    tint = _mix(support, bg, 0.5)
    strong = _mix(support, bg, 0.2)
    accent_soft = _mix(accent, bg, 0.35)
    try:
        fn(slide, accent_soft, tint, strong, bg)
    except Exception:
        pass


def draw(slide, motif: str, support: RGBColor, bg: tuple[int, int, int],
         index: int = 0, on_cover: bool = False) -> None:
    name = str(motif or "").strip().lower()
    fn = _MOTIFS.get(name)
    if fn is None:
        return
    if on_cover and name in ("timeline", "wave", "ornament"):
        return
    tint = _mix(support, bg, 0.62 if on_cover else 0.74)
    strong = _mix(support, bg, 0.34 if on_cover else 0.5)
    try:
        if _add_pattern(slide, name, tint, bg):
            return
        fn(slide, tint, strong, index, bg)
    except Exception:
        pass
