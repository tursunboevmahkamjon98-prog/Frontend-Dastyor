import io
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from app.subject_theme import get_subject_accent_rgb
from app.konspekt_templates import get_template

_GROUP_LABEL_RE = re.compile(
    r"^\s*(?:Группа|Гурӯҳи?|Group|Guruh)\s*\d+\s*[:.\-—]\s*", re.IGNORECASE
)


def _strip_group_label(task: str) -> str:
    return _GROUP_LABEL_RE.sub("", str(task), count=1)


def _set_cell_shading(cell, color_hex: str):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color_hex,
    })
    shading.append(shd)


def _set_paragraph_shading(paragraph, color_hex: str):
    pPr = paragraph._element.get_or_add_pPr()
    shd = pPr.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color_hex,
    })
    pPr.append(shd)


def _add_bottom_border(paragraph, color_hex: str):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pBdr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single',
        qn('w:sz'): '12',
        qn('w:space'): '4',
        qn('w:color'): color_hex,
    })
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_colored_paragraph(doc, text, font_size, color_rgb, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = color_rgb
    run.bold = bold
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    return p


def _add_bullet_runs(p, text, accent_rgb):
    colon_idx = text.find(':')
    if 0 < colon_idx <= 60:
        term, rest = text[:colon_idx + 1], text[colon_idx + 1:]
        run = p.add_run(term)
        run.bold = True
        run.font.color.rgb = accent_rgb
    else:
        term, rest = '', text
    if rest:
        run = p.add_run(rest)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')


def _add_bullet(doc, text, color_rgb):
    p = doc.add_paragraph(style='List Bullet')
    _add_bullet_runs(p, text, color_rgb)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    return p


def _add_section_number(doc, index, color_rgb):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f'  {index + 1}  ')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.bold = True
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    hex_str = str(color_rgb)
    _set_paragraph_shading(p, hex_str)
    return p


def _add_formula_card(doc, formula: str, explanation: str, color_rgb: RGBColor):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    _set_paragraph_shading(p, 'F8FAFC')
    run = p.add_run(formula)
    run.font.size = Pt(17)
    run.font.name = 'Cambria Math'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
    run.font.color.rgb = color_rgb
    run.bold = True
    _add_bottom_border(p, str(color_rgb))

    if explanation:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(5)
        _set_paragraph_shading(p2, 'F8FAFC')
        run2 = p2.add_run(explanation)
        run2.font.size = Pt(10)
        run2.italic = True
        run2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        run2.font.name = 'Calibri'
        run2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    else:
        p.paragraph_format.space_after = Pt(5)


def _add_code_card(doc, code: str, language: str, explanation: str, accent_rgb: RGBColor):
    badge = doc.add_paragraph()
    badge.paragraph_format.space_before = Pt(6)
    badge.paragraph_format.space_after = Pt(0)
    run = badge.add_run(f'  {language.upper()}  ' if language else '  CODE  ')
    run.font.size = Pt(8.5)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = 'Calibri'
    _set_paragraph_shading(badge, str(accent_rgb))

    p = doc.add_paragraph()
    _set_paragraph_shading(p, '1E293B')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Pt(6)
    lines = code.splitlines() or ['']
    for i, line in enumerate(lines):
        run = p.add_run(line if line else ' ')
        run.font.size = Pt(10)
        run.font.name = 'Consolas'
        run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        if i < len(lines) - 1:
            run.add_break()

    if explanation:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(8)
        run2 = p2.add_run(explanation)
        run2.font.size = Pt(10)
        run2.italic = True
        run2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        run2.font.name = 'Calibri'
        run2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    else:
        p.paragraph_format.space_after = Pt(8)


def _add_illustration(doc, image_path: str, caption: str, attribution: str, width_inches: float = 5.0):
    import os
    full_path = os.path.join(os.path.dirname(__file__), "..", image_path.lstrip("/"))
    if not os.path.exists(full_path):
        return
    doc.add_picture(full_path, width=Inches(width_inches))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(caption)
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    if attribution:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(attribution)
        run.font.size = Pt(8)
        run.italic = True
        run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')


def _add_lesson_image_block(doc, image: dict, label: str, accent_rgb: RGBColor):
    import os
    full_path = os.path.join(os.path.dirname(__file__), "..", image.get("path", "").lstrip("/"))
    if not image.get("path") or not os.path.exists(full_path):
        return []
    start = len(doc.paragraphs)
    tag_p = doc.add_paragraph()
    tag_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag_p.paragraph_format.space_before = Pt(8)
    tag_p.paragraph_format.space_after = Pt(2)
    tag_run = tag_p.add_run(label.upper())
    tag_run.font.name = 'Calibri'
    tag_run.font.size = Pt(9)
    tag_run.bold = True
    tag_run.font.color.rgb = accent_rgb
    _add_illustration(doc, image["path"], image.get("caption", ""), image.get("credit", ""),
                      width_inches=4.0)
    explanation = str(image.get("explanation") or "").strip()
    if explanation:
        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.left_indent = Cm(1.5)
        note.paragraph_format.right_indent = Cm(1.5)
        note.paragraph_format.space_after = Pt(8)
        run = note.add_run(explanation)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    return doc.paragraphs[start:]


def _clear_cell(cell):
    cell.text = ''
    return cell.paragraphs[0]


def _fill_concept_card(outer_cell, card: dict, accent_rgb: RGBColor):
    _clear_cell(outer_cell)
    inner = outer_cell.add_table(rows=2, cols=1)
    inner.autofit = True

    header_cell = inner.cell(0, 0)
    _set_cell_shading(header_cell, str(accent_rgb))
    hp = _clear_cell(header_cell)
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_before = Pt(4)
    hp.paragraph_format.space_after = Pt(0)
    hrun = hp.add_run(card.get('title', ''))
    hrun.font.bold = True
    hrun.font.size = Pt(12)
    hrun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    hrun.font.name = 'Calibri'
    tag = card.get('tag')
    if tag:
        tp = header_cell.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tp.paragraph_format.space_after = Pt(4)
        trun = tp.add_run(tag)
        trun.font.size = Pt(9)
        trun.italic = True
        trun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        trun.font.name = 'Calibri'
    else:
        hp.paragraph_format.space_after = Pt(4)

    body_cell = inner.cell(1, 0)
    _set_cell_shading(body_cell, 'F8FAFC')
    bp = _clear_cell(body_cell)
    bp.paragraph_format.space_before = Pt(0)
    bp.paragraph_format.space_after = Pt(0)
    bp.add_run('').font.size = Pt(1)

    table_rows = card.get('table')
    if table_rows:
        n_r, n_c = len(table_rows), max(len(r) for r in table_rows)
        tt = body_cell.add_table(rows=n_r, cols=n_c)
        tt.autofit = True
        for ri, row in enumerate(table_rows):
            for ci in range(n_c):
                val = row[ci] if ci < len(row) else ''
                cell = tt.cell(ri, ci)
                cp = _clear_cell(cell)
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                crun = cp.add_run(str(val))
                crun.font.size = Pt(10)
                crun.font.name = 'Calibri'
                if ri == 0:
                    crun.bold = True
                    _set_cell_shading(cell, 'E2E8F0')

    formula = card.get('formula')
    if formula:
        fp = body_cell.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(4)
        frun = fp.add_run(formula)
        frun.font.bold = True
        frun.font.size = Pt(13)
        frun.font.name = 'Cambria Math'
        frun._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
        frun.font.color.rgb = accent_rgb

    note = card.get('note')
    if note:
        np = body_cell.add_paragraph()
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
        np.paragraph_format.space_before = Pt(4)
        np.paragraph_format.space_after = Pt(6)
        nrun = np.add_run(note)
        nrun.font.size = Pt(8.5)
        nrun.italic = True
        nrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        nrun.font.name = 'Calibri'


def _add_important_note(doc, text: str, accent_rgb: RGBColor):
    p = doc.add_paragraph()
    _set_paragraph_shading(p, 'FEF2F2')
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._element.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    left = pBdr.makeelement(qn('w:left'), {qn('w:val'): 'single', qn('w:sz'): '18', qn('w:space'): '6', qn('w:color'): str(accent_rgb)})
    pBdr.append(left)
    pPr.append(pBdr)
    run = p.add_run('☝  ')
    run.font.size = Pt(11)
    run2 = p.add_run(text)
    run2.font.size = Pt(10.5)
    run2.italic = True
    run2.font.color.rgb = RGBColor(0x7F, 0x1D, 0x1D)
    run2.font.name = 'Calibri'
    run2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')


def _add_quick_check(doc, items: list, accent_rgb: RGBColor, answer_label: str):
    for item in items:
        if not isinstance(item, dict):
            continue
        question = item.get('question', '')
        answer = item.get('answer', '')
        if not question:
            continue
        p = doc.add_paragraph()
        _set_paragraph_shading(p, 'F0FDF4')
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(0 if answer else 6)
        pPr = p._element.get_or_add_pPr()
        pBdr = pPr.makeelement(qn('w:pBdr'), {})
        left = pBdr.makeelement(qn('w:left'), {qn('w:val'): 'single', qn('w:sz'): '18', qn('w:space'): '6', qn('w:color'): str(accent_rgb)})
        pBdr.append(left)
        pPr.append(pBdr)
        run = p.add_run('?  ')
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = accent_rgb
        run2 = p.add_run(question)
        run2.font.size = Pt(10.5)
        run2.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        run2.font.name = 'Calibri'
        run2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        if answer:
            p2 = doc.add_paragraph()
            _set_paragraph_shading(p2, 'F0FDF4')
            p2.paragraph_format.space_after = Pt(6)
            pPr2 = p2._element.get_or_add_pPr()
            pBdr2 = pPr2.makeelement(qn('w:pBdr'), {})
            left2 = pBdr2.makeelement(qn('w:left'), {qn('w:val'): 'single', qn('w:sz'): '18', qn('w:space'): '6', qn('w:color'): str(accent_rgb)})
            pBdr2.append(left2)
            pPr2.append(pBdr2)
            run3 = p2.add_run(f'{answer_label}: ')
            run3.font.size = Pt(10)
            run3.italic = True
            run3.bold = True
            run3.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
            run3.font.name = 'Calibri'
            run4 = p2.add_run(answer)
            run4.font.size = Pt(10)
            run4.italic = True
            run4.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
            run4.font.name = 'Calibri'
            run4._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')


def _add_worked_examples(doc, items: list, accent_rgb: RGBColor, solution_label: str):
    tint = _light_tint_hex(accent_rgb, amount=0.93)
    for i, item in enumerate(items):
        if not isinstance(item, dict):
            continue
        problem = item.get('problem', '')
        solution = item.get('solution', '')
        if not problem:
            continue

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        row = table.rows[0]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(trPr.makeelement(qn('w:cantSplit'), {}))

        badge_cell, body_cell = row.cells
        badge_cell.width = Cm(1.15)
        body_cell.width = Cm(15.0)

        _set_cell_shading(badge_cell, str(accent_rgb))
        bp = _clear_cell(badge_cell)
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bp.paragraph_format.space_before = Pt(6)
        bp.paragraph_format.space_after = Pt(6)
        brun = bp.add_run(str(i + 1))
        brun.font.size = Pt(15)
        brun.bold = True
        brun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        brun.font.name = 'Calibri'
        brun._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

        _set_cell_shading(body_cell, tint)
        pp = _clear_cell(body_cell)
        pp.paragraph_format.space_before = Pt(5)
        pp.paragraph_format.space_after = Pt(2)
        pp.paragraph_format.line_spacing = 1.0
        prun = pp.add_run(problem)
        prun.font.size = Pt(10.5)
        prun.bold = True
        prun.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        prun.font.name = 'Calibri'
        prun._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

        if solution:
            sp = body_cell.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(6)
            sp.paragraph_format.line_spacing = 1.0
            srun = sp.add_run(f'{solution_label}: ')
            srun.font.size = Pt(10)
            srun.bold = True
            srun.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
            srun.font.name = 'Calibri'
            srun._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
            srun2 = sp.add_run(solution)
            srun2.font.size = Pt(10)
            srun2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            srun2.font.name = 'Calibri'
            srun2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        else:
            pp.paragraph_format.space_after = Pt(6)

        gap = doc.add_paragraph()
        gap.paragraph_format.space_before = Pt(0)
        gap.paragraph_format.space_after = Pt(0)
        gap.paragraph_format.line_spacing = 1.0
        gap.add_run('').font.size = Pt(4)


def _add_visual_block(doc, block: dict, accent_rgb: RGBColor, L: dict):
    btype = block.get('type')
    data = block.get('data') or {}
    title = block.get('title') or ''
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        run.font.name = 'Calibri'

    try:
        if btype == 'table' or btype == 'comparison':
            if btype == 'table':
                headers = data.get('headers') or []
                rows = data.get('rows') or []
            else:
                criteria = data.get('criteria') or []
                items = data.get('items') or []
                headers = [''] + criteria
                rows = [[it.get('name', '')] + list(it.get('values') or []) for it in items]
            if not headers or not rows:
                return
            n_cols = len(headers)
            table = doc.add_table(rows=1, cols=n_cols)
            table.autofit = True
            for cell, text in zip(table.rows[0].cells, headers):
                _set_cell_shading(cell, str(accent_rgb))
                cell.text = ''
                run = cell.paragraphs[0].add_run(str(text))
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.name = 'Calibri'
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                if ri % 2 == 1:
                    for cell in cells:
                        _set_cell_shading(cell, 'FEF2F2')
                for ci, cell in enumerate(cells):
                    val = row[ci] if ci < len(row) else ''
                    cell.text = ''
                    run = cell.paragraphs[0].add_run(str(val))
                    run.font.size = Pt(9.5)
                    run.font.name = 'Calibri'
            doc.add_paragraph()

        elif btype == 'timeline':
            for ev in (data.get('events') or []):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                label_bits = [b for b in [ev.get('date'), ev.get('label')] if b]
                run = p.add_run('●  ' + ' — '.join(label_bits))
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = accent_rgb
                run.font.name = 'Calibri'
                if ev.get('description'):
                    p2 = doc.add_paragraph()
                    p2.paragraph_format.left_indent = Pt(18)
                    p2.paragraph_format.space_after = Pt(5)
                    run2 = p2.add_run(ev['description'])
                    run2.font.size = Pt(10)
                    run2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    run2.font.name = 'Calibri'

        elif btype in ('flowchart', 'process'):
            for i, step in enumerate(data.get('steps') or []):
                title_p = doc.add_paragraph()
                title_p.paragraph_format.space_before = Pt(5)
                title_p.paragraph_format.space_after = Pt(1)
                title_p.paragraph_format.keep_with_next = True
                num = title_p.add_run(f'{i + 1}.  ')
                num.bold = True
                num.font.size = Pt(11)
                num.font.color.rgb = accent_rgb
                num.font.name = 'Calibri'
                head = title_p.add_run(str(step.get('title', '')))
                head.bold = True
                head.font.size = Pt(11)
                head.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
                head.font.name = 'Calibri'
                if step.get('description'):
                    desc_p = doc.add_paragraph()
                    desc_p.paragraph_format.left_indent = Cm(0.7)
                    desc_p.paragraph_format.space_after = Pt(3)
                    run2 = desc_p.add_run(str(step['description']))
                    run2.font.size = Pt(10)
                    run2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
                    run2.font.name = 'Calibri'

        elif btype == 'concept_map':
            nodes = {n.get('id'): n.get('label', n.get('id', '')) for n in (data.get('nodes') or [])}
            for edge in (data.get('edges') or []):
                p = doc.add_paragraph(style='List Bullet')
                frm = nodes.get(edge.get('from'), edge.get('from', ''))
                to = nodes.get(edge.get('to'), edge.get('to', ''))
                text = f'{frm}  →  {to}'
                if edge.get('label'):
                    text += f'  ({edge["label"]})'
                run = p.add_run(text)
                run.font.size = Pt(10.5)
                run.font.name = 'Calibri'
    except Exception:
        return


def _add_concept_card_grid(doc, cards: list, accent_rgb: RGBColor, cols: int = 2):
    if not cards:
        return
    n_rows = (len(cards) + cols - 1) // cols
    outer = doc.add_table(rows=n_rows, cols=cols)
    outer.autofit = True
    for idx in range(n_rows * cols):
        r, c = divmod(idx, cols)
        cell = outer.cell(r, c)
        if idx < len(cards):
            _fill_concept_card(cell, cards[idx], accent_rgb)
        else:
            _clear_cell(cell)
    if cols == 2 and len(cards) % cols:
        outer.cell(n_rows - 1, 0).merge(outer.cell(n_rows - 1, 1))
    doc.add_paragraph()


_DOCX_LABELS = {
    'Русский': {
        'objective': 'Цель лекции', 'lecture_plan': 'План лекции', 'references': 'Литература',
        'duration': 'Время урока', 'competencies': 'Компетенции', 'objectives': 'Цели урока',
        'key_concepts': 'Ключевые понятия', 'terms': 'Термины', 'example': 'Пример',
        'homework': 'Домашнее задание', 'summary': 'Итоги урока',
        'lesson_program': 'План урока', 'tools': 'Материалы',
        'warmup': 'Разминка', 'main_content': 'Основное содержание',
        'pair_work': 'Работа в парах', 'consolidation': 'Закрепление', 'group_work': 'Групповая работа', 'group': 'Группа',
        'assessment': 'Оценивание',
        'key_terms': 'Словарь урока', 'fun_facts': 'Интересные факты',
        'real_life_examples': 'Примеры из жизни', 'visual_aid': 'Наглядный материал',
        'common_mistakes': 'Частые ошибки',
        'formulas': 'Формулы', 'map': 'Карта', 'illustration': 'Иллюстрация', 'diagrams': 'Схемы',
        'code_blocks': 'Код', 'quick_check': 'Быстрая проверка', 'answer': 'Ответ',
        'worked_examples': 'Решённые примеры', 'solution': 'Решение',
    },
    'Таджикский': {
        'objective': 'Мақсади лексия', 'lecture_plan': 'Нақшаи лексия', 'references': 'Адабиёт',
        'duration': 'Вақти дарс', 'competencies': 'Салоҳиятҳо', 'objectives': 'Мақсадҳои дарс',
        'key_concepts': 'Мафҳумҳои асосӣ', 'terms': 'Истилоҳот', 'example': 'Намуна',
        'homework': 'Супориши хонагӣ', 'summary': 'Хулосаи дарс',
        'lesson_program': 'Барномаи дарс', 'tools': 'Воситаҳои аёнӣ',
        'warmup': 'Санҷиши дониш', 'main_content': 'Шиносоӣ бо мазмуни мавзӯъ',
        'pair_work': 'Кори дунафара', 'consolidation': 'Мустаҳкамкунии дарс', 'group_work': 'Кори гурӯҳӣ', 'group': 'Гурӯҳи',
        'assessment': 'Арзёбӣ',
        'key_terms': 'Луғати дарс', 'fun_facts': 'Фактҳои ҷолиб',
        'real_life_examples': 'Дар ҳаёти воқеӣ', 'visual_aid': 'Ёрии визуалӣ',
        'common_mistakes': 'Хатогиҳои маъмул',
        'formulas': 'Формулаҳо', 'map': 'Харита', 'illustration': 'Тасвир', 'diagrams': 'Расмҳо',
        'code_blocks': 'Код', 'quick_check': 'Санҷиши зуд', 'answer': 'Ҷавоб',
        'worked_examples': 'Мисолҳои ҳалшуда', 'solution': 'Ҳал',
    },
    'English': {
        'objective': 'Objective', 'lecture_plan': 'Lecture plan', 'references': 'References',
        'duration': 'Duration', 'competencies': 'Competencies', 'objectives': 'Objectives',
        'key_concepts': 'Key Concepts', 'terms': 'Terms', 'example': 'Example',
        'homework': 'Homework', 'summary': 'Summary',
        'lesson_program': 'Lesson Program', 'tools': 'Teaching Tools',
        'warmup': 'Warm-up', 'main_content': 'Main Content',
        'pair_work': 'Pair Work', 'consolidation': 'Consolidation', 'group_work': 'Group Work', 'group': 'Group',
        'assessment': 'Assessment',
        'key_terms': 'Lesson Glossary', 'fun_facts': 'Fun Facts',
        'real_life_examples': 'Real Life', 'visual_aid': 'Visual Aid',
        'common_mistakes': 'Common Mistakes',
        'formulas': 'Formulas', 'map': 'Map', 'illustration': 'Illustration', 'diagrams': 'Diagrams',
        'code_blocks': 'Code', 'quick_check': 'Quick Check', 'answer': 'Answer',
        'worked_examples': 'Worked Examples', 'solution': 'Solution',
    },
}
_DOCX_LABELS['Английский'] = _DOCX_LABELS['English']


def _header_numbered(doc, num, label, color, font_name):
    p = _add_colored_paragraph(doc, f'{num}. {label}', 12, color, bold=True)
    p.paragraph_format.space_before = Pt(8)
    return p


def _header_underline(doc, num, label, color, font_name):
    p = _add_colored_paragraph(doc, label.upper(), 12, RGBColor(0x1F, 0x29, 0x37), bold=True)
    p.paragraph_format.space_before = Pt(10)
    _add_bottom_border(p, str(color))
    return p


def _header_smallcaps(doc, num, label, color, font_name):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label.upper())
    run.font.size = Pt(9.5)
    run.bold = True
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    _add_bottom_border(p, 'E2E8F0')
    return p


def _header_serif(doc, num, label, color, font_name):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(f'{num}. {label}')
    run.font.size = Pt(12.5)
    run.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    _add_bottom_border(p, 'CBD5E1')
    return p


def _header_bar(doc, num, label, color, font_name):
    p = doc.add_paragraph()
    _set_paragraph_shading(p, str(color))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f'  {num}.  {label}  ')
    run.font.size = Pt(12)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = 'Calibri'
    return p


_SECTION_RENDERERS = {
    "numbered": _header_numbered,
    "underline": _header_underline,
    "smallcaps": _header_smallcaps,
    "serif": _header_serif,
    "bar": _header_bar,
}


def _light_tint_hex(rgb: RGBColor, amount: float = 0.88) -> str:
    r, g, b = rgb[0], rgb[1], rgb[2]
    tr = int(r + (255 - r) * amount)
    tg = int(g + (255 - g) * amount)
    tb = int(b + (255 - b) * amount)
    return f'{tr:02X}{tg:02X}{tb:02X}'


def _add_sidebar_panel(doc, content: dict, sections: tuple, accent_rgb: RGBColor, L: dict):
    rows_present = [(key, content.get(key)) for key in sections if content.get(key)]
    if not rows_present:
        return
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    _set_cell_shading(cell, _light_tint_hex(accent_rgb, amount=0.92))
    cell.paragraphs[0].text = ''
    first_section = True
    for key, items in rows_present:
        header_p = cell.paragraphs[0] if first_section else cell.add_paragraph()
        header_p.paragraph_format.space_before = Pt(2) if first_section else Pt(10)
        header_p.paragraph_format.space_after = Pt(2)
        run = header_p.add_run(L.get(key, key).upper())
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = accent_rgb
        run.font.name = 'Calibri'
        first_section = False
        for item in items:
            p2 = cell.add_paragraph()
            p2.paragraph_format.left_indent = Pt(10)
            p2.paragraph_format.space_after = Pt(3)
            bullet_run = p2.add_run('•  ')
            bullet_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            _add_bullet_runs(p2, str(item), accent_rgb)
            for run2 in p2.runs:
                run2.font.size = Pt(9.5)
    doc.add_paragraph()


_LECTURE_SUMMARY_LABEL_DOCX = {
    'Русский': 'Выводы', 'Таджикский': 'Хулосаҳо',
    'English': 'Conclusions',
}
_LECTURE_SUMMARY_LABEL_DOCX['Английский'] = _LECTURE_SUMMARY_LABEL_DOCX['English']


def _add_konspekt_body(doc, content: dict, L: dict):
    ACCENT = RGBColor(*get_subject_accent_rgb(content.get("subject")))
    DARK = RGBColor(0x1F, 0x29, 0x37)
    SECONDARY = RGBColor(0x6B, 0x72, 0x80)
    tmpl = get_template(content.get("template") or "zamonaviy")
    header_font = 'Cambria' if tmpl.font_family == 'serif' else 'Calibri'
    header_renderer = _SECTION_RENDERERS.get(tmpl.header_style, _header_numbered)

    _sec_num = [0]

    def _section(label, color=ACCENT):
        _sec_num[0] += 1
        return header_renderer(doc, _sec_num[0], label, color, header_font)

    def _bullet(text):
        if not tmpl.alt_row_tint:
            _add_bullet(doc, text, ACCENT)
            return
        p = doc.add_paragraph()
        _set_paragraph_shading(p, _light_tint_hex(ACCENT))
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.space_after = Pt(4)
        bullet_run = p.add_run('•  ')
        bullet_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        _add_bullet_runs(p, text, ACCENT)
        for run in p.runs:
            run.font.size = Pt(10.5)
            run.font.name = 'Calibri'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    _visuals_by_position: dict[str, list] = {}
    for _block in (content.get('visual_blocks') or []):
        if isinstance(_block, dict):
            _visuals_by_position.setdefault(_block.get('position_after') or '', []).append(_block)

    _lesson_by_anchor: dict[str, list] = {}
    for _img in (content.get('lesson_images') or []):
        if isinstance(_img, dict) and _img.get('path'):
            _lesson_by_anchor.setdefault(_img.get('position_after') or 'main_content', []).append(_img)

    def _render_lesson_images(position_key):
        for image in _lesson_by_anchor.pop(position_key, []):
            _add_lesson_image_block(doc, image, L.get('illustration', 'Иллюстрация'), ACCENT)

    def _render_visuals(position_key):
        _render_lesson_images(position_key)
        for block in _visuals_by_position.pop(position_key, []):
            _add_visual_block(doc, block, ACCENT, L)
        _render_figures(position_key)

    _figures = [f for f in (content.get('figures') or [])
                if isinstance(f, dict) and f.get('image')]
    _figure_anchors = ['key_concepts', 'main_content', 'real_life_examples', 'consolidation']
    _figures_by_anchor: dict = {}
    for _i, _fig in enumerate(_figures):
        _figures_by_anchor.setdefault(_figure_anchors[_i % len(_figure_anchors)], []).append(_fig)

    def _render_figures(position_key):
        for fig in _figures_by_anchor.pop(position_key, []):
            width = 3.6
            try:
                import os
                from PIL import Image as _PILImg
                _full = os.path.join(os.path.dirname(__file__), '..', fig['image'].lstrip('/'))
                with _PILImg.open(_full) as _im:
                    _iw, _ih = _im.size
                width = min(3.6, 2.9 * _iw / _ih)
            except Exception:
                pass
            _add_illustration(doc, fig['image'], fig.get('caption', ''), '', width_inches=width)

    _add_bottom_border(doc.add_paragraph(), str(ACCENT))

    duration = content.get('duration', '')
    if duration:
        _add_colored_paragraph(doc, f'{L["duration"]}: {duration}', 11, ACCENT, bold=True)

    if tmpl.sidebar_sections:
        _add_sidebar_panel(doc, content, tmpl.sidebar_sections, ACCENT, L)

    competencies = content.get('competencies', [])
    if competencies:
        _section(L['competencies'])
        for c in competencies:
            _bullet(c)

    is_lecture = bool(content.get('lecture_plan') or content.get('objective'))
    if is_lecture:
        L = dict(L)
        L['summary'] = _LECTURE_SUMMARY_LABEL_DOCX.get(
            str(content.get('language') or 'Русский'), L['summary'])

    objective = content.get('objective', '')
    if objective:
        _section(L.get('objective', 'Цель'))
        _bullet(objective)

    lecture_plan = [x for x in (content.get('lecture_plan') or []) if str(x).strip()]
    if lecture_plan:
        _section(L.get('lecture_plan', 'План'))
        for i, point in enumerate(lecture_plan, start=1):
            _bullet(f"{i}. {point}")

    objectives = content.get('objectives', [])
    if objectives:
        _section(L['objectives'])
        for o in objectives:
            _bullet(o)

    key_concepts = content.get('key_concepts', [])
    if key_concepts and 'key_concepts' not in tmpl.sidebar_sections:
        _section(L['key_concepts'])
        for k in key_concepts:
            _bullet(k)
    _render_visuals('key_concepts')

    formulas = content.get('formulas', [])
    if formulas:
        for f in formulas:
            if isinstance(f, dict):
                _add_formula_card(doc, f.get('formula', ''), f.get('explanation', ''), ACCENT)
            else:
                _add_formula_card(doc, str(f), '', ACCENT)

    concept_cards = content.get('concept_cards', [])
    if concept_cards:
        _add_concept_card_grid(doc, concept_cards, ACCENT)

    key_terms = content.get('key_terms', [])
    if key_terms and 'key_terms' not in tmpl.sidebar_sections:
        _section(L['key_terms'])
        for t in key_terms:
            _bullet(t)

    code_blocks = content.get('code_blocks') or []
    for cb in code_blocks:
        if isinstance(cb, dict) and cb.get('code'):
            _add_code_card(doc, cb['code'], cb.get('language', ''), cb.get('explanation', ''), ACCENT)

    def _render_real_life():
        real_life = content.get('real_life_examples', [])
        if real_life:
            _section(L['real_life_examples'])
            for r in real_life:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                _add_bullet_runs(p, r, ACCENT)
        _render_visuals('real_life_examples')

    if not is_lecture:
        _render_real_life()

    map_image = content.get('map_image', '')
    if map_image:
        locations = content.get('map_locations', [])
        caption = ', '.join(locations) if locations else ''
        _add_illustration(doc, map_image, caption, '© OpenStreetMap contributors', width_inches=4.8)

    real_image = content.get('real_image')
    if isinstance(real_image, dict) and real_image.get('path'):
        _add_illustration(doc, real_image['path'], real_image.get('caption', ''), 'Wikipedia', width_inches=2.3)

    visual = content.get('visual_aid', '')
    if visual:
        _section(L['visual_aid'])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(visual)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    _render_visuals('visual_aid')

    important_notes = content.get('important_notes') or []
    if important_notes:
        for note in important_notes:
            _add_important_note(doc, str(note), ACCENT)

    _add_bottom_border(doc.add_paragraph(), 'E5E7EB')

    lesson_program = content.get('lesson_program', '')
    if lesson_program:
        _section(L['lesson_program'])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(lesson_program)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    _render_visuals('lesson_program')

    tools = content.get('tools', [])
    if tools and 'tools' not in tmpl.sidebar_sections:
        _section(L['tools'])
        for t in tools:
            _bullet(t)
    _render_visuals('tools')

    main_content = content.get('main_content', '')
    if main_content:
        _section(L['main_content'], DARK)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(main_content)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    _render_visuals('main_content')

    if is_lecture:
        _render_real_life()

    worked_examples = content.get('worked_examples') or []
    if worked_examples:
        _section(L['worked_examples'])
        _add_worked_examples(doc, worked_examples, ACCENT, L['solution'])

    pair_work = content.get('pair_work', '')
    if pair_work:
        _section(L['pair_work'])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(pair_work)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    _render_visuals('pair_work')

    consolidation = content.get('consolidation', '')
    if consolidation:
        _section(L['consolidation'])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(consolidation)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    _render_visuals('consolidation')
    for _left in list(_figures_by_anchor):
        _render_figures(_left)

    group_work = content.get('group_work') or []
    if group_work:
        _section(L['group_work'])
        for i, task in enumerate(group_work):
            p = doc.add_paragraph()
            _set_paragraph_shading(p, 'EFF6FF')
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.space_before = Pt(4)
            run = p.add_run(f"{L['group']} {i + 1}. ")
            run.font.size = Pt(11)
            run.font.color.rgb = ACCENT
            run.bold = True
            run.font.name = 'Calibri'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
            run = p.add_run(_strip_group_label(task))
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            run.font.name = 'Calibri'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    for _leftover_key in list(_visuals_by_position):
        _render_visuals(_leftover_key)
    for _leftover_key in list(_lesson_by_anchor):
        _render_lesson_images(_leftover_key)

    quick_check = content.get('quick_check') or []
    if quick_check:
        _section(L['quick_check'])
        _add_quick_check(doc, quick_check, ACCENT, L['answer'])

    _add_bottom_border(doc.add_paragraph(), 'E5E7EB')

    summary = content.get('summary', '')
    if summary:
        p = _section(L['summary'])
        p.paragraph_format.space_before = Pt(16)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(summary)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    homework = content.get('homework', '')
    if homework:
        p = _section(L['homework'])
        p.paragraph_format.space_before = Pt(16)
        if isinstance(homework, list):
            for i, h in enumerate(homework):
                p = doc.add_paragraph()
                _set_paragraph_shading(p, 'EFF6FF')
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.space_before = Pt(4)
                run = p.add_run(f'{i + 1}. ')
                run.font.size = Pt(11)
                run.font.color.rgb = ACCENT
                run.bold = True
                run.font.name = 'Calibri'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
                run = p.add_run(h)
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                run.font.name = 'Calibri'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        else:
            p = doc.add_paragraph()
            _set_paragraph_shading(p, 'EFF6FF')
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            run = p.add_run(homework)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            run.font.name = 'Calibri'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    assessment = content.get('assessment', '')
    if assessment:
        _section(L['assessment'])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(11.5)
        run = p.add_run(assessment)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    references = [r for r in (content.get('references') or []) if str(r).strip()]
    if references:
        _section(L.get('references', 'Литература'))
        for i, ref in enumerate(references, start=1):
            _bullet(f"{i}. {ref}")



_PLAN_BG_DOCX = "EAF3FB"
_PLAN_BAR_DOCX = "8FBEDC"
_PLAN_INK_DOCX = RGBColor(0x1B, 0x44, 0x60)
_PLAN_EX_BG_DOCX = "F4F9FD"
_PLAN_MARK_DOCX = "D3E7F6"
_PLAN_FONT = "Cambria"
_PLAN_WIDTH = Cm(17.2)


def _add_thin_border(paragraph, color_hex: str):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pBdr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '4',
        qn('w:space'): '1', qn('w:color'): color_hex,
    })
    pBdr.append(bottom)
    pPr.append(pBdr)


def _plan_run(p, text, size=11, bold=False, italic=False, color=None, mark=False):
    run = p.add_run(text)
    run.font.name = _PLAN_FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), _PLAN_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if mark:
        rPr = run._element.get_or_add_rPr()
        rPr.append(rPr.makeelement(qn('w:shd'), {
            qn('w:val'): 'clear', qn('w:color'): 'auto',
            qn('w:fill'): _PLAN_MARK_DOCX,
        }))
    return run


_MATH_SPAN_DOCX = re.compile(r"\$([^$\n]{1,400}?)\$")


def _normalize_math(text) -> str:
    try:
        from app.math_render import normalize_math
        return normalize_math(text)
    except Exception:
        return str(text or "")


def _plain_to_latex(text) -> str:
    try:
        from app.math_render import plain_to_latex
        return plain_to_latex(text)
    except Exception:
        return ""


def _plan_math_run(p, latex: str) -> bool:
    from app.math_render import omml_element
    el = omml_element(latex)
    if el is None:
        return False
    p._p.append(el)
    return True


def _plan_text_math(p, text, size=11, bold=False, italic=False, color=None,
                    mark_text=False):
    raw = _normalize_math(text)
    pos = 0
    for m in _MATH_SPAN_DOCX.finditer(raw):
        chunk = raw[pos:m.start()]
        if chunk:
            r = _plan_run(p, chunk, size, bold, italic, color, mark=mark_text)
        if not _plan_math_run(p, m.group(1)):
            _plan_run(p, m.group(1), size, bold, italic, color, mark=mark_text)
        pos = m.end()
    rest = raw[pos:]
    if rest or pos == 0:
        _plan_run(p, rest if pos else raw, size, bold, italic, color, mark=mark_text)



_MONO_FONTS = {"Consolas", "Courier New", "Cascadia Mono"}


def _run_is_code(r) -> bool:
    fonts = r.findall(qn('w:rPr') + '/' + qn('w:rFonts'))
    for f in fonts:
        if f.get(qn('w:ascii')) in _MONO_FONTS:
            return True
    return False


def _clone_run(r, text: str, script: str = ""):
    import copy
    new = copy.deepcopy(r)
    for child in list(new):
        if child.tag != qn('w:rPr'):
            new.remove(child)
    if script:
        rPr = new.find(qn('w:rPr'))
        if rPr is None:
            rPr = new.makeelement(qn('w:rPr'), {})
            new.insert(0, rPr)
        for old in rPr.findall(qn('w:vertAlign')):
            rPr.remove(old)
        val = 'superscript' if script == 'sup' else 'subscript'
        rPr.append(rPr.makeelement(qn('w:vertAlign'), {qn('w:val'): val}))
    t = new.makeelement(qn('w:t'), {qn('xml:space'): 'preserve'})
    t.text = text
    new.append(t)
    return new


def _math_replacement_nodes(r, text: str):
    from app.math_render import script_segments, omml_element
    normalized = _normalize_math(text)
    if "$" not in normalized:
        return None
    nodes = []
    pos = 0
    changed = False
    for m in _MATH_SPAN_DOCX.finditer(normalized):
        prose = normalized[pos:m.start()]
        if prose:
            nodes.append(_clone_run(r, prose))
        latex = m.group(1)
        segments = script_segments(latex)
        if segments:
            for seg_text, script in segments:
                nodes.append(_clone_run(r, seg_text, script))
            changed = True
        else:
            el = omml_element(latex)
            if el is not None:
                nodes.append(el)
                changed = True
            else:
                nodes.append(_clone_run(r, latex))
        pos = m.end()
    rest = normalized[pos:]
    if rest:
        nodes.append(_clone_run(r, rest))
    return nodes if changed else None


def _typeset_math_runs(doc) -> None:
    try:
        runs = [r for r in doc.element.body.iter(qn('w:r'))]
    except Exception:
        return
    for r in runs:
        try:
            if _run_is_code(r):
                continue
            texts = r.findall(qn('w:t'))
            if len(texts) != 1 or not (texts[0].text or ""):
                continue
            nodes = _math_replacement_nodes(r, texts[0].text)
            if not nodes:
                continue
            parent = r.getparent()
            if parent is None:
                continue
            index = list(parent).index(r)
            for offset, node in enumerate(nodes):
                parent.insert(index + offset, node)
            parent.remove(r)
        except Exception:
            continue


def _set_cell_left_bar(cell, color_hex: str, size: str = "18"):
    tcPr = cell._element.get_or_add_tcPr()
    borders = tcPr.makeelement(qn('w:tcBorders'), {})
    left = borders.makeelement(qn('w:left'), {
        qn('w:val'): 'single', qn('w:sz'): size,
        qn('w:space'): '0', qn('w:color'): color_hex,
    })
    borders.append(left)
    tcPr.append(borders)


def _plan_box(doc, fill_hex: str, keep_together: bool = True):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if keep_together:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(trPr.makeelement(qn('w:cantSplit'), {}))
    cell = t.cell(0, 0)
    cell.width = _PLAN_WIDTH
    _set_cell_shading(cell, fill_hex)
    _set_cell_left_bar(cell, _PLAN_BAR_DOCX)
    _clear_cell(cell)
    return cell


def _plan_cell_para(cell, space_after=2):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _plan_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    _plan_run(p, text, size=11, bold=True, color=_PLAN_INK_DOCX)
    _add_bottom_border(p, _PLAN_BAR_DOCX)
    return p


def _plan_feature_heading(doc, text, badge=None):
    t = doc.add_table(rows=1, cols=2 if badge else 1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    left = t.cell(0, 0)
    left.width = Cm(13.5) if badge else _PLAN_WIDTH
    _set_cell_shading(left, _PLAN_BG_DOCX)
    _set_cell_left_bar(left, _PLAN_BAR_DOCX, "24")
    _clear_cell(left)
    _plan_run(_plan_cell_para(left, 1), text, size=12, bold=True, color=_PLAN_INK_DOCX)
    if badge:
        right = t.cell(0, 1)
        right.width = Cm(3.7)
        _set_cell_shading(right, _PLAN_BG_DOCX)
        _clear_cell(right)
        p = _plan_cell_para(right, 1)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _plan_run(p, str(badge), size=10, color=_PLAN_INK_DOCX)


def _plan_bullets(cell_or_doc, items, in_cell=False):
    for item in (items if isinstance(items, list) else [items]):
        if not str(item).strip():
            continue
        p = _plan_cell_para(cell_or_doc, 2) if in_cell else cell_or_doc.add_paragraph()
        if not in_cell:
            p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        _plan_run(p, "•  ")
        _plan_text_math(p, item)


def _plan_example_card(doc, label, problem, solution, solution_label):
    cell = _plan_box(doc, _PLAN_EX_BG_DOCX)
    p = _plan_cell_para(cell, 2)
    _plan_run(p, label + " ", bold=True, color=_PLAN_INK_DOCX)
    _plan_text_math(p, problem)
    if solution:
        p2 = _plan_cell_para(cell, 1)
        p2.paragraph_format.left_indent = Cm(0.45)
        _plan_run(p2, solution_label + ": ", size=10.5, italic=True)
        head, marked = _split_answer(str(solution))
        _plan_text_math(p2, head, size=10.5)
        if marked:
            _plan_text_math(p2, marked, size=10.5, bold=True, mark_text=True)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(0)
    gap.add_run().font.size = Pt(4)


_ANSWER_WORDS_DOCX = ("ҷавоб", "җавоб", "жавоб", "javob", "ответ", "answer",
                      "натиҷа", "натича", "натижа", "natija", "результат", "result")


def _split_answer(text: str):
    low = text.lower()
    cut = max((low.rfind(w) for w in _ANSWER_WORDS_DOCX), default=-1)
    if cut < 0:
        return text, ""
    return text[:cut], text[cut:]


def _plan_labels(L: dict, language: str) -> dict:
    from app.export_builder import _PDF_KONSPEKT_LABELS
    pdf = _PDF_KONSPEKT_LABELS.get(language) or _PDF_KONSPEKT_LABELS["Русский"]
    merged = dict(L)
    for key in ("plan_date", "plan_class", "plan_school", "plan_year",
                "plan_note", "example", "solution", "worked_examples"):
        if pdf.get(key):
            merged[key] = pdf[key]
    return merged


def _plan_visual_docx(doc, block: dict, L: dict):
    btype = block.get("type")
    data = block.get("data") or {}

    if block.get("title"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        _plan_run(p, str(block["title"]), size=10, italic=True, color=_PLAN_INK_DOCX)

    if btype in ("table", "comparison"):
        if btype == "table":
            headers = list(data.get("headers") or [])
            rows = [list(r) for r in (data.get("rows") or [])]
        else:
            criteria = list(data.get("criteria") or [])
            headers = [""] + criteria
            rows = [[it.get("name", "")] + list(it.get("values") or [])
                    for it in (data.get("items") or []) if isinstance(it, dict)]
        if not headers or not rows:
            return
        width = len(headers)
        rows = [(r + [""] * width)[:width] for r in rows]
        t = doc.add_table(rows=1 + len(rows), cols=width)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(headers):
            c = t.cell(0, j)
            _clear_cell(c)
            _set_cell_shading(c, _PLAN_BG_DOCX)
            _plan_text_math(_plan_cell_para(c, 1), h, size=9.5, bold=True,
                            color=_PLAN_INK_DOCX)
        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                _clear_cell(c)
                _plan_text_math(_plan_cell_para(c, 1), val, size=9.5)
        tblPr = t._tbl.tblPr
        borders = tblPr.makeelement(qn('w:tblBorders'), {})
        for edge in ("top", "bottom", "insideH"):
            borders.append(borders.makeelement(qn('w:' + edge), {
                qn('w:val'): 'single', qn('w:sz'): '4',
                qn('w:space'): '0', qn('w:color'): _PLAN_BAR_DOCX,
            }))
        for edge in ("left", "right", "insideV"):
            borders.append(borders.makeelement(qn('w:' + edge), {qn('w:val'): 'none'}))
        tblPr.append(borders)

    elif btype in ("process", "flowchart", "timeline"):
        if btype == "timeline":
            items = data.get("events") or []
            def line(k, ev):
                bits = [b for b in (ev.get("date"), ev.get("label")) if b]
                return "•  ", " — ".join(str(b) for b in bits), ev.get("description")
        else:
            items = data.get("steps") or []
            def line(k, st):
                return f"{k}.  ", str(st.get("title", "")), st.get("description")
        wrote = False
        for k, item in enumerate(items, 1):
            if not isinstance(item, dict):
                continue
            prefix, head, desc = line(k, item)
            if not head and not desc:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            _plan_run(p, prefix)
            _plan_text_math(p, head, bold=True)
            if desc:
                _plan_run(p, " — ")
                _plan_text_math(p, desc)
            wrote = True
        if not wrote:
            return
    else:
        return

    note = block.get("description") or block.get("note")
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        _plan_text_math(p, note, size=10, italic=True)


def _add_plan_body_docx(doc, content: dict, L: dict):
    from app.cover_builder import _academic_year
    from app.export_builder import _display_subject, _bare_grade
    grade = str(content.get("grade") or "")
    subject = str(content.get("subject") or "")
    language = str(content.get("language") or "Русский")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    _plan_run(p, f'{L.get("plan_year", "Соли таҳсили")} {_academic_year()}', size=10.5)

    fill = doc.add_table(rows=1, cols=3)
    fill.autofit = False
    grade_text = f'{L.get("plan_class", "Синф")} {_bare_grade(grade) or grade or "________"}'
    for i, (txt, width) in enumerate((
        (f'{L.get("plan_date", "Сана")} ____________', Cm(6.0)),
        (grade_text, Cm(5.2)),
        (f'{L.get("plan_school", "Мактаб")} ____________', Cm(6.0)),
    )):
        c = fill.cell(0, i)
        c.width = width
        _clear_cell(c)
        _plan_run(_plan_cell_para(c, 0), txt, size=10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    if subject:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        _plan_run(p, " ".join(_display_subject(subject, language).upper()), size=8.5, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    _plan_run(p, str(content.get("title") or "").upper(), size=16, bold=True)

    heavy = []

    formulas = content.get("formulas") or []
    if formulas:
        def _formulas():
            _plan_heading(doc, L.get("formulas", "Формулаҳо"))
            cell = _plan_box(doc, _PLAN_BG_DOCX, keep_together=False)
            for f in formulas:
                latex = f.get("latex") if isinstance(f, dict) else None
                plain = f.get("formula", "") if isinstance(f, dict) else str(f)
                if not latex and plain:
                    latex = _plain_to_latex(plain)
                p = _plan_cell_para(cell, 1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not (latex and _plan_math_run(p, latex)):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.left_indent = Cm(0.5)
                    _plan_run(p, "•  " + str(plain), bold=True)
                    if isinstance(f, dict) and f.get("explanation"):
                        _plan_run(p, " — " + str(f["explanation"]))
                    continue
                if isinstance(f, dict) and f.get("explanation"):
                    e = _plan_cell_para(cell, 5)
                    e.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _plan_run(e, str(f["explanation"]), size=9.5)
        heavy.append(_formulas)

    examples = [e for e in (content.get("worked_examples") or [])
                if isinstance(e, dict) and e.get("problem")]
    if examples:
        def _examples():
            _plan_feature_heading(doc, L.get("worked_examples", "Мисолҳо"), str(len(examples)))
            for i, ex in enumerate(examples, 1):
                _plan_example_card(
                    doc, f'{L.get("example", "Мисол")} {i}.', ex["problem"],
                    ex.get("solution"), L.get("solution", "Ҳал"),
                )
        heavy.append(_examples)

    for block in (content.get("visual_blocks") or []):
        if isinstance(block, dict) and block.get("type"):
            def _vis(b=block):
                _plan_visual_docx(doc, b, L)
            heavy.append(_vis)

    for cb in (content.get("code_blocks") or []):
        if isinstance(cb, dict) and cb.get("code"):
            def _code(c=cb):
                _add_code_card(doc, c["code"], c.get("language", ""),
                               c.get("explanation", ""), _PLAN_INK_DOCX)
            heavy.append(_code)

    for figure in (content.get("figures") or []):
        if isinstance(figure, dict) and figure.get("image"):
            def _fig(f=figure):
                width = 3.6
                try:
                    import os
                    from PIL import Image as _PILImg
                    full = os.path.join(os.path.dirname(__file__), "..", f["image"].lstrip("/"))
                    with _PILImg.open(full) as im:
                        width = min(3.6, 2.9 * im.size[0] / im.size[1])
                except Exception:
                    pass
                _add_illustration(doc, f["image"], f.get("caption", ""), "", width_inches=width)
            heavy.append(_fig)

    lesson_by_anchor: dict = {}
    for image in (content.get("lesson_images") or []):
        if isinstance(image, dict) and image.get("path"):
            lesson_by_anchor.setdefault(image.get("position_after") or "main_content", []).append(image)

    PLAN_ORDER = [
        ("competencies", "fill"), ("warmup", "fill"), ("objectives", "bullets"),
        ("key_concepts", "bullets"), ("key_terms", "bullets"), ("tools", "runin"),
        ("main_content", "runin"), ("important_notes", "notes"),
        ("real_life_examples", "bullets"), ("lesson_program", "runin"),
        ("pair_work", "runin"), ("group_work", "bullets"),
        ("consolidation", "runin"), ("homework", "bullets"),
        ("summary", "runin"), ("assessment", "runin"),
    ]
    PANELLED = {"key_concepts", "key_terms"}

    printable = [k for k, kind in PLAN_ORDER if kind == "fill" or content.get(k)]
    slots = printable[1:]
    drop_after: dict = {}
    if slots and heavy:
        for i, item in enumerate(heavy):
            pos = round((i + 1) * len(slots) / (len(heavy) + 1))
            drop_after.setdefault(slots[max(0, min(pos, len(slots) - 1))], []).append(item)

    for key, kind in PLAN_ORDER:
        value = content.get(key)
        label = L.get(key, key)
        if kind == "fill":
            _plan_heading(doc, f"{label}:")
            if value:
                _plan_bullets(doc, value)
            else:
                for _ in range(2):
                    blank = doc.add_paragraph()
                    blank.paragraph_format.space_before = Pt(0)
                    blank.paragraph_format.space_after = Pt(11)
                    _add_thin_border(blank, "9AA5B1")
        elif value and kind == "notes":
            for note in (value if isinstance(value, list) else [value]):
                if not str(note).strip():
                    continue
                cell = _plan_box(doc, _PLAN_BG_DOCX)
                p = _plan_cell_para(cell, 1)
                _plan_run(p, f'{L.get("plan_note", "Диққат")}. ', bold=True, color=_PLAN_INK_DOCX)
                _plan_text_math(p, note)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif value and kind == "bullets":
            _plan_heading(doc, label)
            if key in PANELLED:
                cell = _plan_box(doc, _PLAN_BG_DOCX, keep_together=False)
                _plan_bullets(cell, value, in_cell=True)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                _plan_bullets(doc, value)
        elif value and kind == "runin":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _plan_run(p, f"{label}. ", bold=True, italic=True, color=_PLAN_INK_DOCX)
            _plan_text_math(p, _plan_flatten_docx(value))
        for image in lesson_by_anchor.pop(key, []):
            _add_lesson_image_block(doc, image, L.get("illustration", "Тасвир"), _PLAN_INK_DOCX)
        for item in drop_after.pop(key, []):
            item()

    for images in lesson_by_anchor.values():
        for image in images:
            _add_lesson_image_block(doc, image, L.get("illustration", "Тасвир"), _PLAN_INK_DOCX)
    for leftover in drop_after.values():
        for item in leftover:
            item()


def _plan_flatten_docx(value):
    if isinstance(value, list):
        parts = [str(v).strip().rstrip(".") for v in value if v]
        return "; ".join(parts) + ("." if parts else "")
    return str(value or "")


def build_konspekt_docx(content: dict, language: str = 'Русский') -> io.BytesIO:
    doc = Document()
    L = _DOCX_LABELS.get(language, _DOCX_LABELS['Русский'])

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    title = content.get('title', 'Конспект урока')
    subtitle = content.get('subtitle', '')


    if get_template(content.get('template') or 'zamonaviy').plan_layout:
        for section in doc.sections:
            section.top_margin = Cm(1.4)
            section.bottom_margin = Cm(1.4)
            section.left_margin = Cm(1.9)
            section.right_margin = Cm(1.9)
        try:
            from app.cover_builder import build_cover_image
            png = build_cover_image(
                content.get('subject', ''), title, content.get('grade', ''),
                language, template_id='nakscha',
            )
            import tempfile, os as _os
            fd, tmp = tempfile.mkstemp(suffix='.png')
            with _os.fdopen(fd, 'wb') as fh:
                fh.write(png)
            from docx.enum.text import WD_BREAK
            doc.add_picture(tmp, height=Cm(23.5))
            pic = doc.paragraphs[-1]
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_after = Pt(0)
            pic.add_run().add_break(WD_BREAK.PAGE)
            _os.unlink(tmp)
        except Exception:
            pass
        _add_plan_body_docx(doc, content, _plan_labels(L, language))
        buf = io.BytesIO()
        _typeset_math_runs(doc)
        doc.save(buf)
        buf.seek(0)
        return buf

    for section in doc.sections:
        section.top_margin = Cm(1.1)
        section.bottom_margin = Cm(1.1)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)


    DARK = RGBColor(0x1F, 0x29, 0x37)
    SECONDARY = RGBColor(0x6B, 0x72, 0x80)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.size = Pt(22)
    run.font.color.rgb = DARK
    run.bold = True
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(subtitle)
        run.font.size = Pt(12)
        run.font.color.rgb = SECONDARY
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    _add_konspekt_body(doc, content, L)

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf


def build_lecture_docx(content: dict, language: str = 'Русский') -> io.BytesIO:
    from app.lecture_builder import _words, _split_definition, _split_main_content, _queue_callouts
    from app.subject_theme import get_subject_accent_rgb

    w = _words(language)
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.12

    for section in doc.sections:
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    ar, ag, ab = get_subject_accent_rgb(content.get('subject'))
    ACCENT = RGBColor(ar, ag, ab)
    accent_hex = '%02X%02X%02X' % (ar, ag, ab)
    INK = RGBColor(0x15, 0x1A, 0x22)
    MUTED = RGBColor(0x4A, 0x54, 0x62)

    def para(text, size=10.5, bold=False, italic=False, color=INK, after=4,
             before=0, align=None, indent=None):
        pr = doc.add_paragraph()
        pr.paragraph_format.space_after = Pt(after)
        pr.paragraph_format.space_before = Pt(before)
        if align is not None:
            pr.alignment = align
        if indent is not None:
            pr.paragraph_format.left_indent = Cm(indent)
        run = pr.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color
        run.font.name = 'Calibri'
        return pr

    para('  '.join(w['tag']), size=8, bold=True, color=ACCENT, after=1,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(str(content.get('title') or ''), size=15, bold=True, after=1,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    if content.get('subtitle'):
        para(str(content['subtitle']), size=10, color=MUTED, after=6,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    from app.export_builder import _format_grade
    grade = str(content.get('grade') or '')
    if grade:
        grade = _format_grade(grade, w['grade_suffix'])
    facts = [(w['subject'], str(content.get('subject') or '—')),
             (w['grade_suffix'].capitalize(), grade or '—'),
             (w['duration'], w['duration_value']),
             (w['date'], '________')]
    table = doc.add_table(rows=2, cols=len(facts))
    table.autofit = True
    for i, (label, value) in enumerate(facts):
        lc = table.cell(0, i).paragraphs[0]
        lr = lc.add_run(label)
        lr.font.size = Pt(7.5)
        lr.font.color.rgb = MUTED
        lr.font.name = 'Calibri'
        vc = table.cell(1, i).paragraphs[0]
        vr = vc.add_run(value)
        vr.font.size = Pt(9.5)
        vr.bold = True
        vr.font.color.rgb = INK
        vr.font.name = 'Calibri'
    _add_bottom_border(doc.add_paragraph(), accent_hex)

    section_no = [0]

    def heading(label):
        section_no[0] += 1
        pr = doc.add_paragraph()
        pr.paragraph_format.space_before = Pt(11)
        pr.paragraph_format.space_after = Pt(2)
        run = pr.add_run(f'{section_no[0]}.  {label.upper()}')
        run.font.size = Pt(10.5)
        run.bold = True
        run.font.color.rgb = INK
        run.font.name = 'Calibri'
        _add_bottom_border(pr, accent_hex)

    def note(label, text, colour=ACCENT):
        pr = doc.add_paragraph()
        pr.paragraph_format.left_indent = Cm(0.5)
        pr.paragraph_format.space_before = Pt(3)
        pr.paragraph_format.space_after = Pt(3)
        _set_paragraph_shading(pr, 'F6F8FA')
        run = pr.add_run(f'{label}.  ')
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = colour
        run.font.name = 'Calibri'
        run = pr.add_run(text)
        run.font.size = Pt(9.5)
        run.font.color.rgb = INK
        run.font.name = 'Calibri'

    objective = str(content.get('objective') or '').strip()
    if objective:
        heading(w['objective'])
        para(objective, size=10.5, indent=0.4, after=2)

    plan = [str(x).strip() for x in (content.get('lecture_plan') or []) if str(x).strip()]
    if plan:
        heading(w['plan'])
        for i, item in enumerate(plan, 1):
            pr = doc.add_paragraph()
            pr.paragraph_format.left_indent = Cm(0.9)
            pr.paragraph_format.first_line_indent = Cm(-0.55)
            pr.paragraph_format.space_after = Pt(2)
            run = pr.add_run(f'{i:02d}  ')
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = ACCENT
            run.font.name = 'Calibri'
            run = pr.add_run(item)
            run.font.size = Pt(10.5)
            run.font.color.rgb = INK
            run.font.name = 'Calibri'

    def definitions(items):
        for item in items:
            term, body = _split_definition(item)
            pr = doc.add_paragraph()
            pr.paragraph_format.left_indent = Cm(0.9)
            pr.paragraph_format.first_line_indent = Cm(-0.5)
            pr.paragraph_format.space_after = Pt(3)
            if term:
                run = pr.add_run(f'{term} — ')
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = INK
                run.font.name = 'Calibri'
            run = pr.add_run(body)
            run.font.size = Pt(10.5)
            run.font.color.rgb = INK
            run.font.name = 'Calibri'

    concepts = [c for c in (content.get('key_concepts') or []) if str(c).strip()]
    if concepts:
        heading(w['concepts'])
        definitions(concepts)
    terms = [t for t in (content.get('key_terms') or []) if str(t).strip()]
    if terms:
        heading(w['terms'])
        definitions(terms)

    parts = _split_main_content(content.get('main_content'))
    images = list(content.get('lesson_images') or [])
    visuals = list(content.get('visual_blocks') or [])
    if parts:
        heading(w['body'])
        callouts = _queue_callouts(content, w)
        for i, (head, paragraphs) in enumerate(parts):
            if head:
                para(head, size=11, bold=True, color=ACCENT, before=7, after=2)
            for text in paragraphs:
                pr = para(text, size=10.5, after=4)
                pr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if callouts and i < len(parts) - 1:
                kind, label, text = callouts.pop(0)
                note(label, text,
                     RGBColor(0xB4, 0x23, 0x2A) if kind == 'important' else ACCENT)
            if images and i in (0, len(parts) // 2):
                try:
                    _add_lesson_image_block(doc, images.pop(0), w['figure'], ACCENT)
                except Exception:
                    pass
            if visuals and i == max(0, len(parts) - 2):
                try:
                    _add_visual_block(doc, visuals.pop(0), ACCENT,
                                      _DOCX_LABELS.get(language, _DOCX_LABELS['Русский']))
                except Exception:
                    pass
        for kind, label, text in callouts:
            note(label, text,
                 RGBColor(0xB4, 0x23, 0x2A) if kind == 'important' else ACCENT)
    for image in images:
        try:
            _add_lesson_image_block(doc, image, w['figure'], ACCENT)
        except Exception:
            pass
    for block in visuals:
        try:
            _add_visual_block(doc, block, ACCENT,
                              _DOCX_LABELS.get(language, _DOCX_LABELS['Русский']))
        except Exception:
            pass

    formulas = [f for f in (content.get('formulas') or []) if isinstance(f, dict)]
    if formulas:
        heading(w['formulas'])
        for item in formulas:
            formula = str(item.get('formula') or '').strip()
            if not formula:
                continue
            pr = para(formula, size=12, bold=True, after=1,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_paragraph_shading(pr, 'F6F8FA')
            if item.get('explanation'):
                para(str(item['explanation']), size=9.5, color=MUTED, after=6,
                     align=WD_ALIGN_PARAGRAPH.CENTER)

    checks = [q for q in (content.get('quick_check') or []) if q]
    if checks:
        heading(w['selfcheck'])
        for i, item in enumerate(checks, 1):
            if isinstance(item, dict):
                question = str(item.get('question') or '').strip()
                answer = str(item.get('answer') or '').strip()
            else:
                question, answer = str(item).strip(), ''
            if not question:
                continue
            pr = doc.add_paragraph()
            pr.paragraph_format.left_indent = Cm(0.9)
            pr.paragraph_format.first_line_indent = Cm(-0.5)
            pr.paragraph_format.space_after = Pt(2)
            run = pr.add_run(f'{i}.  ')
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = INK
            run.font.name = 'Calibri'
            run = pr.add_run(question)
            run.font.size = Pt(10.5)
            run.font.color.rgb = INK
            run.font.name = 'Calibri'
            if answer:
                para(f"{w['answer']}: {answer}", size=9.5, color=MUTED,
                     indent=0.9, after=4)

    summary = str(content.get('summary') or '').strip()
    if summary:
        heading(w['summary'])
        pr = para(summary, size=10.5, after=4, indent=0.4)
        pr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    refs = [str(r).strip() for r in (content.get('references') or []) if str(r).strip()]
    if refs:
        heading(w['references'])
        for i, ref in enumerate(refs, 1):
            para(f'{i}.  {ref}', size=9.5, color=MUTED, indent=0.9, after=2)

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf


_ROADMAP_LABELS = {
    'Русский': {'day': 'День', 'date': 'Дата', 'topic': 'Тема', 'type': 'Тип', 'lesson': 'Урок', 'exam': 'Экзамен', 'roadmap': 'Дорожная карта курса', 'overview': 'Обзор', 'course': 'Курс'},
    'Таджикский': {'day': 'Рӯз', 'date': 'Сана', 'topic': 'Мавзӯъ', 'type': 'Навъ', 'lesson': 'Дарс', 'exam': 'Имтиҳон', 'roadmap': 'Харитаи курс', 'overview': 'Дида баромадан', 'course': 'Курс'},
    'English': {'day': 'Day', 'date': 'Date', 'topic': 'Topic', 'type': 'Type', 'lesson': 'Lesson', 'exam': 'Exam', 'roadmap': 'Course Roadmap', 'overview': 'Overview', 'course': 'Course'},
}
_ROADMAP_LABELS['Английский'] = _ROADMAP_LABELS['English']


def _mono_heading(doc, text, size=14, bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    run.bold = bold
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    return p


def _mono_para(doc, text, size=11, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    run.italic = italic
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    return p


def _mono_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    return p


def _mono_divider(doc):
    _add_bottom_border(doc.add_paragraph(), 'CCCCCC')


def _add_lesson_content(doc, content: dict, L: dict):
    duration = content.get('duration', '')
    if duration:
        _mono_para(doc, f'{L["duration"]}: {duration}', italic=True)

    for key, is_list in [
        ('competencies', True), ('objectives', True), ('key_concepts', True),
        ('key_terms', True), ('lesson_program', False), ('tools', True),
        ('warmup', False), ('main_content', False), ('real_life_examples', True),
        ('visual_aid', False), ('pair_work', False), ('consolidation', False),
        ('fun_facts', True), ('common_mistakes', True), ('summary', False),
        ('assessment', False),
    ]:
        value = content.get(key)
        if not value:
            continue
        label = L.get(key, key)
        _mono_heading(doc, label, size=12)
        if is_list and isinstance(value, list):
            for item in value:
                _mono_bullet(doc, str(item))
        else:
            _mono_para(doc, str(value))

    group_work = content.get('group_work')
    if group_work:
        _mono_heading(doc, L.get('group_work', 'Group Work'), size=12)
        for i, task in enumerate(group_work):
            _mono_para(doc, f"{L.get('group', 'Group')} {i + 1}. {_strip_group_label(task)}")

    homework = content.get('homework')
    if homework:
        _mono_heading(doc, L.get('homework', 'Homework'), size=12)
        if isinstance(homework, list):
            for i, h in enumerate(homework):
                _mono_para(doc, f'{i + 1}. {h}')
        else:
            _mono_para(doc, str(homework))


def _add_exam_content(doc, content: dict, R: dict):
    description = content.get('description', '')
    if description:
        _mono_para(doc, description, italic=True)
    for i, q in enumerate(content.get('questions', [])):
        q_type = q.get('type', 'multiple_choice')
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f'{i + 1}. {q.get("question", "")}')
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        run.bold = True
        run.font.name = 'Calibri'
        img = q.get('image') if isinstance(q.get('image'), dict) else None
        if img and img.get('path'):
            _add_illustration(doc, img.get('path', ''), '', img.get('credit', ''), width_inches=2.6)
        if q_type == 'open_ended':
            if q.get('model_answer'):
                _mono_para(doc, f'{R.get("answer", "Answer")}: {q["model_answer"]}')
        else:
            options = q.get('options', [])
            correct_indices = set(q.get('correct_indices') or [])
            correct = q.get('correct_index', 0)
            for j, opt in enumerate(options):
                is_correct = (j in correct_indices) if q_type == 'multiple_select' else (j == correct)
                mark = '✓ ' if is_correct else '   '
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(f'{mark}{chr(65 + j)}) {opt}')
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                run.bold = is_correct
                run.font.name = 'Calibri'
        if q.get('explanation'):
            _mono_para(doc, f'{R.get("explanation", "Explanation")}: {q["explanation"]}', size=10, italic=True)


def build_konspekt_docx_mono(content: dict, language: str = 'Русский') -> io.BytesIO:
    doc = Document()
    L = _DOCX_LABELS.get(language, _DOCX_LABELS['Русский'])

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(2)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = content.get('title', 'Конспект урока')
    subtitle = content.get('subtitle', '')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run.bold = True
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(subtitle)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        run.font.name = 'Calibri'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    _mono_divider(doc)
    _add_lesson_content(doc, content, L)

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf


def build_test_docx_mono(content: dict, language: str = 'Русский') -> io.BytesIO:
    doc = Document()
    R = _ROADMAP_LABELS.get(language, _ROADMAP_LABELS['Русский'])

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(2)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = content.get('title', 'Тест')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run.bold = True
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    _mono_divider(doc)
    _add_exam_content(doc, content, R)

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf


def build_test_docx(content: dict) -> io.BytesIO:
    from app.export_builder import _test_ui

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    DARK = RGBColor(0x1F, 0x29, 0x37)
    GREY = RGBColor(0x6B, 0x72, 0x80)
    language = str(content.get('language') or 'Русский')
    ui = _test_ui(language)
    try:
        from app.export_builder import get_subject_accent_rgb
        ar, ag, ab = get_subject_accent_rgb(content.get('subject', ''))
        ACCENT = RGBColor(ar, ag, ab)
    except Exception:
        ACCENT = RGBColor(0x10, 0xB9, 0x81)

    title = content.get('title', 'Тест')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(22)
    run.font.color.rgb = DARK
    run.bold = True
    run.font.name = 'Calibri'

    description = content.get('description', '')
    if description:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(description)
        run.font.size = Pt(12)
        run.font.color.rgb = GREY
        run.font.name = 'Calibri'

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(f'{ui["name"]} ________________________________     '
                    f'{ui["klass"]} _______     {ui["date"]} __________     '
                    f'{ui["mark"]} _______')
    run.font.size = Pt(10.5)
    run.font.color.rgb = GREY
    run.font.name = 'Calibri'

    _add_bottom_border(doc.add_paragraph(), '374151')

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(ui['hint'])
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    run.italic = True
    run.font.name = 'Calibri'

    questions = content.get('questions', [])
    _LETTERS = [chr(65 + i) for i in range(8)]

    for i, q in enumerate(questions):
        q_type = q.get('type', 'multiple_choice')
        options = q.get('options', [])

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f'{i + 1}.  ')
        run.font.size = Pt(12)
        run.font.color.rgb = ACCENT
        run.bold = True
        run.font.name = 'Calibri'
        run = p.add_run(q.get('question', ''))
        run.font.size = Pt(12)
        run.font.color.rgb = DARK
        run.bold = True
        run.font.name = 'Calibri'
        if q_type == 'multiple_select':
            run = p.add_run(f'  ({ui["select_hint"]})')
            run.font.size = Pt(9)
            run.font.color.rgb = GREY
            run.italic = True
            run.font.name = 'Calibri'

        img = q.get('image') if isinstance(q.get('image'), dict) else None
        if img and img.get('path'):
            _add_illustration(doc, img.get('path', ''), '', img.get('credit', ''), width_inches=2.6)

        if q_type == 'open_ended':
            for _ in range(3):
                line = doc.add_paragraph()
                line.paragraph_format.left_indent = Cm(1)
                line.paragraph_format.space_after = Pt(10)
                _add_bottom_border(line, 'D1D5DB')
        else:
            for j, opt in enumerate(options):
                letter = _LETTERS[j] if j < len(_LETTERS) else str(j)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_after = Pt(2)
                box = '[   ]  ' if q_type == 'multiple_select' else ''
                run = p.add_run(f'{box}{letter})  ')
                run.font.size = Pt(11)
                run.font.color.rgb = GREY
                run.bold = True
                run.font.name = 'Calibri'
                run = p.add_run(str(opt))
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
                run.font.name = 'Calibri'

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{ui["total"]}: {len(questions)}     ·     '
                    f'{ui["max_score"]}: {len(questions)}')
    run.font.size = Pt(9.5)
    run.font.color.rgb = GREY
    run.font.name = 'Calibri'

    if questions:
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(ui['key_title'])
        run.font.size = Pt(16)
        run.font.color.rgb = DARK
        run.bold = True
        run.font.name = 'Calibri'
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(ui['key_hint'])
        run.font.size = Pt(9)
        run.font.color.rgb = GREY
        run.italic = True
        run.font.name = 'Calibri'

        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        table.autofit = False
        col_widths = (Cm(1.2), Cm(3.8), Cm(11.0))
        accent_hex = '%02X%02X%02X' % (ACCENT[0], ACCENT[1], ACCENT[2])
        for cell, label, width in zip(table.rows[0].cells,
                                      (ui['col_no'], ui['col_answer'], ui['col_why']),
                                      col_widths):
            cell.width = width
            cell.text = ''
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Calibri'
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _set_cell_shading(cell, accent_hex)

        for i, q in enumerate(questions):
            q_type = q.get('type', 'multiple_choice')
            options = q.get('options') or []
            if q_type == 'open_ended':
                answer = ui['open']
            elif q_type == 'multiple_select':
                idx = sorted(int(x) for x in (q.get('correct_indices') or [])
                             if isinstance(x, (int, float)) and 0 <= int(x) < len(options))
                answer = ', '.join(_LETTERS[x] for x in idx) if idx else '—'
            else:
                ci = q.get('correct_index')
                answer = ('—' if not isinstance(ci, int) or not (0 <= ci < len(options))
                          else (f'{_LETTERS[ci]} · {options[ci]}' if q_type == 'true_false'
                                else _LETTERS[ci]))

            note = str(q.get('explanation') or '').strip()
            model = str(q.get('model_answer') or '').strip()

            row = table.add_row().cells
            for cell, width in zip(row, col_widths):
                cell.width = width
            r = row[0].paragraphs[0].add_run(str(i + 1))
            r.bold = True
            r.font.size = Pt(10)
            r.font.name = 'Calibri'
            r = row[1].paragraphs[0].add_run(answer)
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = ACCENT
            r.font.name = 'Calibri'
            cell = row[2]
            if model:
                r = cell.paragraphs[0].add_run(f'{ui["model"]}: ')
                r.bold = True
                r.font.size = Pt(9.5)
                r.font.name = 'Calibri'
                r = cell.paragraphs[0].add_run(model)
                r.font.size = Pt(9.5)
                r.font.name = 'Calibri'
                if note:
                    cell.add_paragraph()
            if note:
                target = cell.paragraphs[-1]
                r = target.add_run(note)
                r.font.size = Pt(9.5)
                r.font.color.rgb = GREY
                r.font.name = 'Calibri'

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf


_PRACTICAL_DOCX_UI = {
    'Русский': {'individual': 'Индивидуальные задания', 'group': 'Групповые задания',
               'difficulty': {'easy': 'лёгкое', 'medium': 'среднее', 'hard': 'сложное'},
               'roles': 'Роли в группе', 'outcome': 'Результат', 'group_size': 'Состав группы'},
    'Таджикский': {'individual': 'Супоришҳои инфиродӣ', 'group': 'Супоришҳои гурӯҳӣ',
                  'difficulty': {'easy': 'сабук', 'medium': 'миёна', 'hard': 'душвор'},
                  'roles': 'Нақшҳо дар гурӯҳ', 'outcome': 'Натиҷа', 'group_size': 'Таркиби гурӯҳ'},
    'English': {'individual': 'Individual tasks', 'group': 'Group tasks',
               'difficulty': {'easy': 'easy', 'medium': 'medium', 'hard': 'hard'},
               'roles': 'Roles in the group', 'outcome': 'Expected outcome', 'group_size': 'Group size'},
}


def build_practical_docx(content: dict) -> io.BytesIO:
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    DARK = RGBColor(0x1F, 0x29, 0x37)
    GREY = RGBColor(0x6B, 0x72, 0x80)
    language = str(content.get('language') or 'Русский')
    ui = _PRACTICAL_DOCX_UI.get(language, _PRACTICAL_DOCX_UI['Русский'])
    try:
        from app.export_builder import get_subject_accent_rgb
        ar, ag, ab = get_subject_accent_rgb(content.get('subject', ''))
        ACCENT = RGBColor(ar, ag, ab)
    except Exception:
        ACCENT = RGBColor(0x10, 0xB9, 0x81)
    DIFF_COLORS = {'easy': RGBColor(0x16, 0xA3, 0x4A), 'medium': RGBColor(0xD9, 0x77, 0x06),
                   'hard': RGBColor(0xDC, 0x26, 0x26)}

    title = content.get('title', 'Амалӣ супоришҳо')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(22)
    run.font.color.rgb = DARK
    run.bold = True
    run.font.name = 'Calibri'

    description = content.get('description', '')
    if description:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(description)
        run.font.size = Pt(12)
        run.font.color.rgb = GREY
        run.font.name = 'Calibri'

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run('_________________________________     _______     __________')
    run.font.size = Pt(10.5)
    run.font.color.rgb = GREY
    run.font.name = 'Calibri'
    _add_bottom_border(doc.add_paragraph(), '374151')

    def add_task(i, task, section_ui_label=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f'{i}.  ')
        run.font.size = Pt(12)
        run.font.color.rgb = ACCENT
        run.bold = True
        run.font.name = 'Calibri'
        run = p.add_run(task.get('title', ''))
        run.font.size = Pt(12)
        run.font.color.rgb = DARK
        run.bold = True
        run.font.name = 'Calibri'
        diff = str(task.get('difficulty') or '').lower()
        if diff in ui['difficulty']:
            run = p.add_run(f"   [{ui['difficulty'][diff].upper()}]")
            run.font.size = Pt(9)
            run.font.color.rgb = DIFF_COLORS.get(diff, GREY)
            run.bold = True
            run.font.name = 'Calibri'
        if task.get('group_size'):
            gp = doc.add_paragraph()
            gp.paragraph_format.left_indent = Cm(0.8)
            r = gp.add_run(f"{ui['group_size']}: {task['group_size']}")
            r.font.size = Pt(9.5)
            r.font.color.rgb = GREY
            r.italic = True
            r.font.name = 'Calibri'
        ip = doc.add_paragraph()
        ip.paragraph_format.left_indent = Cm(0.8)
        ip.paragraph_format.space_after = Pt(3)
        r = ip.add_run(task.get('instructions', ''))
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        r.font.name = 'Calibri'
        roles = task.get('roles') or []
        if roles:
            rp = doc.add_paragraph()
            rp.paragraph_format.left_indent = Cm(0.8)
            rr = rp.add_run(ui['roles'])
            rr.font.size = Pt(9)
            rr.font.color.rgb = GREY
            rr.bold = True
            rr.font.name = 'Calibri'
            for role in roles:
                role_p = doc.add_paragraph()
                role_p.paragraph_format.left_indent = Cm(1.2)
                role_p.paragraph_format.space_after = Pt(1)
                rrun = role_p.add_run(f'• {role}')
                rrun.font.size = Pt(10)
                rrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                rrun.font.name = 'Calibri'
        if task.get('expected_outcome'):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm(0.8)
            orun = op.add_run(f"{ui['outcome']}: {task['expected_outcome']}")
            orun.font.size = Pt(9.5)
            orun.italic = True
            orun.font.color.rgb = GREY
            orun.font.name = 'Calibri'

    individual = content.get('individual_tasks') or []
    if individual:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(8)
        hr = hp.add_run(ui['individual'])
        hr.font.size = Pt(14)
        hr.bold = True
        hr.font.color.rgb = ACCENT
        hr.font.name = 'Calibri'
        for i, task in enumerate(individual):
            if isinstance(task, dict):
                add_task(i + 1, task)

    group = content.get('group_tasks') or []
    if group:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(16)
        hr = hp.add_run(ui['group'])
        hr.font.size = Pt(14)
        hr.bold = True
        hr.font.color.rgb = ACCENT
        hr.font.name = 'Calibri'
        for i, task in enumerate(group):
            if isinstance(task, dict):
                add_task(i + 1, task)

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf


def build_presentation_docx(content: dict) -> io.BytesIO:
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    ACCENT = RGBColor(0x37, 0x41, 0x51)
    DARK = RGBColor(0x1F, 0x29, 0x37)
    GREY = RGBColor(0x6B, 0x72, 0x80)

    title = content.get('title', 'Презентация')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(22)
    run.font.color.rgb = DARK
    run.bold = True
    run.font.name = 'Calibri'

    _add_bottom_border(doc.add_paragraph(), '374151')

    slides = content.get('slides', [])
    for i, slide in enumerate(slides):
        slide_title = slide.get('title', '')
        bullets = slide.get('bullet_points', [])
        notes = slide.get('speaker_notes', '')

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f'  {i + 1}  ')
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.bold = True
        run.font.name = 'Calibri'
        _set_paragraph_shading(p, '374151')

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(slide_title)
        run.font.size = Pt(14)
        run.font.color.rgb = DARK
        run.bold = True
        run.font.name = 'Calibri'

        for b in bullets:
            _add_bullet(doc, b, ACCENT)

        if notes:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            _set_paragraph_shading(p, 'F3F4F6')
            run = p.add_run('Заметки:  ')
            run.font.size = Pt(10)
            run.font.color.rgb = GREY
            run.bold = True
            run.font.name = 'Calibri'
            run = p.add_run(notes)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
            run.italic = True
            run.font.name = 'Calibri'

    buf = io.BytesIO()
    _typeset_math_runs(doc)
    doc.save(buf)
    buf.seek(0)
    return buf

