
import io

from app.ai_service import fix_legacy_tajik_glyphs
from app.logger import get_logger

logger = get_logger(__name__)

MAX_SOURCE_CHARS = 60000


def extract_text_from_docx(raw: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(raw))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text.strip())
    return fix_legacy_tajik_glyphs("\n".join(lines))


def extract_text_from_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception as e:
            logger.warning(f"PDF page extraction failed, skipping page: {e}")
            text = ""
        if text.strip():
            pages.append(text.strip())
    return fix_legacy_tajik_glyphs("\n\n".join(pages))


def extract_text_from_txt(raw: bytes) -> str:
    for encoding in ("utf-8", "cp1251", "latin-1"):
        try:
            return fix_legacy_tajik_glyphs(raw.decode(encoding))
        except UnicodeDecodeError:
            continue
    return fix_legacy_tajik_glyphs(raw.decode("utf-8", errors="replace"))


_EXTRACTORS = {
    ".pdf": extract_text_from_pdf,
    ".docx": extract_text_from_docx,
    ".txt": extract_text_from_txt,
}


def extract_source_text(filename: str, raw: bytes) -> tuple[str, bool]:
    name = (filename or "").lower()
    ext = next((e for e in _EXTRACTORS if name.endswith(e)), None)
    if ext is None:
        raise ValueError(f"Unsupported file type: {filename}")

    try:
        text = _EXTRACTORS[ext](raw)
    except ValueError:
        raise
    except Exception as e:
        logger.error(f"Document extraction failed for {filename}: {e}", exc_info=True)
        raise ValueError("parse_error") from e

    text = text.strip()
    if not text:
        raise ValueError("empty_document")

    truncated = len(text) > MAX_SOURCE_CHARS
    if truncated:
        text = text[:MAX_SOURCE_CHARS]
    return text, truncated
