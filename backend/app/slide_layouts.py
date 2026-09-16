
from __future__ import annotations

import math
import re
from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app import pptx_shapes, slide_characters


TOP_IN = 1.55
BOTTOM_IN = 6.95
LEFT_IN = 0.9
RIGHT_IN = 12.4
WIDTH_IN = RIGHT_IN - LEFT_IN
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5


@dataclass
class Ctx:
    slide: object
    accent: RGBColor
    accent_dark: RGBColor
    accent_soft: RGBColor
    support: RGBColor
    ink: RGBColor
    muted: RGBColor
    bg: tuple
    title_font: str
    body_font: str
    template_id: str
    grade_tier: str
    language: str
    index: int
    total: int
    image_path: str | None = None
    image_credit: str = ""
    character_prop: str = "pointer"
    prop_color: RGBColor | None = None
    topic: str = ""




def _text(ctx, x, y, w, h, s, size, bold=False, color=None, align=PP_ALIGN.LEFT,
          font=None):
    from app.export_builder import _add_text
    return _add_text(ctx.slide, Inches(x), Inches(y), Inches(w), Inches(h),
                     s, size, bold, color if color is not None else ctx.ink,
                     align, font_name=font or ctx.body_font)


def _shape(ctx, shape_type, x, y, w, h, fill=None, line=None, line_pt=1.0,
           rotation=None, adjust=None):
    shp = ctx.slide.shapes.add_shape(shape_type, Inches(x), Inches(y),
                                     Inches(max(w, 0.01)), Inches(max(h, 0.01)))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.fill.solid()
        shp.line.fill.fore_color.rgb = line
        shp.line.width = Pt(line_pt)
    if rotation is not None:
        shp.rotation = rotation
    if adjust is not None:
        try:
            shp.adjustments[0] = adjust
        except Exception:
            pass
    return pptx_shapes.flatten(shp)


def _lines(s, width_in, pt):
    from app.export_builder import _estimate_pptx_lines
    return _estimate_pptx_lines(s, width_in, pt)


def _fit(s, width_in, sizes, max_h_in):
    for pt in sizes:
        h = _lines(s, width_in, pt) * pt * 1.34 / 72.0
        if h <= max_h_in:
            return pt, h
    pt = sizes[-1]
    return pt, _lines(s, width_in, pt) * pt * 1.34 / 72.0


def _character(ctx, x, y, h, pose):
    if ctx.grade_tier == "senior" and ctx.index % 4 != 1:
        return
    if ctx.grade_tier == "middle" and ctx.index % 2 == 1:
        return
    slide_characters.draw(ctx.slide, x, y, h, ctx.accent,
                          ctx.prop_color or ctx.support,
                          pose=pose, prop=ctx.character_prop)



_CALLOUT_LABEL = {
    "important": {"Русский": "ВАЖНО", "Таджикский": "МУҲИМ", "English": "IMPORTANT",
                  "Английский": "IMPORTANT"},
    "remember": {"Русский": "ЗАПОМНИ", "Таджикский": "ДАР ХОТИР ДОР",
                 "English": "REMEMBER", "Английский": "REMEMBER"},
    "fact": {"Русский": "ИНТЕРЕСНЫЙ ФАКТ", "Таджикский": "ДАЛЕЛИ ҶОЛИБ",
             "English": "DID YOU KNOW", "Английский": "DID YOU KNOW"},
    "why": {"Русский": "ПОЧЕМУ?", "Таджикский": "ЧАРО?", "English": "WHY?",
            "Английский": "WHY?"},
    "example": {"Русский": "ПРИМЕР", "Таджикский": "МИСОЛ", "English": "EXAMPLE",
                "Английский": "EXAMPLE"},
}


def draw_callout(ctx, kind, text, x, y, w) -> float:
    kind = str(kind or "important").lower()
    label = _CALLOUT_LABEL.get(kind, _CALLOUT_LABEL["important"]).get(
        ctx.language, _CALLOUT_LABEL.get(kind, _CALLOUT_LABEL["important"])["Русский"])
    pt, text_h = _fit(text, w - 0.7, (14, 13, 12, 11), 1.5)
    h = max(0.78, text_h + 0.62)

    if kind == "important":
        _shape(ctx, MSO_SHAPE.RECTANGLE, x, y, w, h, ctx.accent_soft)
        _shape(ctx, MSO_SHAPE.RECTANGLE, x, y, Inches(0.06).inches, h, ctx.accent)
    elif kind == "remember":
        _shape(ctx, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, None,
               line=ctx.accent, line_pt=1.5, adjust=0.12)
    elif kind == "fact":
        _shape(ctx, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, ctx.accent_soft, adjust=0.2)
        _shape(ctx, MSO_SHAPE.OVAL, x + w - 0.42, y - 0.14, 0.32, 0.32, ctx.accent)
    elif kind == "why":
        _shape(ctx, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, y, w, h, None,
               line=ctx.support, line_pt=1.5)
        _shape(ctx, MSO_SHAPE.RECTANGLE, x, y + h - 0.035, w, 0.035, ctx.accent)
    else:
        _shape(ctx, MSO_SHAPE.SNIP_2_DIAG_RECTANGLE, x, y, w, h, ctx.accent_soft)

    _text(ctx, x + 0.26, y + 0.13, w - 0.5, 0.26, label, 10, True, ctx.accent)
    _text(ctx, x + 0.26, y + 0.4, w - 0.5, text_h + 0.2, text, pt, False, ctx.ink)
    return h


def _callout_of(sd):
    c = sd.get("callout")
    if not isinstance(c, dict):
        return None
    text = _clean(c.get("text"))
    if not text:
        return None
    kind = str(c.get("kind") or "important").strip().lower()
    return (kind if kind in _CALLOUT_LABEL else "important"), text


def _callout(ctx, sd, x, y, w) -> float:
    found = _callout_of(sd)
    if not found:
        return 0.0
    if y + 0.85 > BOTTOM_IN:
        return 0.0
    return draw_callout(ctx, found[0], found[1], x, y, w)




def _hero(ctx, sd) -> bool:
    body = _clean(sd.get("body"))
    bullets = _strings(sd.get("bullet_points"))
    lead = bullets[0] if bullets else body
    if not lead:
        return False
    rest = body if (bullets and body) else ""

    text_w = 8.2
    pt, h = _fit(lead, text_w, (54, 46, 40, 34, 30), 3.0)
    top = TOP_IN + max(0.2, (BOTTOM_IN - TOP_IN - h - (1.0 if rest else 0)) / 2)

    _shape(ctx, MSO_SHAPE.RECTANGLE, LEFT_IN, top - 0.28, 1.3, 0.05, ctx.accent)
    _text(ctx, LEFT_IN, top, text_w, h + 0.3, lead, pt, True, ctx.ink,
          font=ctx.title_font)
    y = top + h + 0.3
    if rest:
        rpt, rh = _fit(rest, text_w, (17, 16, 15, 14), 1.5)
        _text(ctx, LEFT_IN, y, text_w, rh + 0.25, rest, rpt, False, ctx.muted)
        y += rh + 0.2

    _callout(ctx, sd, LEFT_IN, y + 0.25, 7.0)
    _character(ctx, 10.9, BOTTOM_IN - 2.6, 2.4, "present")
    return True


def _central(ctx, sd) -> bool:
    bullets = _strings(sd.get("bullet_points"))
    if not (3 <= len(bullets) <= 6):
        return False

    cx, cy = 6.55, (TOP_IN + BOTTOM_IN) / 2
    core_d = 2.3
    hub = _shorten(ctx.topic or sd.get("title") or "", 22)

    _shape(ctx, MSO_SHAPE.OVAL, cx - core_d / 2 - 0.22, cy - core_d / 2 - 0.22,
           core_d + 0.44, core_d + 0.44, ctx.accent_soft)
    _shape(ctx, MSO_SHAPE.OVAL, cx - core_d / 2, cy - core_d / 2, core_d, core_d,
           ctx.accent)
    hpt, hh = _fit(hub, core_d - 0.5, (18, 16, 14, 12), 1.4)
    from app.export_builder import _TEXT_WHITE
    _text(ctx, cx - core_d / 2 + 0.25, cy - hh / 2, core_d - 0.5, hh + 0.3,
          hub, hpt, True, _TEXT_WHITE, PP_ALIGN.CENTER, font=ctx.title_font)

    n = len(bullets)
    rx, ry = 4.35, 2.0
    node_w, node_h = 2.75, 0.86
    start = -math.pi / 2 + (math.pi / n if n % 2 == 0 else 0)
    for i, bp in enumerate(bullets):
        a = start + i * (2 * math.pi / n)
        px, py = cx + rx * math.cos(a), cy + ry * math.sin(a)
        ex, ey = cx + (core_d / 2 + 0.1) * math.cos(a), cy + (core_d / 2 + 0.1) * math.sin(a)
        length = math.hypot(px - ex, py - ey)
        ang = math.degrees(math.atan2(py - ey, px - ex))
        _shape(ctx, MSO_SHAPE.RECTANGLE, (ex + px) / 2 - length / 2,
               (ey + py) / 2, length, 0.018, ctx.accent_soft, rotation=ang)
        bpt, bh = _fit(bp, node_w - 0.4, (13, 12, 11), node_h - 0.24)
        _shape(ctx, MSO_SHAPE.ROUNDED_RECTANGLE, px - node_w / 2, py - node_h / 2,
               node_w, max(node_h, bh + 0.24), None, line=ctx.accent, line_pt=1.25,
               adjust=0.2)
        _shape(ctx, MSO_SHAPE.OVAL, px - node_w / 2 - 0.11, py - 0.11, 0.22, 0.22,
               ctx.accent)
        _text(ctx, px - node_w / 2 + 0.2, py - bh / 2 - 0.02, node_w - 0.4,
              bh + 0.25, bp, bpt, False, ctx.ink, PP_ALIGN.CENTER)
    return True


def _flow(ctx, sd) -> bool:
    bullets = _strings(sd.get("bullet_points"))
    if not (3 <= len(bullets) <= 5):
        return False

    n = len(bullets)
    gap = 0.22
    point = 0.18
    step_w = (WIDTH_IN - point - gap * (n - 1)) / n
    step_h = 1.7
    body_probe = _clean(sd.get("body"))
    body_h_probe = 0.0
    if body_probe:
        _, body_h_probe = _fit(body_probe, WIDTH_IN - 3.0, (15, 14, 13, 12), 2.0)
        body_h_probe += 0.42
    top = TOP_IN + max(0.2, (BOTTOM_IN - TOP_IN - step_h - body_h_probe) / 2)
    for i, bp in enumerate(bullets):
        x = LEFT_IN + i * (step_w + gap)
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        _shape(ctx, shp, x, top, step_w + (point if i else 0), step_h,
               ctx.accent_soft if i % 2 == 0 else None,
               line=None if i % 2 == 0 else ctx.accent, line_pt=1.25)
        _shape(ctx, MSO_SHAPE.OVAL, x + step_w / 2 - 0.21, top + 0.2, 0.42, 0.42,
               ctx.accent)
        from app.export_builder import _TEXT_WHITE
        _text(ctx, x + step_w / 2 - 0.21, top + 0.26, 0.42, 0.32, str(i + 1), 13,
              True, _TEXT_WHITE, PP_ALIGN.CENTER)
        bpt, bh = _fit(bp, step_w - 0.6, (14, 13, 12, 11), 0.9)
        _text(ctx, x + 0.3, top + 0.78 + (0.62 - bh) / 2, step_w - 0.6, bh + 0.3,
              bp, bpt, True, ctx.ink, PP_ALIGN.CENTER)

    body = _clean(sd.get("body"))
    y = top + step_h + 0.42
    if body:
        bpt, bh = _fit(body, WIDTH_IN - 3.0, (15, 14, 13, 12), BOTTOM_IN - y - 0.1)
        _text(ctx, LEFT_IN, y, WIDTH_IN - 3.0, bh + 0.3, body, bpt, False, ctx.muted)
    _character(ctx, 11.15, BOTTOM_IN - 1.9, 1.75, "point_left")
    return True


def _bignumber(ctx, sd) -> bool:
    bullets = _strings(sd.get("bullet_points"))
    num, rest = _extract_number(bullets)
    if not num:
        return False

    _text(ctx, LEFT_IN, TOP_IN + 0.5, 5.4, 2.8, num, 150 if len(num) <= 3 else 110,
          True, ctx.accent, PP_ALIGN.LEFT, font=ctx.title_font)
    _shape(ctx, MSO_SHAPE.RECTANGLE, 6.55, TOP_IN + 0.7, 0.045, 3.4, ctx.accent_soft)

    y = TOP_IN + 0.7
    for bp in rest[:4]:
        bpt, bh = _fit(bp, 5.1, (17, 16, 15, 14), 1.2)
        _text(ctx, 6.95, y, 5.1, bh + 0.3, bp, bpt, False, ctx.ink)
        y += bh + 0.42
    body = _clean(sd.get("body"))
    if body and y < BOTTOM_IN - 0.9:
        bpt, bh = _fit(body, 5.1, (14, 13, 12), BOTTOM_IN - y - 0.2)
        _text(ctx, 6.95, y + 0.1, 5.1, bh + 0.3, body, bpt, False, ctx.muted)
    _character(ctx, 2.6, BOTTOM_IN - 2.2, 2.1, "point_right")
    return True


def _split_image(ctx, sd) -> bool:
    if not ctx.image_path:
        return False
    from PIL import Image as PILImage

    on_right = ctx.index % 2 == 0
    half_w = 5.6
    img_x = SLIDE_W_IN - half_w if on_right else 0.0
    band_top, band_h = 1.5, SLIDE_H_IN - 1.5 - 0.55

    _shape(ctx, MSO_SHAPE.RECTANGLE, img_x, band_top, half_w, band_h, ctx.accent_soft)
    try:
        with PILImage.open(ctx.image_path) as im:
            iw, ih = im.size
        scale = max(half_w / iw, band_h / ih)
        tw, th = iw * scale, ih * scale
        px = img_x + (half_w - tw) / 2
        py = band_top + (band_h - th) / 2
        if tw > half_w * 1.6 or th > band_h * 1.6:
            scale = min(half_w / iw, band_h / ih)
            tw, th = iw * scale, ih * scale
            px, py = img_x + (half_w - tw) / 2, band_top + (band_h - th) / 2
        ctx.slide.shapes.add_picture(ctx.image_path, Inches(px), Inches(py),
                                     Inches(tw), Inches(th))
    except Exception:
        return False

    tx = LEFT_IN if on_right else half_w + 0.6
    tw_in = SLIDE_W_IN - half_w - LEFT_IN - 0.75

    y = TOP_IN + 0.35
    bullets = _strings(sd.get("bullet_points"))[:4]
    for i, bp in enumerate(bullets):
        bpt, bh = _fit(bp, tw_in - 0.45, (19, 18, 17, 16), 1.3)
        _shape(ctx, MSO_SHAPE.RECTANGLE, tx, y + 0.06, 0.045,
               max(0.22, bh - 0.04), ctx.accent)
        _text(ctx, tx + 0.26, y, tw_in - 0.45, bh + 0.3, bp, bpt, False, ctx.ink)
        y += bh + 0.34
    body = _clean(sd.get("body"))
    if body and y < BOTTOM_IN - 0.7:
        bpt, bh = _fit(body, tw_in, (15, 14, 13, 12), BOTTOM_IN - y - 0.35)
        _text(ctx, tx, y + 0.12, tw_in, bh + 0.3, body, bpt, False, ctx.muted)
        y += bh + 0.3
    _callout(ctx, sd, tx, y + 0.2, tw_in)
    if ctx.image_credit:
        _text(ctx, img_x + 0.12, SLIDE_H_IN - 0.62, half_w - 0.24, 0.28,
              ctx.image_credit, 8, False, ctx.muted,
              PP_ALIGN.RIGHT if on_right else PP_ALIGN.LEFT)
    return True


def _timeline(ctx, sd) -> bool:
    bullets = _strings(sd.get("bullet_points"))
    if not (3 <= len(bullets) <= 5):
        return False

    y = (TOP_IN + BOTTOM_IN) / 2
    _shape(ctx, MSO_SHAPE.RECTANGLE, LEFT_IN, y, WIDTH_IN, 0.035, ctx.accent)
    n = len(bullets)
    step = WIDTH_IN / n
    for i, bp in enumerate(bullets):
        cx = LEFT_IN + step * (i + 0.5)
        above = i % 2 == 0
        _shape(ctx, MSO_SHAPE.OVAL, cx - 0.15, y - 0.15 + 0.017, 0.3, 0.3, ctx.accent)
        stem_h = 0.85
        _shape(ctx, MSO_SHAPE.RECTANGLE, cx - 0.016,
               y - stem_h if above else y + 0.035, 0.032, stem_h, ctx.accent)
        card_w = step - 0.3
        bpt, bh = _fit(bp, card_w - 0.5, (16, 15, 14, 13), 1.7)
        card_h = max(1.15, bh + 0.6)
        card_y = y - stem_h - card_h if above else y + 0.035 + stem_h
        _shape(ctx, MSO_SHAPE.RECTANGLE, cx - card_w / 2, card_y, card_w, card_h,
               None, line=ctx.accent_soft, line_pt=1.25)
        _shape(ctx, MSO_SHAPE.RECTANGLE, cx - card_w / 2, card_y, card_w, 0.045,
               ctx.accent)
        _text(ctx, cx - card_w / 2 + 0.25, card_y + (card_h - bh) / 2,
              card_w - 0.5, bh + 0.25, bp, bpt, False, ctx.ink, PP_ALIGN.CENTER)
    return True


def _comparison(ctx, sd) -> bool:
    bullets = _strings(sd.get("bullet_points"))
    if len(bullets) not in (2, 4):
        return False
    half = len(bullets) // 2
    left, right = bullets[:half], bullets[half:]

    col_w = (WIDTH_IN - 0.9) / 2
    caps = _split_on_vs(sd.get("title"))

    measured = []
    for items in (left, right):
        h_items = 0.0
        for bp in items:
            _, bh = _fit(bp, col_w - 0.9, (18, 17, 16, 15), 1.6)
            h_items += bh + 0.42
        measured.append(h_items)
    head_h = 0.62 if caps else 0.0
    h = min(BOTTOM_IN - TOP_IN - 0.5,
            max(2.2, max(measured) + head_h + 0.8))
    top = TOP_IN + max(0.25, (BOTTOM_IN - TOP_IN - h) / 2)

    for col, (items, tint) in enumerate(((left, ctx.accent_soft), (right, None))):
        x = LEFT_IN + col * (col_w + 0.9)
        _shape(ctx, MSO_SHAPE.ROUNDED_RECTANGLE, x, top, col_w, h,
               tint, line=None if tint else ctx.accent, line_pt=1.4, adjust=0.06)
        y = top + 0.45
        if caps:
            cpt, chh = _fit(caps[col], col_w - 0.7, (19, 18, 17, 16), 0.5)
            _text(ctx, x + 0.35, y, col_w - 0.7, chh + 0.25, caps[col].capitalize(),
                  cpt, True, ctx.accent_dark, font=ctx.title_font)
            _shape(ctx, MSO_SHAPE.RECTANGLE, x + 0.35, y + chh + 0.14,
                   col_w - 0.7, 0.02, ctx.accent)
            y += chh + 0.42
        for bp in items:
            bpt, bh = _fit(bp, col_w - 0.9, (18, 17, 16, 15), 1.6)
            _shape(ctx, MSO_SHAPE.OVAL, x + 0.35, y + 0.06, 0.2, 0.2, ctx.accent)
            _text(ctx, x + 0.68, y, col_w - 0.9, bh + 0.3, bp, bpt, False, ctx.ink)
            y += bh + 0.42

    mid = LEFT_IN + col_w + 0.45
    _shape(ctx, MSO_SHAPE.RECTANGLE, mid - 0.012, top + 0.2, 0.024, h - 0.4,
           ctx.accent_soft)
    _shape(ctx, MSO_SHAPE.OVAL, mid - 0.3, (top + h / 2) - 0.3, 0.6, 0.6, ctx.accent)
    from app.export_builder import _TEXT_WHITE
    _text(ctx, mid - 0.3, (top + h / 2) - 0.2, 0.6, 0.4, "VS", 13, True,
          _TEXT_WHITE, PP_ALIGN.CENTER)
    return True


def _magazine(ctx, sd) -> bool:
    body = _clean(sd.get("body"))
    if len(body) < 180:
        return False
    bullets = _strings(sd.get("bullet_points"))

    col_w = (WIDTH_IN - 0.7) / 2
    top = TOP_IN + 0.35
    cut = _split_sentences(body)
    first, second = cut

    _text(ctx, LEFT_IN, top - 0.16, 0.7, 1.0, first[:1], 58, True, ctx.accent,
          font=ctx.title_font)
    pt, h1 = _fit(first[1:], col_w - 0.62, (15, 14, 13), BOTTOM_IN - top)
    _text(ctx, LEFT_IN + 0.62, top, col_w - 0.62, h1 + 0.3, first[1:], pt, False,
          ctx.ink)
    if second:
        _, h2 = _fit(second, col_w, (pt,), BOTTOM_IN - top)
        _text(ctx, LEFT_IN + col_w + 0.7, top, col_w, h2 + 0.3, second, pt, False,
              ctx.ink)

    if _callout_of(sd):
        _callout(ctx, sd, LEFT_IN, max(top + max(h1, 1.6) + 0.45, BOTTOM_IN - 1.35),
                 WIDTH_IN)
    elif bullets:
        q = bullets[0]
        qy = max(top + max(h1, 1.6) + 0.45, BOTTOM_IN - 1.5)
        if qy + 1.1 <= BOTTOM_IN:
            _shape(ctx, MSO_SHAPE.RECTANGLE, LEFT_IN, qy, WIDTH_IN, 0.035,
                   ctx.accent_soft)
            qpt, qh = _fit(q, WIDTH_IN - 1.6, (22, 20, 18, 16), 1.0)
            _text(ctx, LEFT_IN + 0.1, qy + 0.22, WIDTH_IN - 1.6, qh + 0.3, q, qpt,
                  True, ctx.accent_dark, font=ctx.title_font)
    return True


def _code(ctx, sd) -> bool:
    code = sd.get("code")
    if not isinstance(code, dict):
        return False
    snippet = str(code.get("snippet") or "").strip()
    if not snippet:
        return False
    from app.export_builder import _TEXT_WHITE

    lines = snippet.replace("\r\n", "\n").split("\n")[:10]
    lang = str(code.get("language") or "").strip() or "code"

    panel_w = 7.4
    line_h = 0.32
    pad = 0.35
    header_h = 0.5
    panel_h = min(BOTTOM_IN - TOP_IN - 0.3, header_h + len(lines) * line_h + pad * 1.6)
    top = TOP_IN + max(0.1, (BOTTOM_IN - TOP_IN - panel_h) / 2)

    dark = RGBColor(0x1E, 0x1E, 0x2E)
    _shape(ctx, MSO_SHAPE.ROUNDED_RECTANGLE, LEFT_IN, top, panel_w, panel_h, dark, adjust=0.04)

    for n, c in enumerate((RGBColor(0xEF, 0x44, 0x44), RGBColor(0xEA, 0xB3, 0x08),
                           RGBColor(0x22, 0xC5, 0x5E))):
        _shape(ctx, MSO_SHAPE.OVAL, LEFT_IN + pad + n * 0.22, top + 0.17, 0.11, 0.11, c)
    _text(ctx, LEFT_IN + panel_w - 1.5, top + 0.1, 1.15, 0.3, lang, 10, True,
          ctx.accent, PP_ALIGN.RIGHT, font="Consolas")

    y = top + header_h
    gutter_w = 0.4
    for i, line in enumerate(lines):
        _text(ctx, LEFT_IN + pad - 0.08, y, gutter_w, line_h,
              str(i + 1), 11, False, RGBColor(0x6B, 0x72, 0x80), font="Consolas")
        _text(ctx, LEFT_IN + pad + gutter_w, y, panel_w - pad * 2 - gutter_w, line_h,
              line if line.strip() else " ", 12, False, _TEXT_WHITE, font="Consolas")
        y += line_h

    tx = LEFT_IN + panel_w + 0.45
    tw = RIGHT_IN - tx
    explanation = _clean(code.get("explanation")) or _clean(sd.get("body"))
    bullets = _strings(sd.get("bullet_points"))
    by = top + 0.1
    for bp in bullets[:3]:
        bpt, bh = _fit(bp, tw - 0.3, (15, 14, 13), 1.3)
        _shape(ctx, MSO_SHAPE.OVAL, tx, by + 0.08, 0.14, 0.14, ctx.accent)
        _text(ctx, tx + 0.28, by, tw - 0.28, bh + 0.25, bp, bpt, False, ctx.ink)
        by += bh + 0.32
    if explanation:
        bpt, bh = _fit(explanation, tw, (14, 13, 12), max(0.8, top + panel_h - by))
        _text(ctx, tx, by + 0.15, tw, bh + 0.3, explanation, bpt, False, ctx.muted)

    _callout(ctx, sd, LEFT_IN, top + panel_h + 0.3, panel_w)
    return True


def _cards(ctx, sd) -> bool:
    return False




def _strings(value) -> list[str]:
    if not isinstance(value, (list, tuple)):
        return []
    out = []
    for v in value:
        s = re.sub(r"\s+", " ", str(v or "")).strip()
        if s:
            out.append(s)
    return out


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _shorten(s, n):
    s = _clean(s)
    if len(s) <= n:
        return s
    cut = s[:n]
    space = cut.rfind(" ")
    if space >= n * 0.5:
        cut = cut[:space]
    return cut.rstrip(" ,.;:—-") + "…"


def _split_sentences(body):
    parts = re.split(r"(?<=[.!?])\s+", body)
    if len(parts) < 2:
        return body, ""
    half = len(body) / 2
    acc, i = 0, 0
    for i, p in enumerate(parts):
        acc += len(p)
        if acc >= half:
            break
    return " ".join(parts[: i + 1]), " ".join(parts[i + 1:])


_VS_WORDS = (" vs ", " и ", " или ", " ва ", " ё ", " versus ")


def _split_on_vs(title):
    t = _clean(title)
    if not t:
        return None
    low = " " + t.lower() + " "
    for w in _VS_WORDS:
        i = low.find(w)
        if i <= 0:
            continue
        left = t[: i - 1].strip(" ,")
        right = t[i - 1 + len(w):].strip(" ,")
        if not left or not right:
            continue
        rw = right.split()
        if len(rw) > 1:
            right = " ".join(rw[:-1])
        if len(left) > 26 or len(right) > 26:
            return None
        return left, right
    return None


def _looks_like_comparison(sd) -> bool:
    bullets = _strings(sd.get("bullet_points"))
    if len(bullets) not in (2, 4):
        return False
    title = " " + _clean(sd.get("title")).lower() + " "
    return any(w in title for w in _VS_WORDS)


_NUM_RE = re.compile(r"^[^\d]{0,12}?(\d[\d\s.,]{0,9}\s*%?)\s*(.*)$")


def _extract_number(bullets):
    for i, bp in enumerate(bullets):
        m = _NUM_RE.match(bp)
        if not m:
            continue
        num = m.group(1).strip()
        if len(bp) > 42 or not num or len(num) > 7:
            continue
        rest = [b for j, b in enumerate(bullets) if j != i]
        if m.group(2).strip():
            rest = [m.group(2).strip()] + rest
        if not rest:
            continue
        return num, rest
    return "", []



_BY_KIND = {
    "intro":       ("split_image", "hero", "magazine"),
    "concepts":    ("central", "split_image", "comparison"),
    "explanation": ("split_image", "magazine", "flow", "central"),
    "example":     ("flow", "bignumber", "split_image"),
    "task":        ("flow", "comparison"),
    "summary":     ("timeline", "central", "hero"),
}

_BY_SUBJECT = {
    "history":       ("timeline",),
    "history_world": ("timeline",),
    "geography":     ("split_image",),
    "mathematics":   ("flow", "bignumber"),
    "algebra":       ("flow",),
    "geometry":      ("split_image",),
    "informatics":   ("flow",),
    "social_studies": ("comparison", "bignumber"),
    "ecology":       ("central", "flow"),
}

_LAYOUTS = {
    "hero": _hero,
    "central": _central,
    "flow": _flow,
    "bignumber": _bignumber,
    "split_image": _split_image,
    "timeline": _timeline,
    "comparison": _comparison,
    "magazine": _magazine,
    "code": _code,
    "cards": _cards,
}


def draw(ctx: Ctx, sd: dict, last_layout: str | None) -> str | None:
    kind = str(sd.get("kind") or "").strip().lower()

    order: list[str] = []
    for name in _BY_SUBJECT.get(ctx.template_id, ()):
        if name not in order:
            order.append(name)
    for name in _BY_KIND.get(kind, ()):
        if name not in order:
            order.append(name)
    for name in ("split_image", "central", "flow", "comparison", "timeline",
                 "bignumber", "magazine", "hero"):
        if name not in order:
            order.append(name)

    if _extract_number(_strings(sd.get("bullet_points")))[0] and "bignumber" in order:
        order.remove("bignumber")
        order.insert(0, "bignumber")

    if _looks_like_comparison(sd) and "comparison" in order:
        order.remove("comparison")
        order.insert(0, "comparison")

    if ctx.image_path and "split_image" in order:
        order.remove("split_image")
        order.insert(0, "split_image")

    if isinstance(sd.get("code"), dict) and str(sd["code"].get("snippet") or "").strip():
        order = ["code"] + [n for n in order if n != "code"]

    if last_layout and last_layout in order and len(order) > 1:
        order.remove(last_layout)
        order.append(last_layout)

    if ctx.grade_tier == "senior" and "magazine" in order:
        order.remove("magazine")
        order.insert(1 if ctx.image_path else 0, "magazine")
    elif ctx.grade_tier == "junior":
        for name in ("magazine",):
            if name in order:
                order.remove(name)

    for name in order:
        fn = _LAYOUTS.get(name)
        if fn is None:
            continue
        try:
            if fn(ctx, sd):
                return name
        except Exception:
            continue
    return None


def grade_tier(grade) -> str:
    m = re.search(r"\d+", str(grade or ""))
    if not m:
        return "middle"
    try:
        n = int(m.group())
    except ValueError:
        return "middle"
    if n <= 4:
        return "junior"
    return "senior" if n >= 10 else "middle"
