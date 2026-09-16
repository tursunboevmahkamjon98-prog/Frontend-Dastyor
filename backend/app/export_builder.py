import hashlib
import io
import math
import os
import re
import platform as _platform
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.platypus import (SimpleDocTemplate, Paragraph as _RLParagraph,
                                Spacer, Table, TableStyle, KeepTogether)
from reportlab.platypus.flowables import Flowable as _RLFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader




def _xml_escape(s: str) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


_RL_IMG_TAG = re.compile(r"<img\b[^>]*/?>", re.IGNORECASE)
_RL_STYLE_TAG = re.compile(r"</?(?:b|i|super|sub|font|para)\b[^>]*>", re.IGNORECASE)


def _strip_reportlab_markup(text: str) -> str:
    text = _RL_IMG_TAG.sub("", str(text or ""))
    text = _RL_STYLE_TAG.sub("", text)
    return text


def _math_markup(text, size: float = 11, bg: str | None = None) -> str:
    raw = _normalize_math(text)
    if "$" not in raw:
        return raw

    def repl(m):
        from app.math_render import script_segments, is_literal_operator
        latex = m.group(1)
        if is_literal_operator(latex):
            return _xml_escape(latex)
        try:
            segments = script_segments(latex)
        except Exception:
            segments = None
        if segments:
            out = []
            for seg_text, kind in segments:
                esc = _xml_escape(seg_text)
                if kind == "sup":
                    out.append(f"<super>{esc}</super>")
                elif kind == "sub":
                    out.append(f"<sub>{esc}</sub>")
                else:
                    out.append(esc)
            return "".join(out)
        got = _math_png(latex, size, bg)
        if not got:
            return _xml_escape(latex)
        path, w, asc, desc = got
        return (f'<img src="{path}" width="{w + 2 * _MATH_PAD:.1f}" '
                f'height="{asc + desc + 2 * _MATH_PAD:.1f}" '
                f'valign="{-(desc + _MATH_PAD):.1f}"/>')
    return _MATH_SPAN.sub(repl, raw)


class Paragraph(_RLParagraph):

    def __init__(self, text, style=None, *args, **kwargs):
        size = 11
        try:
            size = float(getattr(style, "fontSize", 11) or 11)
        except Exception:
            pass
        is_code = getattr(style, "fontName", None) == CODE_FONT
        original_text = text
        marked_up = text
        try:
            if not is_code:
                marked_up = _math_markup(text, size)
        except Exception:
            marked_up = text
        try:
            super().__init__(marked_up, style, *args, **kwargs)
        except Exception as e:
            logger.warning(f"Paragraph markup failed to parse (falling back to plain text): {e}")
            super().__init__(_xml_escape(_strip_reportlab_markup(original_text)), style, *args, **kwargs)

from app.logger import get_logger
from app.subject_theme import get_subject_accent_hex, get_subject_accent_rgb, get_subject_illustration_path
from app import slide_decor, pptx_shapes, slide_layouts, slide_characters
from app.konspekt_templates import get_template

logger = get_logger(__name__)

_GROUP_LABEL_RE = re.compile(
    r"^\s*(?:Группа|Гурӯҳи?|Group|Guruh)\s*\d+\s*[:.\-—]\s*", re.IGNORECASE
)


def _strip_group_label(task: str) -> str:
    return _GROUP_LABEL_RE.sub("", str(task), count=1)


def _register_unicode_font():
    font_name = 'Arial'
    font_name_bold = 'Arial-Bold'
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name, font_name_bold
    try:
        font_paths = []
        font_bold_paths = []
        if _platform.system() == 'Windows':
            windir = os.environ.get('WINDIR', r'C:\Windows')
            font_paths = [os.path.join(windir, 'Fonts', 'arial.ttf')]
            font_bold_paths = [os.path.join(windir, 'Fonts', 'arialbd.ttf')]
        elif _platform.system() == 'Darwin':
            font_paths = ['/Library/Fonts/Arial.ttf']
            font_bold_paths = ['/Library/Fonts/Arial Bold.ttf']
        else:
            font_paths = ['/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/TTF/DejaVuSans.ttf']
            font_bold_paths = ['/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', '/usr/share/fonts/TTF/DejaVuSans-Bold.ttf']

        for path in font_paths:
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont(font_name, path))
                break

        bold_registered = False
        for path in font_bold_paths:
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont(font_name_bold, path))
                bold_registered = True
                break

        if not bold_registered:
            font_name_bold = font_name

        from reportlab.pdfbase.pdfmetrics import registerFontFamily
        registerFontFamily(font_name, normal=font_name, bold=font_name_bold)

        return font_name, font_name_bold
    except Exception:
        return 'Helvetica', 'Helvetica-Bold'


FONT_NAME, FONT_NAME_BOLD = _register_unicode_font()


def _register_math_font():
    name = 'CambriaMath'
    if name in pdfmetrics.getRegisteredFontNames():
        return name, name
    try:
        windir = os.environ.get('WINDIR', r'C:\Windows') if _platform.system() == 'Windows' else None
        if windir:
            regular_path = os.path.join(windir, 'Fonts', 'cambria.ttc')
            if os.path.exists(regular_path):
                pdfmetrics.registerFont(TTFont(name, regular_path))
                return name, name
    except Exception:
        pass
    return FONT_NAME, FONT_NAME_BOLD


MATH_FONT, MATH_FONT_BOLD = _register_math_font()


def _register_code_font():
    name = 'Consolas'
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    try:
        windir = os.environ.get('WINDIR', r'C:\Windows') if _platform.system() == 'Windows' else None
        if windir:
            path = os.path.join(windir, 'Fonts', 'consola.ttf')
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont(name, path))
                return name
    except Exception:
        pass
    return FONT_NAME


CODE_FONT = _register_code_font()


_TAJIK_PROBE = "ғқҳҷӣӯ"


def verify_pdf_fonts() -> list[str]:
    problems: list[str] = []
    try:
        registered = pdfmetrics.getRegisteredFontNames()
        if FONT_NAME not in registered:
            return [
                f"the PDF font '{FONT_NAME}' was never registered — falling back "
                f"to a built-in face with no Cyrillic. On Linux install "
                f"fonts-dejavu-core."
            ]
        face = pdfmetrics.getFont(FONT_NAME).face
        missing = [c for c in _TAJIK_PROBE
                   if getattr(face, "charToGlyph", {}).get(ord(c), 0) == 0]
        if missing:
            problems.append(
                f"the PDF font '{FONT_NAME}' cannot draw {''.join(missing)} — "
                f"Tajik PDFs will contain empty boxes. Install fonts-dejavu-core."
            )
    except Exception as e:
        problems.append(f"could not verify the PDF fonts: {e}")
    return problems


def verify_pptx_renderer() -> str | None:
    try:
        from app import pptx_pdf
        if pptx_pdf.find_binary() is None:
            return ("LibreOffice was not found — presentation previews will fall "
                    "back to the A4 rendering instead of the real slides. "
                    "Install libreoffice-impress.")
    except Exception as e:
        return f"could not check for LibreOffice: {e}"
    return None


from pptx.enum.shapes import MSO_SHAPE



_ACCENT = RGBColor(0x3B, 0x82, 0xF6)
_ACCENT_DARK = RGBColor(0x1E, 0x3A, 0x8A)
_ACCENT_SOFT = RGBColor(0xDB, 0xEA, 0xFE)
_INK = RGBColor(0x1E, 0x29, 0x37)


def _pptx_accent_shades(rgb: tuple[int, int, int]) -> tuple[RGBColor, RGBColor, RGBColor]:
    r, g, b = rgb
    dark = RGBColor(int(r * 0.55), int(g * 0.55), int(b * 0.55))
    soft = RGBColor(int(r + (255 - r) * 0.88), int(g + (255 - g) * 0.88), int(b + (255 - b) * 0.88))
    return RGBColor(r, g, b), dark, soft
_MUTED = RGBColor(0x6B, 0x72, 0x80)
_HAIRLINE = RGBColor(0xE2, 0xE8, 0xF0)
_TEXT_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_TEXT_SOFT = RGBColor(0xDB, 0xEA, 0xFE)

_ICON_SHAPES = [MSO_SHAPE.LIGHTNING_BOLT, MSO_SHAPE.STAR_5_POINT, MSO_SHAPE.HEXAGON, MSO_SHAPE.DIAMOND, MSO_SHAPE.CHEVRON]
_PLAYFUL_ICON_SHAPES = [MSO_SHAPE.STAR_5_POINT, MSO_SHAPE.HEART, MSO_SHAPE.CLOUD, MSO_SHAPE.SUN, MSO_SHAPE.SMILEY_FACE]

_PLAYFUL_PALETTE = [
    RGBColor(0xBF, 0xDB, 0xFE),
    RGBColor(0xBB, 0xF7, 0xD0),
    RGBColor(0xDD, 0xD6, 0xFE),
    RGBColor(0xFB, 0xCF, 0xE8),
    RGBColor(0xFD, 0xE6, 0x8A),
]
_PLAYFUL_PALETTE_DARK = [
    RGBColor(0x1D, 0x4E, 0xD8),
    RGBColor(0x15, 0x80, 0x3D),
    RGBColor(0x6D, 0x28, 0xD9),
    RGBColor(0xBE, 0x18, 0x5D),
    RGBColor(0x92, 0x6B, 0x00),
]



_DECK_THEMES = {
    "playful": {
        "bg": (0xFF, 0xFD, 0xF7),
        "ink": RGBColor(0x2D, 0x2A, 0x4A),
        "muted": RGBColor(0x6B, 0x63, 0x8C),
        "title_font": "Comic Sans MS",
        "body_font": "Arial",
        "title_caps": False,
        "header": "rule",
        "cards": True,
        "badge": True,
        "pastel": True,
        "spiral": True,
    },
    "google": {
        "bg": (0xFF, 0xFF, 0xFF),
        "ink": RGBColor(0x20, 0x21, 0x24),
        "muted": RGBColor(0x5F, 0x63, 0x68),
        "title_font": "Arial",
        "body_font": "Arial",
        "title_caps": False,
        "header": "underline",
        "cards": False,
        "badge": False,
    },
    "zamonaviy": {
        "bg": (0xFF, 0xFF, 0xFF),
        "ink": RGBColor(0x1E, 0x29, 0x37),
        "muted": RGBColor(0x6B, 0x72, 0x80),
        "title_font": "Arial",
        "body_font": "Arial",
        "title_caps": False,
        "header": "rule",
        "cards": True,
        "badge": True,
    },
    "klassik": {
        "bg": (0xFC, 0xFA, 0xF5),
        "ink": RGBColor(0x1F, 0x1B, 0x16),
        "muted": RGBColor(0x7A, 0x6E, 0x60),
        "title_font": "Georgia",
        "body_font": "Georgia",
        "title_caps": False,
        "header": "underline",
        "cards": False,
        "badge": False,
    },
    "rangli": {
        "bg": (0xFF, 0xFF, 0xFF),
        "ink": RGBColor(0x11, 0x18, 0x27),
        "muted": RGBColor(0x64, 0x74, 0x8B),
        "title_font": "Arial",
        "body_font": "Arial",
        "title_caps": False,
        "header": "band",
        "cards": True,
        "badge": True,
    },
    "minimal": {
        "bg": (0xFF, 0xFF, 0xFF),
        "ink": RGBColor(0x0F, 0x17, 0x2A),
        "muted": RGBColor(0x94, 0xA3, 0xB8),
        "title_font": "Arial",
        "body_font": "Arial",
        "title_caps": True,
        "header": "plain",
        "cards": False,
        "badge": False,
    },
}
_DECK_THEME_DEFAULT = "playful"


def deck_theme(template_id) -> dict:
    return _DECK_THEMES.get(str(template_id or "").strip().lower(), _DECK_THEMES[_DECK_THEME_DEFAULT])


def _hex_rgb(value) -> RGBColor:
    if isinstance(value, RGBColor):
        return value
    h = str(value).lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _hex_tuple(value) -> tuple[int, int, int]:
    if isinstance(value, (tuple, list)):
        return (int(value[0]), int(value[1]), int(value[2]))
    h = str(value).lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _normalise_theme(theme: dict) -> dict:
    out = dict(theme)
    out["bg"] = _hex_tuple(theme["bg"])
    out["ink"] = _hex_rgb(theme["ink"])
    out["muted"] = _hex_rgb(theme["muted"])
    return out


def resolve_deck_theme(content: dict) -> dict:
    from app import subject_templates

    stored = str(content.get("template") or "").strip().lower()
    if stored in _DECK_THEMES:
        theme = dict(_DECK_THEMES[stored])
        theme.setdefault("subject_template", "")
        theme["accent"] = None
        return _normalise_theme(theme)

    if stored in subject_templates.TEMPLATES:
        tpl = subject_templates.TEMPLATES[stored]
        theme = tpl.as_deck_theme()
        language = content.get("language")
        theme["title_font"] = subject_templates.safe_font(theme["title_font"], language)
        theme["body_font"] = subject_templates.safe_font(theme["body_font"], language)
        return _normalise_theme(theme)

    return _normalise_theme(subject_templates.deck_theme_for(
        content.get("subject"), content.get("grade"), content.get("language")))


_CONTENT_ICON_KEYWORDS: list[tuple[tuple[str, ...], "MSO_SHAPE"]] = [
    (("ғоя", "идея", "g'oya", "goya", "idea", "мисол"), MSO_SHAPE.SUN),
    (("таърих", "тарих", "история", "tarix", "history", "созанда", "эҷод", "yaratil", "создан"), MSO_SHAPE.FLOWCHART_DOCUMENT),
    (("оянда", "будущ", "kelajak", "future", "рушд", "развит", "rivojlan", "growth", "ўсиш", "усиш"), MSO_SHAPE.UP_ARROW),
    (("савол", "вопрос", "savol", "question", "муаммо", "проблем", "muammo", "problem", "чаро", "почему", "nima uchun", "why"), MSO_SHAPE.ACTION_BUTTON_HELP),
    (("маълумот", "информ", "ma'lumot", "malumot", "information", "факт", "fact", "статистик"), MSO_SHAPE.ACTION_BUTTON_INFORMATION),
    (("раванд", "процесс", "jarayon", "process", "механизм", "mexanizm", "mechanism", "алгоритм", "algorithm", "кор мекунад", "ишлайди", "работает", "works"), MSO_SHAPE.GEAR_9),
    (("суръат", "скорост", "tezlik", "speed", "қувва", "энерги", "kuch", "energy", "quvvat"), MSO_SHAPE.LIGHTNING_BOLT),
    (("хатар", "опасн", "xavf", "danger", "risk", "амният", "безопасност", "xavfsizlik", "security"), MSO_SHAPE.NO_SYMBOL),
    (("шабака", "интернет", "сеть", "tarmoq", "internet", "network", "онлайн", "online", "abr", "cloud", "облак"), MSO_SHAPE.CLOUD),
    (("китобхона", "библиотек", "kutubxona", "library", "асбоб", "инструмент", "vosita", "tool"), MSO_SHAPE.FLOWCHART_PREDEFINED_PROCESS),
    (("сохт", "структур", "tuzilish", "structure", "модел", "model", "куб", "cube"), MSO_SHAPE.CUBE),
    (("бозӣ", "игра", "o'yin", "oyin", "game", "fun", "фан", "лаззат"), MSO_SHAPE.SMILEY_FACE),
    (("муҳаббат", "любов", "sevgi", "love", "care", "ғамхорӣ", "забота"), MSO_SHAPE.HEART),
    (("дастовард", "достижен", "yutuq", "achievement", "муваффақ", "успех", "muvaffaqiyat", "success", "беҳтарин", "лучш", "eng yaxshi", "best"), MSO_SHAPE.STAR_5_POINT),
    (("интишор", "выпуск", "release", "e'lon", "launch", "оғоз", "начал", "boshlash", "start"), MSO_SHAPE.RIGHT_ARROW),
]
def _pick_content_icon(*texts: str):
    joined = " ".join(t for t in texts if t).lower()
    for keywords, shape in _CONTENT_ICON_KEYWORDS:
        if any(kw in joined for kw in keywords):
            return shape
    return None


_DECK_AGENDA_LABEL = {
    "Русский": "В ЭТОМ УРОКЕ", "Таджикский": "ДАР ИН ДАРС",
    "English": "IN THIS LESSON",
    "Английский": "IN THIS LESSON",
}

_DECK_AGENDA_MORE = {
    "Русский": "ещё {n} {word}", "Таджикский": "боз {n} слайд",
    "English": "{n} more {word}",
    "Английский": "{n} more {word}",
}


def _agenda_more_word(language: str, n: int) -> str:
    lang = str(language or "Русский")
    if lang == "Русский":
        if n % 100 in (11, 12, 13, 14):
            return "слайдов"
        last = n % 10
        if last == 1:
            return "слайд"
        if last in (2, 3, 4):
            return "слайда"
        return "слайдов"
    if lang in ("English", "Английский"):
        return "slide" if n == 1 else "slides"
    return ""

_SLIDE_KIND_LABEL = {
    "intro": {"Русский": "ВВЕДЕНИЕ", "Таджикский": "МУҚАДДИМА",
              "English": "INTRODUCTION",
              "Английский": "INTRODUCTION"},
    "concepts": {"Русский": "КЛЮЧЕВЫЕ ПОНЯТИЯ", "Таджикский": "МАФҲУМҲОИ АСОСӢ",
                 "English": "KEY CONCEPTS",
                 "Английский": "KEY CONCEPTS"},
    "example": {"Русский": "ПРИМЕР", "Таджикский": "МИСОЛ",
                "English": "EXAMPLE",
                "Английский": "EXAMPLE"},
    "task": {"Русский": "ЗАДАНИЕ", "Таджикский": "ВАЗИФА",
             "English": "TASK",
             "Английский": "TASK"},
    "summary": {"Русский": "ИТОГ", "Таджикский": "ХУЛОСА",
                "English": "SUMMARY",
                "Английский": "SUMMARY"},
}

_KIND_PANEL = {
    "example": "tinted",
    "task": "dashed",
    "summary": "banded",
}


def _kind_label(kind: str, language) -> str:
    table = _SLIDE_KIND_LABEL.get(kind)
    if not table:
        return ""
    return table.get(str(language or "Русский"), table.get("Русский", ""))


def _draw_kind_panel(slide, style: str, top_in: float, height_in: float,
                     accent: RGBColor, accent_soft: RGBColor) -> None:
    x_in, w_in = 0.66, 11.98
    pad = 0.28
    y_in = max(1.45, top_in - pad)
    h_in = min(6.98 - y_in, height_in + pad * 2)
    if h_in <= 0.3:
        return
    if style == "tinted":
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_in), Inches(y_in),
                   Inches(w_in), Inches(h_in), accent_soft)
    elif style == "dashed":
        shp = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_in), Inches(y_in),
                         Inches(w_in), Inches(h_in), _TEXT_WHITE,
                         line_rgb=accent, line_width=Pt(1.25))
        try:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        except Exception:
            pass
    elif style == "banded":
        _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x_in), Inches(y_in),
                   Inches(w_in), Inches(h_in), accent_soft)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x_in), Inches(y_in),
                   Pt(5), Inches(h_in), accent)


_LOGO_PATH = os.path.join(os.path.dirname(__file__), "static", "logo.png")

_SLIDE_W = Inches(13.333)
_SLIDE_H = Inches(7.5)


def _set_slide_bg(slide, r, g, b):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(r, g, b)


def _add_shape(slide, shape_type, left, top, width, height, fill_rgb, line_rgb=None, line_width=Pt(1)):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = line_rgb
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    pptx_shapes.flatten(shape)
    return shape


def _set_shape_alpha(shape, alpha_pct: int):
    from pptx.oxml.ns import qn
    sp_pr = shape.fill._xPr
    solid_fill = sp_pr.find(qn('a:solidFill'))
    if solid_fill is None:
        return
    color_el = solid_fill.find(qn('a:srgbClr'))
    if color_el is None:
        return
    alpha_el = color_el.makeelement(qn('a:alpha'), {'val': str(alpha_pct * 1000)})
    color_el.append(alpha_el)


def _add_gradient_fill(shape, rgb1, rgb2, angle=45):
    fill = shape.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = rgb1
    stops[0].position = 0.0
    stops[-1].color.rgb = rgb2
    stops[-1].position = 1.0
    fill.gradient_angle = angle
    return shape


def _trim_text_to_lines(text: str, width_in: float, font_size: float, max_lines: int) -> str:
    words = str(text or "").split()
    if not words or _estimate_pptx_lines(text, width_in, font_size) <= max_lines:
        return text
    kept: list[str] = []
    for word in words:
        trial = " ".join(kept + [word]) + "…"
        if _estimate_pptx_lines(trial, width_in, font_size) > max_lines:
            break
        kept.append(word)
    return (" ".join(kept) + "…") if kept else ""


_CYRILLIC_RE = re.compile(r"[Ѐ-ӿ]")


def _estimate_pptx_lines(text, width_in, font_size):
    factor = 0.62 if _CYRILLIC_RE.search(str(text)) else 0.52
    avg_char_width_in = (font_size * factor) / 72.0
    chars_per_line = max(8, int(width_in / avg_char_width_in))
    return max(1, -(-len(text) // chars_per_line))


def _add_logo_badge(slide, on_dark_bg=False):
    if not os.path.exists(_LOGO_PATH):
        return
    diameter = Inches(0.62)
    left = _SLIDE_W - diameter - Inches(0.4)
    top = _SLIDE_H - diameter - Inches(0.35)
    line_rgb = _ACCENT_SOFT if not on_dark_bg else None
    _add_shape(slide, MSO_SHAPE.OVAL, left, top, diameter, diameter, _TEXT_WHITE, line_rgb=line_rgb, line_width=Pt(0.75))
    pad = Emu(int(diameter * 0.17))
    slide.shapes.add_picture(_LOGO_PATH, left + pad, top + pad, width=diameter - pad * 2, height=diameter - pad * 2)


def _pptx_sticker_decorations(slide) -> None:
    book_colors = [RGBColor(0xF9, 0xA8, 0xD4), RGBColor(0x93, 0xC5, 0xFD), RGBColor(0x86, 0xEF, 0xAC)]
    x0, y0, w0 = 1.0, 5.35, 1.5
    y = y0
    for i, color in enumerate(book_colors):
        h = 0.22
        w = w0 - i * 0.12
        book = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0), Inches(y), Inches(w), Inches(h),
                           color, line_rgb=RGBColor(0xFF, 0xFF, 0xFF), line_width=Pt(1.5))
        book.rotation = -4 + i * 3
        y += h - 0.03
    heart = _add_shape(slide, MSO_SHAPE.HEART, Inches(x0 + w0 + 0.25), Inches(y0 - 0.15), Inches(0.42), Inches(0.42),
                        RGBColor(0xFB, 0x71, 0x85))
    heart.rotation = -12

    pencil_x, pencil_y = x0 + w0 + 0.95, y0 + 0.05
    body = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(pencil_x), Inches(pencil_y), Inches(0.85), Inches(0.16),
                       RGBColor(0xFD, 0xE0, 0x47), line_rgb=RGBColor(0xB4, 0x8A, 0x00), line_width=Pt(1.0))
    body.rotation = 35
    tip = _add_shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(pencil_x + 0.62), Inches(pencil_y + 0.30),
                      Inches(0.18), Inches(0.16), RGBColor(0x8B, 0x5E, 0x34))
    tip.rotation = 125


def _pptx_washi_tape(slide, x_in: float, y_in: float, color: RGBColor, rotation: float = -6) -> None:
    tape = _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x_in), Inches(y_in), Inches(0.85), Inches(0.24), color)
    tape.rotation = rotation
    _set_shape_alpha(tape, 60)


def _pptx_notebook_lines(slide) -> None:
    line_color = RGBColor(0xEE, 0xE4, 0xC8)
    y = 0.55
    while y < 7.3:
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, Inches(y), _SLIDE_W, Pt(0.75), line_color)
        y += 0.32


def _pptx_spiral_margin(slide) -> None:
    hole_d = Inches(0.14)
    hole_x = Inches(0.28) - hole_d / 2
    hole_line = RGBColor(0xD8, 0xD2, 0xC4)
    sparkle_shapes = [MSO_SHAPE.HEART, MSO_SHAPE.STAR_5_POINT]
    sparkle_colors = [RGBColor(0xFB, 0x71, 0x85), RGBColor(0xFB, 0xBF, 0x24)]
    y = 0.35
    n = 0
    while y < 7.15:
        if n % 4 == 3:
            d = Inches(0.16)
            k = (n // 4) % len(sparkle_shapes)
            spark = _add_shape(slide, sparkle_shapes[k], Inches(0.28) - d / 2, Inches(y) - d / 2 + hole_d / 2,
                               d, d, sparkle_colors[k])
            spark.rotation = -10 if k == 0 else 0
        else:
            _add_shape(slide, MSO_SHAPE.OVAL, hole_x, Inches(y), hole_d, hole_d,
                       RGBColor(0xFF, 0xFD, 0xF7), line_rgb=hole_line, line_width=Pt(1.0))
        y += 0.55
        n += 1
    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(0.15), Pt(1.25), Inches(7.2),
               RGBColor(0xF3, 0xC6, 0xC6))


_MARKER_BADGE_SHAPE = {
    "number": MSO_SHAPE.OVAL,
    "circle": MSO_SHAPE.OVAL,
    "dot": MSO_SHAPE.OVAL,
    "square": MSO_SHAPE.ROUNDED_RECTANGLE,
    "hexagon": MSO_SHAPE.HEXAGON,
    "diamond": MSO_SHAPE.DIAMOND,
    "chevron": MSO_SHAPE.PENTAGON,
    "pin": MSO_SHAPE.OVAL,
    "roman": MSO_SHAPE.RECTANGLE,
    "bracket": MSO_SHAPE.ROUNDED_RECTANGLE,
    "dash": MSO_SHAPE.OVAL,
    "star": MSO_SHAPE.OVAL,
    "leaf": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
}

_MARKER_GLYPH_SHAPE = {
    "chevron": MSO_SHAPE.CHEVRON,
    "diamond": MSO_SHAPE.DIAMOND,
    "hexagon": MSO_SHAPE.HEXAGON,
    "star": MSO_SHAPE.STAR_5_POINT,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "leaf": MSO_SHAPE.DIAMOND,
    "pin": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "square": MSO_SHAPE.RECTANGLE,
}

_ROMAN = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII")


def _marker_label(marker: str, n: int) -> str:
    if marker == "roman":
        return _ROMAN[(n - 1) % len(_ROMAN)]
    if marker == "bracket":
        return f"[{n}]"
    if marker == "dash":
        return "—"
    return str(n)


def _add_icon_badge(slide, cx, cy, diameter, shape_type, bg_rgb=_ACCENT,
                    icon_rgb=_TEXT_WHITE, badge_shape=MSO_SHAPE.OVAL):
    _add_shape(slide, badge_shape, cx - diameter / 2, cy - diameter / 2, diameter, diameter, bg_rgb)
    inner = diameter * 0.42
    _add_shape(slide, shape_type, cx - inner / 2, cy - inner / 2, inner, inner, icon_rgb)


def _add_number_badge(slide, cx, cy, diameter, number: int, bg_rgb=_ACCENT,
                      text_rgb=_TEXT_WHITE, badge_shape=MSO_SHAPE.OVAL,
                      label: str | None = None):
    text = label if label is not None else str(number)
    _add_shape(slide, badge_shape, cx - diameter / 2, cy - diameter / 2,
               diameter, diameter, bg_rgb)
    size = 13 if len(text) <= 2 else (11 if len(text) == 3 else 9)
    _add_text(slide, cx - diameter / 2, cy - diameter / 2 + Emu(int(diameter * 0.12)),
              diameter, diameter, text, size, True, text_rgb, PP_ALIGN.CENTER)


def _add_transition(slide, duration_ms=600):
    from pptx.oxml.ns import qn
    sld = slide._element
    existing = sld.find(qn('p:transition'))
    if existing is not None:
        sld.remove(existing)
    transition = sld.makeelement(qn('p:transition'), {'spd': 'med', 'dur': str(duration_ms)})
    fade = transition.makeelement(qn('p:fade'), {})
    transition.append(fade)
    sld.append(transition)


def _add_text(slide, left, top, width, height, text, font_size=14, bold=False, color=_INK, align=PP_ALIGN.LEFT, wrap=True, font_name="Arial"):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf


def _add_multiline(slide, left, top, width, height, lines, font_size=14, color=_INK, line_spacing=Pt(8)):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = line_spacing
    return tf


_GRID_BADGE_D_IN = 0.56
_GRID_BADGE_PROTRUDE_IN = 0.18


_CONTENT_X_IN = 0.9
_CONTENT_W_IN = 11.5


def _grid_columns(n: int, kind: str = "") -> int:
    if kind == "concepts" and n >= 3:
        return 3
    return 3 if n in (3, 6) else 2


def _pptx_bullet_grid_natural_heights(bullets: list[str], col_w_in: float, font_size: int = 14, kind: str = "") -> list[float]:
    pad_in = 0.22
    n = len(bullets)
    cols = _grid_columns(n, kind)
    if cols == 3:
        col_w_in = (col_w_in * 2 + 0.3 - 0.3 * 2) / 3
    text_w_in = col_w_in - pad_in * 2
    line_h_in = font_size * 1.3 / 72.0
    rows = -(-n // cols)
    row_heights = []
    for r in range(rows):
        pair = bullets[r * cols:r * cols + cols]
        max_lines = max(_estimate_pptx_lines(bp, text_w_in, font_size) for bp in pair)
        text_h_in = max_lines * line_h_in
        row_heights.append((_GRID_BADGE_D_IN - _GRID_BADGE_PROTRUDE_IN) + 0.16 + text_h_in + pad_in)
    return row_heights


def _draw_bullet_card(slide, card: str, x_in: float, y_in: float, w_in: float,
                      h_in: float, accent: RGBColor, border: RGBColor,
                      support: RGBColor) -> None:
    L, T, W, H = Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in)

    if card == "sharp":
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, H, _TEXT_WHITE,
                   line_rgb=border, line_width=Pt(1.25))
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, Pt(3.5), H, accent)
    elif card == "tab":
        shp = _add_shape(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, L, T, W, H,
                         _TEXT_WHITE, line_rgb=border, line_width=Pt(1.25))
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T + H - Pt(3), W, Pt(3), accent)
    elif card == "hex":
        _add_shape(slide, MSO_SHAPE.SNIP_2_DIAG_RECTANGLE, L, T, W, H,
                   _TEXT_WHITE, line_rgb=border, line_width=Pt(1.25))
    elif card == "pill":
        shp = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, L, T, W, H,
                         _TEXT_WHITE, line_rgb=border, line_width=Pt(1.25))
        try:
            shp.adjustments[0] = 0.5
        except Exception:
            pass
    elif card == "legend":
        shp = _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, H, _TEXT_WHITE,
                         line_rgb=accent, line_width=Pt(1.0))
        try:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        except Exception:
            pass
    elif card == "plaque":
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, H, _TEXT_WHITE,
                   line_rgb=border, line_width=Pt(1.0))
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, Pt(4), accent)
    elif card == "window":
        _add_shape(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, L, T, W, H,
                   _TEXT_WHITE, line_rgb=border, line_width=Pt(1.25))
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T + Inches(0.26), W, Pt(1), border)
    elif card == "quote":
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, H, _TEXT_WHITE)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, Pt(5), H, support)
    elif card == "framed":
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, H, _TEXT_WHITE,
                   line_rgb=border, line_width=Pt(1.5))
        inset = Inches(0.055)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L + inset, T + inset,
                   W - inset * 2, H - inset * 2, _TEXT_WHITE,
                   line_rgb=support, line_width=Pt(0.75))
    elif card == "bubble":
        shp = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, L, T, W, H,
                         _TEXT_WHITE, line_rgb=border, line_width=Pt(1.25))
        try:
            shp.adjustments[0] = 0.22
        except Exception:
            pass
        _add_shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, L + Inches(0.45),
                   T + H - Inches(0.02), Inches(0.26), Inches(0.2), _TEXT_WHITE,
                   line_rgb=border, line_width=Pt(1.0))
    elif card == "ornate":
        _add_shape(slide, MSO_SHAPE.RECTANGLE, L, T, W, H, _TEXT_WHITE,
                   line_rgb=border, line_width=Pt(1.25))
        d = Inches(0.11)
        for cx, cy in ((L, T), (L + W - d, T), (L, T + H - d), (L + W - d, T + H - d)):
            _add_shape(slide, MSO_SHAPE.DIAMOND, cx, cy, d, d, support)
    else:
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, L, T, W, H, _TEXT_WHITE,
                   line_rgb=border, line_width=Pt(1.25))


def _pptx_bullet_grid(slide, top_in: float, bullets: list[str], row_heights: list[float], accent: RGBColor, accent_soft: RGBColor, ink: RGBColor, pastel: bool = False, card: str = "rounded", marker: str = "number", kind: str = "") -> None:
    left_in = _CONTENT_X_IN
    total_w_in = _CONTENT_W_IN
    col_gap_in = 0.3
    cols = _grid_columns(len(bullets), kind)
    col_w_in = (total_w_in - col_gap_in * (cols - 1)) / cols
    pad_in = 0.22
    badge_d_in = _GRID_BADGE_D_IN
    protrude_in = _GRID_BADGE_PROTRUDE_IN
    font_size = 14
    row_gap_in = 0.35
    text_w_in = col_w_in - pad_in * 2
    _hex = str(accent)
    _r, _g, _b = int(_hex[0:2], 16), int(_hex[2:4], 16), int(_hex[4:6], 16)
    border_color = RGBColor(int(_r + (255 - _r) * 0.55), int(_g + (255 - _g) * 0.55), int(_b + (255 - _b) * 0.55))

    y_in = top_in
    for r, row_h_in in enumerate(row_heights):
        for c in range(cols):
            j = r * cols + c
            if j >= len(bullets):
                break
            bp = bullets[j]
            x_in = left_in + c * (col_w_in + col_gap_in)
            this_w_in = col_w_in
            if c == 0 and j == len(bullets) - 1 and len(bullets) % cols == 1:
                this_w_in = total_w_in
            text_w_in = this_w_in - pad_in * 2
            if pastel:
                card_fill = _PLAYFUL_PALETTE[j % len(_PLAYFUL_PALETTE)]
                card_badge_color = _PLAYFUL_PALETTE_DARK[j % len(_PLAYFUL_PALETTE_DARK)]
                card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_in), Inches(y_in), Inches(this_w_in), Inches(row_h_in),
                                  card_fill, line_rgb=card_badge_color, line_width=Pt(1.25))
                card.rotation = 1.4 if j % 2 == 0 else -1.4
                card.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            else:
                card_badge_color = accent
                _draw_bullet_card(slide, card, x_in, y_in, this_w_in, row_h_in,
                                  accent, border_color, accent_soft)

            badge_top_in = y_in - protrude_in
            icon = _pick_content_icon(bp)
            if pastel and icon is None:
                icon = _PLAYFUL_ICON_SHAPES[j % len(_PLAYFUL_ICON_SHAPES)]
            if icon is None:
                icon = _MARKER_GLYPH_SHAPE.get(marker)
            badge_shape = _MARKER_BADGE_SHAPE.get(marker, MSO_SHAPE.OVAL)
            badge_cx = Inches(x_in + pad_in + badge_d_in / 2)
            badge_cy = Inches(badge_top_in + badge_d_in / 2)
            if icon is not None:
                _add_icon_badge(slide, badge_cx, badge_cy, Inches(badge_d_in),
                                icon, card_badge_color, _TEXT_WHITE,
                                badge_shape=badge_shape)
            else:
                _add_number_badge(slide, badge_cx, badge_cy, Inches(badge_d_in),
                                  j + 1, card_badge_color, _TEXT_WHITE,
                                  badge_shape=badge_shape,
                                  label=_marker_label(marker, j + 1))

            text_y_in = badge_top_in + badge_d_in + 0.16
            _add_text(slide, Inches(x_in + pad_in), Inches(text_y_in), Inches(text_w_in), Inches(row_h_in - (text_y_in - y_in)),
                      bp, font_size, False, ink)
        y_in += row_h_in + row_gap_in


def _estimate_pptx_visual_height(block: dict) -> float:
    btype = block.get("type")
    data = block.get("data") or {}
    if btype in ("table", "comparison"):
        if btype == "table":
            headers = data.get("headers") or []
            rows = data.get("rows") or []
        else:
            headers = [""] + list(data.get("criteria") or [])
            rows = data.get("items") or []
        if not headers or not rows:
            return 0.0
        return 0.42 * (1 + min(5, len(rows)))
    elif btype == "process":
        steps = (data.get("steps") or [])[:6]
        if not steps:
            return 0.0
        return _pptx_process_layout(steps, 10.1)[1]
    elif btype == "chart":
        return 3.0 if data.get("categories") and data.get("series") else 0.0
    elif btype == "figure":
        image_path = block.get("image")
        if not image_path:
            return 0.0
        full_path = os.path.join(os.path.dirname(__file__), "..", str(image_path).lstrip("/"))
        if not os.path.exists(full_path):
            return 0.0
        try:
            with PILImage.open(full_path) as im:
                iw, ih = im.size
            scale = min(10.1 / iw, 3.6 / ih)
            h_in = ih * scale
        except Exception:
            return 0.0
        if str(block.get("caption") or ""):
            h_in += 0.45
        return h_in
    return 0.0


_PROCESS_TITLE_PT = 15
_PROCESS_DESC_PT = 12
_PROCESS_ITEM_GAP_IN = 0.20
_PROCESS_NUM_COL_IN = 0.42


def _pptx_process_layout(steps: list, width_in: float) -> tuple[list[tuple[float, float]], float]:
    text_w_in = width_in - _PROCESS_NUM_COL_IN
    rows: list[tuple[float, float]] = []
    total = 0.0
    for step in steps:
        title = str((step or {}).get("title", "")).strip()
        desc = str((step or {}).get("description", "")).strip()
        title_h = (_estimate_pptx_lines(title, text_w_in, _PROCESS_TITLE_PT)
                   * _PROCESS_TITLE_PT * 1.30 / 72.0) if title else 0.0
        desc_h = (_estimate_pptx_lines(desc, text_w_in, _PROCESS_DESC_PT)
                  * _PROCESS_DESC_PT * 1.34 / 72.0 + 0.04) if desc else 0.0
        rows.append((title_h, desc_h))
        total += title_h + desc_h
    total += _PROCESS_ITEM_GAP_IN * max(0, len(rows) - 1)
    return rows, total


def _chart_color_shades(accent: RGBColor, n: int) -> list[RGBColor]:
    hex_str = str(accent)
    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
    if n <= 1:
        return [accent]
    shades = []
    for i in range(n):
        t = i / (n - 1)
        factor = 0.62 + t * 0.7
        shades.append(RGBColor(min(255, int(r * factor)), min(255, int(g * factor)), min(255, int(b * factor))))
    return shades


def _pptx_visual_block(slide, top_in: float, width_in: float, block: dict, accent: RGBColor, accent_soft: RGBColor, ink: RGBColor, pastel: bool = False) -> float:
    btype = block.get("type")
    data = block.get("data") or {}
    left = Inches(0.9)

    try:
        if btype in ("table", "comparison"):
            if btype == "table":
                headers = data.get("headers") or []
                rows = data.get("rows") or []
            else:
                criteria = data.get("criteria") or []
                items = data.get("items") or []
                headers = [""] + list(criteria)
                rows = [[it.get("name", "")] + list(it.get("values") or []) for it in items]
            if not headers or not rows:
                return 0.0
            rows = rows[:5]
            n_cols = len(headers)
            col_w_in = width_in / n_cols
            row_h_in = 0.42
            y = top_in
            for c, h in enumerate(headers):
                if pastel:
                    head_fill = _PLAYFUL_PALETTE[c % len(_PLAYFUL_PALETTE)]
                    head_text_color = _PLAYFUL_PALETTE_DARK[c % len(_PLAYFUL_PALETTE_DARK)]
                else:
                    head_fill, head_text_color = accent, _TEXT_WHITE
                _add_shape(slide, MSO_SHAPE.RECTANGLE, left + Inches(c * col_w_in), Inches(y), Inches(col_w_in), Inches(row_h_in), head_fill)
                _add_text(slide, left + Inches(c * col_w_in) + Inches(0.05), Inches(y + 0.06), Inches(col_w_in - 0.1), Inches(row_h_in - 0.1),
                          str(h), 11, True, head_text_color, PP_ALIGN.CENTER)
            y += row_h_in
            for r_i, row in enumerate(rows):
                if r_i % 2 == 0:
                    stripe = _PLAYFUL_PALETTE[(r_i + 1) % len(_PLAYFUL_PALETTE)] if pastel else accent_soft
                    _add_shape(slide, MSO_SHAPE.RECTANGLE, left, Inches(y), Inches(width_in), Inches(row_h_in), stripe)
                for c, cell in enumerate(list(row)[:n_cols]):
                    _add_text(slide, left + Inches(c * col_w_in) + Inches(0.05), Inches(y + 0.06), Inches(col_w_in - 0.1), Inches(row_h_in - 0.1),
                              str(cell), 10, c == 0 and btype == "comparison", ink, PP_ALIGN.CENTER)
                y += row_h_in
            return y - top_in

        elif btype == "process":
            steps = (data.get("steps") or [])[:6]
            if not steps:
                return 0.0
            text_x = left + Inches(_PROCESS_NUM_COL_IN)
            text_w_in = width_in - _PROCESS_NUM_COL_IN
            rows, total_h = _pptx_process_layout(steps, width_in)
            y = top_in
            for i, (step, (title_h, desc_h)) in enumerate(zip(steps, rows)):
                title = str((step or {}).get("title", "")).strip()
                desc = str((step or {}).get("description", "")).strip()
                if title:
                    _add_text(slide, left, Inches(y), Inches(_PROCESS_NUM_COL_IN),
                              Inches(title_h), f"{i + 1}.", _PROCESS_TITLE_PT, True, accent)
                    _add_text(slide, text_x, Inches(y), Inches(text_w_in),
                              Inches(title_h), title, _PROCESS_TITLE_PT, True, ink)
                    y += title_h
                if desc:
                    _add_text(slide, text_x, Inches(y), Inches(text_w_in),
                              Inches(desc_h), desc, _PROCESS_DESC_PT, False, _MUTED)
                    y += desc_h
                if i < len(steps) - 1:
                    y += _PROCESS_ITEM_GAP_IN
            return total_h

        elif btype == "chart":
            categories = [str(c) for c in (data.get("categories") or [])][:8]
            series_in = (data.get("series") or [])[:4]
            if not categories or not series_in:
                return 0.0
            chart_data = CategoryChartData()
            chart_data.categories = categories
            for s in series_in:
                raw_vals = list(s.get("values") or [])[:len(categories)]
                vals = []
                for v in raw_vals:
                    try:
                        vals.append(float(v))
                    except (TypeError, ValueError):
                        vals.append(0)
                vals += [0] * (len(categories) - len(vals))
                chart_data.add_series(str(s.get("name", "")), vals)

            chart_kind = str(data.get("chart_type") or "bar").lower()
            xl_type = {
                "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "pie": XL_CHART_TYPE.PIE,
                "line": XL_CHART_TYPE.LINE_MARKERS,
            }.get(chart_kind, XL_CHART_TYPE.COLUMN_CLUSTERED)

            chart_h_in = 3.0
            graphic_frame = slide.shapes.add_chart(xl_type, left, Inches(top_in), Inches(width_in), Inches(chart_h_in), chart_data)
            chart = graphic_frame.chart
            chart.has_title = False

            is_pie = chart_kind == "pie"
            chart.has_legend = is_pie or len(series_in) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                chart.legend.font.size = Pt(10)

            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(9)
            plot.data_labels.font.color.rgb = _TEXT_WHITE if is_pie else ink

            if is_pie:
                wedge_colors = _chart_color_shades(accent, len(categories))
                for i, point in enumerate(plot.series[0].points):
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = wedge_colors[i]
            else:
                series_colors = _chart_color_shades(accent, len(series_in))
                for i, s in enumerate(plot.series):
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = series_colors[i]
                    if chart_kind == "line":
                        s.format.line.color.rgb = series_colors[i]
                        s.format.line.width = Pt(2.25)
                chart.category_axis.tick_labels.font.size = Pt(10)
                chart.value_axis.tick_labels.font.size = Pt(10)
                chart.value_axis.has_major_gridlines = False

            return chart_h_in

        elif btype == "figure":
            image_path = block.get("image")
            if not image_path:
                return 0.0
            full_path = os.path.join(os.path.dirname(__file__), "..", image_path.lstrip("/"))
            if not os.path.exists(full_path):
                return 0.0
            max_w_in, max_h_in = width_in, 3.6
            try:
                with PILImage.open(full_path) as im:
                    iw, ih = im.size
                scale = min(max_w_in / iw, max_h_in / ih)
                w_in, h_in = iw * scale, ih * scale
            except Exception:
                return 0.0
            slide.shapes.add_picture(full_path, left + Inches((width_in - w_in) / 2), Inches(top_in), Inches(w_in), Inches(h_in))
            used_h = h_in
            caption = str(block.get("caption") or "")
            if caption:
                _add_text(slide, left, Inches(top_in + h_in + 0.08), Inches(width_in), Inches(0.35),
                          caption, 11, False, _MUTED, PP_ALIGN.CENTER)
                used_h += 0.45
            return used_h
    except Exception:
        return 0.0
    return 0.0



_PPTX_BASELINE = {"sup": "30000", "sub": "-25000"}

_A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _pptx_math_element(latex: str, template_run, fallback_pieces):
    import copy
    from lxml import etree
    from pptx.oxml.ns import qn as _pqn
    from app.math_render import to_omml

    try:
        omml = to_omml(latex)
        if not omml:
            return None
        math_el = etree.fromstring(omml)
    except Exception:
        return None

    try:
        alt = etree.SubElement(etree.Element("root"), f"{{{_MC_NS}}}AlternateContent",
                               nsmap={"mc": _MC_NS})
        choice = etree.SubElement(alt, f"{{{_MC_NS}}}Choice", nsmap={"a14": _A14_NS})
        choice.set("Requires", "a14")
        holder = etree.SubElement(choice, f"{{{_A14_NS}}}m")
        holder.append(math_el)

        fallback = etree.SubElement(alt, f"{{{_MC_NS}}}Fallback")
        for piece_text, kind in (fallback_pieces or [(latex, "")]):
            if not piece_text:
                continue
            run = copy.deepcopy(template_run)
            text_el = run.find(_pqn('a:t'))
            if text_el is None:
                continue
            text_el.text = piece_text
            if kind:
                rPr = run.find(_pqn('a:rPr'))
                if rPr is None:
                    rPr = run.makeelement(_pqn('a:rPr'), {})
                    run.insert(0, rPr)
                rPr.set('baseline', _PPTX_BASELINE[kind])
            fallback.append(run)
        if len(fallback) == 0:
            return None
        return alt
    except Exception:
        return None


def _pptx_math_runs(paragraph) -> None:
    import copy
    from pptx.oxml.ns import qn as _pqn
    from app.math_render import script_segments, is_literal_operator

    for run in list(paragraph.runs):
        text = run.text or ""
        normalized = _normalize_math(text)
        if "$" not in normalized:
            continue
        pieces: list[tuple] = []
        pos = 0
        for m in _MATH_SPAN.finditer(normalized):
            prose = normalized[pos:m.start()]
            if prose:
                pieces.append(("text", prose, ""))
            latex = m.group(1)
            if is_literal_operator(latex):
                pieces.append(("text", latex, ""))
                pos = m.end()
                continue
            segments = script_segments(latex)
            if segments is None:
                pieces.append(("equation", latex, [(latex, "")]))
            elif any(kind for _, kind in segments):
                pieces.append(("equation", latex, segments))
            else:
                pieces.extend(("text", t, k) for t, k in segments)
            pos = m.end()
        rest = normalized[pos:]
        if rest:
            pieces.append(("text", rest, ""))
        if not pieces:
            continue

        r = run._r
        parent = r.getparent()
        index = list(parent).index(r)
        offset = 0
        for piece in pieces:
            if piece[0] == "equation":
                node = _pptx_math_element(piece[1], r, piece[2])
                if node is None:
                    node = copy.deepcopy(r)
                    t = node.find(_pqn('a:t'))
                    if t is not None:
                        t.text = piece[1]
            else:
                node = copy.deepcopy(r)
                t = node.find(_pqn('a:t'))
                if t is None:
                    continue
                t.text = piece[1]
                if piece[2]:
                    rPr = node.find(_pqn('a:rPr'))
                    if rPr is None:
                        rPr = node.makeelement(_pqn('a:rPr'), {})
                        node.insert(0, rPr)
                    rPr.set('baseline', _PPTX_BASELINE[piece[2]])
            parent.insert(index + offset, node)
            offset += 1
        parent.remove(r)


def _typeset_math_pptx(prs) -> None:
    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    _pptx_math_runs(paragraph)
            except Exception:
                continue


def build_presentation_pptx(content: dict) -> io.BytesIO:
    prs = Presentation()
    prs.slide_width = _SLIDE_W
    prs.slide_height = _SLIDE_H

    slides_data = content.get("slides", [])
    title_text = content.get("title", "")
    desc = content.get("description", "")

    total_slides = len(slides_data)

    theme = resolve_deck_theme(content)
    _accent_rgb = (_hex_tuple(theme["accent"]) if theme.get("accent")
                   else get_subject_accent_rgb(content.get("subject")))
    _ACCENT, _ACCENT_DARK, _ACCENT_SOFT = _pptx_accent_shades(_accent_rgb)
    _SUPPORT = _hex_rgb(theme["support"]) if theme.get("support") else _ACCENT_SOFT
    _DECOR = str(theme.get("decor") or "none")
    _CARD = str(theme.get("card") or ("rounded" if theme.get("cards") else "none"))
    _MARKER = str(theme.get("marker") or "number")
    _COVER = str(theme.get("cover") or "standard")
    title_font = theme["title_font"]
    _INK = theme["ink"]
    _MUTED = theme["muted"]
    _BG = theme["bg"]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, *_BG)
    if theme.get("spiral"):
        _pptx_notebook_lines(slide)
        _pptx_spiral_margin(slide)

    _cover_w_in = 6.7 if len([1 for sd in slides_data if sd.get("title")]) >= 4 else 10.8
    _cover_pt_full = int(theme.get("cover_title_pt") or 46)
    _cover_title_pt = int(_cover_pt_full * 0.87) if _cover_w_in < 10 else _cover_pt_full
    _cover_line_in = _cover_title_pt * 1.22 / 72.0
    _cover_lines = _estimate_pptx_lines(title_text, _cover_w_in, _cover_title_pt)
    if _cover_lines >= 2:
        _cover_lines += 1
    _cover_title_h = min(3.3, max(0.95, _cover_lines * _cover_line_in + 0.15))
    _rule_y_in = max(4.35, 2.4 + _cover_title_h + 0.18)

    slide_decor.draw(slide, _DECOR, _SUPPORT, _BG, 0, on_cover=True)
    if not (desc and _rule_y_in + 0.25 > 5.05):
        slide_decor.draw_cover(slide, _COVER, _ACCENT, _SUPPORT, _BG)


    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(0.85), Inches(0.5), Pt(2.5), _ACCENT)
    _add_text(slide, Inches(1.6), Inches(0.72), Inches(6), Inches(0.35),
              "ПРЕЗЕНТАЦИЯ", 12, True, _ACCENT)

    badge_w, badge_h = Inches(1.7), Inches(1.05)
    badge_x, badge_y = _SLIDE_W - badge_w - Inches(0.8), Inches(0.6)
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, badge_x, badge_y, badge_w, badge_h, _ACCENT)
    _add_text(slide, badge_x, badge_y + Inches(0.16), badge_w, Inches(0.35), f"{total_slides + 1} СЛАЙДОВ", 11, True, _TEXT_WHITE, PP_ALIGN.CENTER)
    _add_text(slide, badge_x, badge_y + Inches(0.55), badge_w, Inches(0.4), "Dastyor", 15, True, _TEXT_WHITE, PP_ALIGN.CENTER)

    _illustration_path = (None if theme.get("subject_template")
                          else get_subject_illustration_path(content.get("subject")))
    if _illustration_path:
        try:
            with PILImage.open(_illustration_path) as _im:
                _iw, _ih = _im.size
            _icon_box_in = 0.65
            _scale = min(_icon_box_in / _iw, _icon_box_in / _ih)
            _icon_w_in, _icon_h_in = _iw * _scale, _ih * _scale
            _icon_x = badge_x + (badge_w - Inches(_icon_w_in)) / 2
            _icon_y = badge_y + badge_h + Inches(0.1)
            slide.shapes.add_picture(_illustration_path, _icon_x, _icon_y, Inches(_icon_w_in), Inches(_icon_h_in))
        except Exception:
            pass

    _add_text(slide, Inches(0.95), Inches(2.4), Inches(_cover_w_in), Inches(_cover_title_h),
              title_text, _cover_title_pt, True, _INK, font_name=title_font)

    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(_rule_y_in),
               Inches(0.9), Pt(3), _ACCENT)

    if desc:
        _add_text(slide, Inches(1.0), Inches(_rule_y_in + 0.25), Inches(_cover_w_in),
                  Inches(1.4), desc, 15, False, _MUTED)

    agenda = [str(sd.get("title") or "").strip()
              for sd in slides_data if str(sd.get("title") or "").strip()]
    if len(agenda) >= 4:
        plan_x_in, plan_top_in = 8.05, 2.45
        _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(plan_x_in), Inches(plan_top_in),
                   Pt(2.5), Inches(3.5), _ACCENT_SOFT)
        _add_text(slide, Inches(plan_x_in + 0.3), Inches(plan_top_in), Inches(4.0),
                  Inches(0.3), _DECK_AGENDA_LABEL.get(
                      str(content.get("language") or "Русский"), "В ЭТОМ УРОКЕ"),
                  10, True, _ACCENT)
        y_in = plan_top_in + 0.42
        for n, item in enumerate(agenda[:5], 1):
            lines = _estimate_pptx_lines(item, 3.5, 12)
            _add_text(slide, Inches(plan_x_in + 0.3), Inches(y_in), Inches(0.35),
                      Inches(0.3), f"{n}", 12, True, _ACCENT)
            _add_text(slide, Inches(plan_x_in + 0.72), Inches(y_in), Inches(3.6),
                      Inches(lines * 0.24), item, 12, False, _INK)
            y_in += max(0.34, lines * 0.24 + 0.12)
        if len(agenda) > 5:
            _lang = str(content.get("language") or "Русский")
            _more_n = len(agenda) - 5
            more = _DECK_AGENDA_MORE.get(_lang, "ещё {n} {word}")
            _add_text(slide, Inches(plan_x_in + 0.72), Inches(y_in), Inches(3.6),
                      Inches(0.3),
                      more.format(n=_more_n, word=_agenda_more_word(_lang, _more_n)).strip(),
                      11, False, _MUTED)

    if theme.get("pastel"):
        _pptx_sticker_decorations(slide)

    _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.85), Inches(11.5), Pt(0.75), _HAIRLINE)
    _add_text(slide, Inches(1.0), Inches(7.0), Inches(9.3), Inches(0.4),
              f"{content.get('subject', '')}  ·  {content.get('grade', '')}".strip(" ·"),
              11, False, _MUTED)

    _add_logo_badge(slide, on_dark_bg=False)
    _add_transition(slide)

    def _name_slide(slide, layout_name):
        try:
            slide._element.cSld.set("name", str(layout_name))
        except Exception:
            pass

    def _finish_slide(slide, i, notes):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

        footer_w_in = _CONTENT_W_IN
        _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(7.1), Inches(footer_w_in), Pt(2.25), _HAIRLINE)
        progress_w_in = footer_w_in * (i + 1) / total_slides
        if progress_w_in > 0.05:
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(7.1), Inches(progress_w_in), Pt(2.25), _ACCENT)

        _add_logo_badge(slide, on_dark_bg=False)
        _add_transition(slide)

    _grade_tier = slide_layouts.grade_tier(content.get("grade"))
    _char_prop = slide_characters.prop_for(str(theme.get("subject_template") or ""))
    _PROP_COLOR = (_hex_rgb(theme["prop_color"]) if theme.get("prop_color")
                   else _ACCENT_DARK)
    _last_layout: str | None = None

    for i, sd in enumerate(slides_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_slide_bg(slide, *_BG)
        if theme.get("spiral"):
            _pptx_notebook_lines(slide)
            _pptx_spiral_margin(slide)
            _pptx_washi_tape(slide, 0.75, 0.1, _PLAYFUL_PALETTE[i % len(_PLAYFUL_PALETTE)])
        slide_decor.draw(slide, _DECOR, _SUPPORT, _BG, i)

        title_str = sd.get("title", f"Слайд {i + 1}")

        band_h_in = 1.15
        header_style = theme["header"]
        title_ink = _TEXT_WHITE if header_style == "band" else _INK
        counter_ink = _TEXT_WHITE if header_style == "band" else _MUTED

        if header_style == "band":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, _SLIDE_W, Inches(1.5), _ACCENT)
        elif header_style == "hexband":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, _SLIDE_W, Inches(0.1), _ACCENT)
            _add_shape(slide, MSO_SHAPE.HEXAGON, Inches(0.62), Inches(0.37),
                       Inches(0.2), Inches(0.2), _ACCENT)
        elif header_style == "index":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.66), Inches(0.33),
                       Pt(3), Inches(0.95), _ACCENT)
        elif header_style == "lozenge":
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(0.33),
                       Inches(0.62), Inches(0.3), _ACCENT_SOFT)
        elif header_style == "initial":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.58),
                       Pt(5), Inches(0.78), _ACCENT)

        _number_x = 0.9 if header_style not in ("prompt",) else 1.12
        if header_style == "prompt":
            _add_text(slide, Inches(0.86), Inches(0.33), Inches(0.3), Inches(0.3),
                      ">", 14, True, _ACCENT, font_name="Consolas")
        _add_text(slide, Inches(_number_x), Inches(0.35), Inches(8.6), Inches(0.3),
                  f"{i + 1:02d}", 13, True,
                  _TEXT_WHITE if header_style == "band" else _ACCENT)
        _kind_text = _kind_label(str(sd.get("kind") or "").strip().lower(),
                                 content.get("language"))
        if _kind_text:
            _add_text(slide, Inches(_number_x + 0.45), Inches(0.37), Inches(4.2), Inches(0.3),
                      f"· {_kind_text}", 10, True,
                      _TEXT_WHITE if header_style == "band" else _MUTED)
        _add_text(slide, Inches(0.9), Inches(0.62), Inches(9.6), Inches(0.7),
                  title_str.upper() if theme["title_caps"] else title_str,
                  int(theme.get("title_pt") or (24 if theme["title_caps"] else 27)),
                  True, title_ink, font_name=title_font)

        if header_style == "rule":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.32), Inches(1.1), Pt(3), _ACCENT)
        elif header_style == "underline":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.36), Inches(11.5), Pt(1.25), _HAIRLINE)
        elif header_style == "index":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.34), Inches(2.2), Pt(1.5), _ACCENT_SOFT)
        elif header_style == "measure":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.36), Inches(11.5), Pt(1.25), _HAIRLINE)
            for _t in range(12):
                _tall = _t % 4 == 0
                _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9 + _t * (11.5 / 12)),
                           Inches(1.36 - (0.09 if _tall else 0.05)), Pt(1.25),
                           Inches(0.09 if _tall else 0.05),
                           _ACCENT if _tall else _HAIRLINE)
        elif header_style == "hexband":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.34), Inches(1.6), Pt(2.5), _ACCENT)
        elif header_style == "lozenge":
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.3),
                       Inches(1.5), Pt(5), _ACCENT)
        elif header_style == "masthead":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.3), Inches(11.5), Pt(2.5), _ACCENT)
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.4), Inches(11.5), Pt(0.75), _ACCENT_SOFT)
        elif header_style == "prompt":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.36), Inches(11.5), Pt(1.25), _HAIRLINE)
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.33), Inches(0.75), Pt(2.5), _ACCENT)
        elif header_style == "ornamental":
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.37), Inches(4.6), Pt(1), _ACCENT_SOFT)
            for _d in range(3):
                _add_shape(slide, MSO_SHAPE.DIAMOND, Inches(4.75 + _d * 0.22), Inches(1.31),
                           Inches(0.11), Inches(0.11), _ACCENT if _d == 1 else _ACCENT_SOFT)
        _add_text(slide, Inches(9.9), Inches(0.35), Inches(1.7), Inches(0.3),
                  f"{i + 1} / {total_slides}", 10, False, counter_ink, PP_ALIGN.RIGHT)
        if theme["badge"] and header_style != "band":
            badge_d = Inches(0.6)
            badge_cx, badge_cy = Inches(12.15), Inches(0.65)
            _add_shape(slide, MSO_SHAPE.OVAL, badge_cx - badge_d / 2, badge_cy - badge_d / 2, badge_d, badge_d, _ACCENT_SOFT)
            icon_d = badge_d * 0.42
            _marker_glyph = _MARKER_GLYPH_SHAPE.get(_MARKER)
            if theme.get("pastel"):
                _corner_shape = _PLAYFUL_ICON_SHAPES[i % len(_PLAYFUL_ICON_SHAPES)]
            elif _marker_glyph is not None:
                _corner_shape = _marker_glyph
            else:
                _corner_shape = _ICON_SHAPES[i % len(_ICON_SHAPES)]
            _add_shape(slide, _corner_shape, badge_cx - icon_d / 2, badge_cy - icon_d / 2, icon_d, icon_d, _ACCENT)

        bullets = sd.get("bullet_points", [])
        notes = sd.get("speaker_notes", "")
        visual = sd.get("visual")

        slide_kind = str(sd.get("kind") or "").strip().lower()
        slide_image = sd.get("image") if isinstance(sd.get("image"), dict) else None
        image_path = None
        if slide_image and slide_image.get("path"):
            candidate = os.path.join(os.path.dirname(__file__), "..",
                                     str(slide_image["path"]).lstrip("/"))
            if os.path.exists(candidate):
                image_path = candidate
        _layout_used = None
        if (theme.get("subject_template") and not visual
                and (bullets or str(sd.get("body") or "").strip())
                and str(sd.get("kind") or "").strip().lower() != "formula"):
            _ctx = slide_layouts.Ctx(
                slide=slide, accent=_ACCENT, accent_dark=_ACCENT_DARK,
                accent_soft=_ACCENT_SOFT, support=_SUPPORT, ink=_INK,
                muted=_MUTED, bg=_BG, title_font=title_font,
                body_font=theme["body_font"],
                template_id=str(theme.get("subject_template") or ""),
                grade_tier=_grade_tier,
                language=str(content.get("language") or "Русский"),
                index=i, total=total_slides, image_path=image_path,
                image_credit=str((slide_image or {}).get("credit") or ""),
                character_prop=_char_prop, prop_color=_PROP_COLOR,
                topic=str(content.get("title") or ""),
            )
            _layout_used = slide_layouts.draw(_ctx, sd, _last_layout)
        if _layout_used:
            _last_layout = _layout_used
            _name_slide(slide, _layout_used)
            _finish_slide(slide, i, notes)
            continue

        is_formula_slide = not image_path and (
            slide_kind == "formula" or (
                bool(bullets) and not visual
                and sum(1 for b in bullets if "$" in str(b)) >= max(1, len(bullets) - 1)))
        image_on_right = (i % 2 == 0)

        bullet_width_in = 5.35 if image_path else _CONTENT_W_IN
        bullet_width = Inches(bullet_width_in)
        bullet_font = 20

        use_grid = (theme["cards"] and len(bullets) >= 1
                    and not image_path and not is_formula_slide)
        available_top_in = band_h_in + 0.5
        available_h_in = 7.0 - available_top_in

        if use_grid:
            grid_row_gap_in = 0.35
            natural_row_heights = _pptx_bullet_grid_natural_heights(bullets, (bullet_width_in - 0.3) / 2, kind=slide_kind)
            n_rows = len(natural_row_heights)
            row_heights = [min(1.7, max(h, 1.15)) for h in natural_row_heights]
            grid_h = sum(row_heights) + grid_row_gap_in * max(0, n_rows - 1)
        else:
            line_h_in = bullet_font * 1.32 / 72.0
            bullet_gap_in = 0.22
            chip = Inches(0.42)
            row_heights = [max(0.55, _estimate_pptx_lines(bp, bullet_width_in, bullet_font) * line_h_in + 0.24) for bp in bullets]
            grid_h = sum(row_heights) + bullet_gap_in * max(0, len(bullets) - 1)

        body_text = re.sub(r"\s+", " ", str(sd.get("body") or "")).strip()
        body_font = 15
        body_width_in = bullet_width_in
        body_x_override = None
        body_h_in = 0.0
        if body_text:
            body_h_in = (_estimate_pptx_lines(body_text, body_width_in - 0.2, body_font)
                         * body_font * 1.38 / 72.0) + 0.30

        content_h = grid_h + (body_h_in + 0.28 if body_h_in else 0.0)
        visual_gap_in = 0.35 if bullets and visual else 0.0
        if visual:
            content_h += visual_gap_in + _estimate_pptx_visual_height(visual)

        max_bottom_in = 6.9
        overflow_in = (available_top_in + content_h) - max_bottom_in
        if overflow_in > 0 and body_h_in:
            content_h -= min(overflow_in, body_h_in + 0.28)

        _offset_cap = 0.85 if bullets else 1.8
        offset_in = min(_offset_cap, max(0.0, (available_h_in - content_h) / 2))
        cursor = available_top_in + offset_in

        if (slide_kind == "intro" and body_text and bullets
                and not image_path and not is_formula_slide and not visual):
            lead_pt = 17
            lead_h_in = (_estimate_pptx_lines(body_text, _CONTENT_W_IN - 0.4, lead_pt)
                         * lead_pt * 1.42 / 72.0) + 0.12
            if cursor + lead_h_in + 0.34 + grid_h <= 6.9:
                _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(_CONTENT_X_IN),
                           Inches(cursor + 0.04), Pt(4), Inches(max(0.3, lead_h_in - 0.1)),
                           _ACCENT)
                _add_text(slide, Inches(_CONTENT_X_IN + 0.24), Inches(cursor),
                          Inches(_CONTENT_W_IN - 0.4), Inches(lead_h_in),
                          body_text, lead_pt, False, _INK,
                          font_name=theme["body_font"])
                cursor += lead_h_in + 0.34
                body_text, body_h_in = "", 0.0

        if is_formula_slide and bullets:
            formula_text = next((str(b) for b in bullets if "$" in str(b)), str(bullets[0]))
            rest = [str(b) for b in bullets if str(b) != formula_text]
            from app.math_render import measure as _measure
            latex_src = formula_text.strip().strip("$").strip()
            avail_pt = 9.5 * 72
            f_size = 32
            width_pt = _measure(latex_src, f_size)[0]
            if width_pt > avail_pt:
                f_size = max(18, int(f_size * avail_pt / width_pt))
                width_pt = _measure(latex_src, f_size)[0]
            f_lines = max(1, math.ceil(width_pt / avail_pt)) if width_pt else 1
            text_h_in = f_lines * (f_size * 1.5 / 72)
            panel_h_in = max(1.6, text_h_in + 0.7)
            panel_top = Inches(available_top_in + 0.35)
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), panel_top,
                       Inches(10.1), Inches(panel_h_in), _ACCENT_SOFT)
            _add_text(slide, Inches(1.9),
                      panel_top + Inches((panel_h_in - text_h_in) / 2),
                      Inches(9.5), Inches(text_h_in),
                      formula_text, f_size, True, _ACCENT_DARK, PP_ALIGN.CENTER)
            below = available_top_in + 0.35 + panel_h_in + 0.45
            for line in rest[:2]:
                _add_text(slide, Inches(1.9), Inches(below), Inches(9.5), Inches(0.5),
                          line, 18, False, _MUTED, PP_ALIGN.CENTER)
                below += 0.55
            cursor = below
            body_width_in = 9.5
            body_x_override = 1.9

        elif image_path:
            img_w_in, img_area_h_in = 5.9, 4.4
            text_x_in = 0.9 if image_on_right else 7.05
            img_x_in = 6.75 if image_on_right else 0.75
            try:
                with PILImage.open(image_path) as im:
                    iw, ih = im.size
                scale = min(img_w_in / iw, img_area_h_in / ih)
                tw, th = iw * scale, ih * scale
            except Exception:
                tw, th = img_w_in, img_area_h_in * 0.7
            img_top_in = available_top_in + max(0.0, (img_area_h_in - th) / 2)
            slide.shapes.add_picture(image_path, Inches(img_x_in + (img_w_in - tw) / 2),
                                     Inches(img_top_in), Inches(tw), Inches(th))
            credit = str(slide_image.get("credit") or "")
            if credit:
                _add_text(slide, Inches(img_x_in), Inches(img_top_in + th + 0.12),
                          Inches(img_w_in), Inches(0.3), credit, 9, False, _MUTED, PP_ALIGN.CENTER)
            bullet_pt, bullet_gap_in = 19, 0.18
            by = available_top_in + 0.15
            for bp in bullets[:4]:
                lines = _estimate_pptx_lines(str(bp), bullet_width_in - 0.3, bullet_pt)
                row_h_in = lines * bullet_pt * 1.32 / 72.0 + 0.14
                _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(text_x_in), Inches(by + 0.08),
                           Pt(3), Inches(max(0.22, row_h_in - 0.16)), _ACCENT)
                _add_text(slide, Inches(text_x_in + 0.22), Inches(by), Inches(bullet_width_in - 0.3),
                          Inches(row_h_in), str(bp), bullet_pt, False, _INK)
                by += row_h_in + bullet_gap_in
            cursor = by - bullet_gap_in

        elif use_grid:
            _panel_style = _KIND_PANEL.get(slide_kind)
            if _panel_style:
                _draw_kind_panel(slide, _panel_style, cursor, grid_h, _ACCENT, _ACCENT_SOFT)
            _pptx_bullet_grid(slide, cursor, bullets, row_heights, _ACCENT, _ACCENT_SOFT, _INK, pastel=bool(theme.get("pastel")), card=_CARD, marker=_MARKER, kind=slide_kind)
            cursor += grid_h
        else:
            for j, bp in enumerate(bullets):
                by = cursor
                row_h = row_heights[j]

                _marker_text = _marker_label(_MARKER, j + 1)
                _add_shape(slide, _MARKER_BADGE_SHAPE.get(_MARKER, MSO_SHAPE.OVAL),
                           Inches(0.9), Inches(by), chip, chip, _TEXT_WHITE,
                           line_rgb=_ACCENT, line_width=Pt(1.5))
                _add_text(slide, Inches(0.9), Inches(by) + Emu(int(chip * 0.14)), chip, chip,
                          _marker_text, 13 if len(_marker_text) <= 2 else 10, True,
                          _ACCENT, PP_ALIGN.CENTER)

                _add_text(slide, Inches(1.55), Inches(by - 0.07), bullet_width, Inches(row_h),
                          bp, bullet_font, False, _INK)
                cursor += row_h + bullet_gap_in
            if bullets:
                cursor -= bullet_gap_in

        if body_text:
            reserved_in = (visual_gap_in + _estimate_pptx_visual_height(visual)) if visual else 0.0
            room_in = 6.9 - (cursor + 0.28) - reserved_in
            for candidate_pt in (15, 14, 13, 12, 11):
                line_in = candidate_pt * 1.38 / 72.0
                needed = (_estimate_pptx_lines(body_text, body_width_in - 0.2, candidate_pt)
                          * line_in) + 0.30
                if needed <= room_in:
                    body_font, body_h_in = candidate_pt, needed
                    break
            else:
                if visual:
                    body_text, body_h_in = "", 0.0
                else:
                    body_font = 11
                    line_in = body_font * 1.38 / 72.0
                    fit_lines = int((room_in - 0.30) / line_in) if room_in > 0 else 0
                    if fit_lines < 2:
                        body_text, body_h_in = "", 0.0
                    else:
                        body_text = _trim_text_to_lines(body_text, body_width_in - 0.2,
                                                        body_font, fit_lines)
                        body_h_in = fit_lines * line_in + 0.30

        _name_slide(slide, "formula" if is_formula_slide
                    else "picture" if image_path
                    else "visual" if visual
                    else "cards")

        if body_text:
            body_x_in = (body_x_override if body_x_override is not None
                         else (0.9 if (not image_path or image_on_right) else 7.05))
            cursor += 0.28
            _add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(body_x_in), Inches(cursor + 0.06),
                       Pt(2.5), Inches(max(0.3, body_h_in - 0.24)), _ACCENT_SOFT)
            _add_text(slide, Inches(body_x_in + 0.18), Inches(cursor), Inches(body_width_in - 0.2),
                      Inches(body_h_in), body_text, body_font, False, _MUTED,
                      font_name=theme["body_font"])
            cursor += body_h_in

        if visual:
            if bullets:
                cursor += visual_gap_in
            visual_h_in = _estimate_pptx_visual_height(visual)
            cursor = max(available_top_in, min(cursor, 6.95 - visual_h_in))
            _pptx_visual_block(slide, cursor, bullet_width_in, visual, _ACCENT, _ACCENT_SOFT, _INK, pastel=bool(theme.get("pastel")))

        _finish_slide(slide, i, notes)


    buf = io.BytesIO()
    _typeset_math_pptx(prs)
    prs.save(buf)
    buf.seek(0)
    return buf



from reportlab.platypus import HRFlowable, Image as _PdfFlowableImage
from reportlab.graphics.shapes import Drawing, Rect, String, Circle, Image as _PdfImage
from reportlab.graphics import renderPDF


_PDF_ACCENT = '#3B82F6'
_PDF_DARK = '#1E293B'
_PDF_LIGHT = '#F1F5F9'
_PDF_GREEN = '#10B981'
_PDF_ORANGE = '#F59E0B'
_PDF_RED = '#EF4444'
_PDF_PURPLE = '#8B5CF6'

_PDF_BRAND = '#DC2626'
_PDF_BRAND_DARK = '#7F1D1D'
_PDF_BRAND_SOFT = '#FEE2E2'
_PDF_ACCENT_DARK = '#1E3A8A'
_PDF_ACCENT_SOFT = '#DBEAFE'

_LOGO_PATH_PDF = os.path.join(os.path.dirname(__file__), "static", "logo_pdf.png")


def _wrap_pdf_title(text, font_name, font_size, max_width):
    words = text.split()
    if not words:
        return [text]
    lines, current = [], ""
    for w in words:
        trial = f"{current} {w}".strip()
        if pdfmetrics.stringWidth(trial, font_name, font_size) <= max_width:
            current = trial
            continue
        if current:
            lines.append(current)
        if pdfmetrics.stringWidth(w, font_name, font_size) <= max_width:
            current = w
        else:
            chunk = ""
            for ch in w:
                if pdfmetrics.stringWidth(chunk + ch, font_name, font_size) <= max_width:
                    chunk += ch
                else:
                    lines.append(chunk)
                    chunk = ch
            current = chunk
    if current:
        lines.append(current)
    return lines


_REAL_IMAGE_BOX = {
    "logo": 110,
    "cutout": 155,
    "diagram": 210,
}
_REAL_IMAGE_BOX_DEFAULT = 155


_GRADE_WORDS = ("класс", "синф", "grade", "sinf")


def _bare_grade(grade: str) -> str:
    lowered = grade.lower()
    for word in _GRADE_WORDS:
        idx = lowered.find(word)
        if idx != -1:
            return grade[:idx].strip()
    return grade


def _format_grade(grade: str, suffix: str) -> str:
    stripped = grade
    lowered = grade.lower()
    for word in _GRADE_WORDS:
        idx = lowered.find(word)
        if idx != -1:
            stripped = grade[:idx].strip()
            break
    return f"{stripped} {suffix}" if stripped else grade


_SUBJECT_DISPLAY_NAMES: dict[str, dict[str, str]] = {
    "Информатика": {"Таджикский": "Технологияи иттилоотӣ", "Английский": "Computer Science"},
    "Русский язык": {"Таджикский": "Забони русӣ", "Английский": "Russian Language"},
    "Английский язык": {"Таджикский": "Забони англисӣ", "Английский": "English Language"},
    "Таджикский язык": {"Таджикский": "Забони тоҷикӣ", "Английский": "Tajik Language"},
    "Таджикская литература": {"Таджикский": "Адабиёти тоҷик", "Английский": "Tajik Literature"},
    "История Таджикистана": {"Таджикский": "Таърихи Тоҷикистон", "Английский": "History of Tajikistan"},
    "Всемирная история": {"Таджикский": "Таърихи умумиҷаҳонӣ", "Английский": "World History"},
    "Математика": {"Английский": "Mathematics"},
    "Алгебра": {"Английский": "Algebra"},
    "Геометрия": {"Английский": "Geometry"},
    "Физика": {"Английский": "Physics"},
    "Химия": {"Английский": "Chemistry"},
    "Биология": {"Английский": "Biology"},
    "География": {"Английский": "Geography"},
}


def _display_subject(subject: str, language: str) -> str:
    return _SUBJECT_DISPLAY_NAMES.get(subject, {}).get(language, subject)



_LECTURE_INK = "#141821"
_LECTURE_RULE = "#C9CFDA"
_LECTURE_MUTED = "#5A6472"


def _pdf_lecture_masthead(title, subtitle, subject, grade, language, grade_suffix,
                          accent=None):
    accent = accent or get_subject_accent_hex(subject)
    display_subject = _display_subject(str(subject or ""), language)
    meta_bits = [b for b in (display_subject.upper(),
                             _format_grade(grade, grade_suffix).upper() if grade else "") if b]
    elements = [Spacer(1, 10)]
    if meta_bits:
        elements.append(Paragraph(
            "&nbsp;&nbsp;·&nbsp;&nbsp;".join("&nbsp;".join(part) for part in meta_bits),
            ParagraphStyle(name="LecMeta", fontName=FONT_NAME_BOLD, fontSize=8.5,
                           leading=11, alignment=TA_CENTER,
                           textColor=HexColor(accent), spaceAfter=10)))
    elements.append(Paragraph(
        str(title or ""),
        ParagraphStyle(name="LecTitle", fontName=MATH_FONT_BOLD, fontSize=23, leading=28,
                       alignment=TA_CENTER, textColor=HexColor(_LECTURE_INK), spaceAfter=8)))
    rule = Table([[""]], colWidths=[110], rowHeights=[2.2])
    rule.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(accent)),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    elements.append(Table([[rule]], colWidths=[500], style=TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ])))
    if subtitle:
        elements.append(Paragraph(
            f"<i>{subtitle}</i>",
            ParagraphStyle(name="LecSub", fontName=MATH_FONT, fontSize=11, leading=15,
                           alignment=TA_CENTER, textColor=HexColor(_LECTURE_MUTED),
                           spaceAfter=16)))
    return elements


def _pdf_lecture_section_header(label, color, number, font_name=None):
    accent = color or _PDF_ACCENT
    number_bit = (f'<font color="{accent}">{number}.</font>&nbsp;&nbsp;' if number else "")
    head = Paragraph(
        f"{number_bit}{str(label).upper()}",
        ParagraphStyle(name=f"LecHead{abs(hash(label)) % 99999}", fontName=MATH_FONT_BOLD,
                       fontSize=12, leading=15, textColor=HexColor(_LECTURE_INK),
                       spaceBefore=12, spaceAfter=2))
    t = Table([[head]], colWidths=[500])
    t.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, -1), 1.1, HexColor(accent)),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return [t, Spacer(1, 5)]


def _pdf_header(title, subtitle=None, tag=None, accent=_PDF_ACCENT, minimal=False):
    elements = []
    elements.append(Spacer(1, 6))

    if minimal:
        if tag:
            elements.append(Paragraph(
                f'<font color="{accent}" size="8"><b>{"&nbsp;".join(tag.upper())}</b></font>',
                ParagraphStyle(name='MinTag', fontName=FONT_NAME_BOLD, fontSize=8,
                               leading=11, textColor=HexColor(accent), spaceAfter=6),
            ))
        elements.append(Paragraph(
            title,
            ParagraphStyle(name='MinTitle', fontName=FONT_NAME_BOLD, fontSize=21,
                           leading=25, textColor=HexColor(_PDF_DARK), spaceAfter=6),
        ))
        rule = Table([[""]], colWidths=[54], rowHeights=[2.2])
        rule.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), HexColor(accent)),
            ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ]))
        elements.append(rule)
        if subtitle:
            elements.append(Spacer(1, 8))
            elements.append(Paragraph(f'<font color="#64748B">{subtitle}</font>', _styles()['DocSubtitle']))
        elements.append(Spacer(1, 14))
        return elements

    max_title_width = 465
    title_font_size = 18
    while title_font_size > 13 and pdfmetrics.stringWidth(title, FONT_NAME_BOLD, title_font_size) > max_title_width:
        title_font_size -= 1
    title_lines = _wrap_pdf_title(title, FONT_NAME_BOLD, title_font_size, max_title_width)
    line_h = title_font_size + 7
    extra_height = line_h * (len(title_lines) - 1)

    height = (66 if tag else 55) + extra_height
    d = Drawing(500, height)
    d.add(Rect(0, 0, 500, height, fillColor=HexColor(_PDF_DARK), strokeColor=None))
    d.add(Rect(0, height - 3, 500, 3, fillColor=HexColor(accent), strokeColor=None))
    if tag:
        d.add(Rect(10, 10, len(tag) * 7 + 20, 22, fillColor=HexColor(accent), strokeColor=None, rx=11, ry=11))
        d.add(String(20, 16, tag, fontName=FONT_NAME_BOLD, fontSize=9, fillColor=HexColor('#FFFFFF')))
        y_title = 40
    else:
        y_title = 22
    for i, line in enumerate(title_lines):
        y = y_title + line_h * (len(title_lines) - 1 - i)
        d.add(String(15, y, line, fontName=FONT_NAME_BOLD, fontSize=title_font_size, fillColor=HexColor('#FFFFFF')))
    elements.append(d)
    if subtitle:
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(f'<font color="#64748B">{subtitle}</font>', _styles()['DocSubtitle']))
    elements.append(Spacer(1, 6))
    return elements


def _pdf_section_header(label, color=_PDF_ACCENT, number=None):
    elements = []
    elements.append(Spacer(1, 8))
    prefix = f'{number}. ' if number else ''
    elements.append(Paragraph(f'<font color="{color}"><b>{prefix}{label}</b></font>', _styles()['SlideTitle']))
    elements.append(Spacer(1, 3))
    return elements


def _pdf_section_header_underline(label, color, number, font_name):
    style = ParagraphStyle(
        name=f'HdrUL_{abs(hash(label)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=12.5,
        textColor=HexColor(_PDF_DARK), spaceAfter=3,
    )
    return [
        Spacer(1, 5),
        Paragraph(label.upper(), style),
        HRFlowable(width=70, thickness=2.4, color=HexColor(color), spaceAfter=3, hAlign='LEFT'),
    ]


def _pdf_section_header_smallcaps(label, color, number, font_name):
    style = ParagraphStyle(
        name=f'HdrSC_{abs(hash(label)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=8.5,
        textColor=HexColor(color), spaceAfter=2,
    )
    return [
        Spacer(1, 9),
        Paragraph(label.upper(), style),
        HRFlowable(width='100%', thickness=0.6, color=HexColor('#E2E8F0'), spaceAfter=4),
    ]


def _pdf_section_header_serif(label, color, number, font_name):
    style = ParagraphStyle(
        name=f'HdrSerif_{abs(hash(label)) % 100000}', fontName=font_name, fontSize=13,
        textColor=HexColor('#334155'), spaceAfter=3,
    )
    prefix = f'{number}. ' if number else ''
    return [
        Spacer(1, 9),
        Paragraph(f'<b>{prefix}{label}</b>', style),
        HRFlowable(width='100%', thickness=0.75, color=HexColor('#CBD5E1'), spaceAfter=4),
    ]


def _pdf_section_header_bar(label, color, number, font_name):
    style = ParagraphStyle(
        name=f'HdrBar_{abs(hash(label)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=11.5,
        textColor=HexColor('#FFFFFF'),
    )
    prefix = f'{number}.  ' if number else ''
    t = Table([[Paragraph(f'{prefix}{label}', style)]], colWidths=[470])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), HexColor(color)),
        ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 10), ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    return [Spacer(1, 8), t, Spacer(1, 3)]


_PDF_SECTION_RENDERERS = {
    "underline": _pdf_section_header_underline,
    "smallcaps": _pdf_section_header_smallcaps,
    "serif": _pdf_section_header_serif,
    "bar": _pdf_section_header_bar,
}


def _pdf_light_tint_hex(hex_color: str, amount: float = 0.88) -> str:
    h = hex_color.lstrip('#')
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    tr = int(r + (255 - r) * amount)
    tg = int(g + (255 - g) * amount)
    tb = int(b + (255 - b) * amount)
    return f'#{tr:02X}{tg:02X}{tb:02X}'


def _pdf_dark_shade_hex(hex_color: str, amount: float = 0.55) -> str:
    h = hex_color.lstrip('#')
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f'#{int(r * amount):02X}{int(g * amount):02X}{int(b * amount):02X}'


def _pdf_bullet(text, color=_PDF_ACCENT, indent=16, serif: bool = False):
    style = ParagraphStyle(
        name=f'Bullet_{color}' + ('S' if serif else ''),
        fontName=MATH_FONT if serif else FONT_NAME,
        fontSize=11 if serif else 10.5,
        leading=15 if serif else 13,
        textColor=HexColor(_LECTURE_INK if serif else _PDF_DARK),
        leftIndent=indent,
        spaceAfter=2,
    )
    colored_dot = f'<font color="{color}"><b>●</b></font>'
    return Paragraph(f'{colored_dot}&nbsp;&nbsp;{text}', style)


def _pdf_prose_paragraph(text, serif: bool = False):
    style = ParagraphStyle(
        name='ProseBody' + ('Serif' if serif else ''),
        fontName=MATH_FONT if serif else FONT_NAME,
        fontSize=11 if serif else 10.5, leading=15 if serif else 13.5,
        textColor=HexColor(_LECTURE_INK if serif else _PDF_DARK),
        alignment=TA_JUSTIFY, spaceAfter=6 if serif else 5,
    )
    return [Paragraph(p.strip(), style) for p in text.split('\n\n') if p.strip()]


def _pdf_numbered_bullet(num, text, color=_PDF_ACCENT, serif: bool = False):
    style = ParagraphStyle(
        name=f'NumBullet_{num}' + ('S' if serif else ''),
        fontName=MATH_FONT if serif else FONT_NAME,
        fontSize=11 if serif else 10.5,
        leading=15 if serif else 13,
        textColor=HexColor(_LECTURE_INK if serif else _PDF_DARK),
        leftIndent=30,
        firstLineIndent=-14,
        spaceAfter=2,
    )
    return Paragraph(f'<font color="{color}"><b>{num}.</b></font>&nbsp;&nbsp;{text}', style)


def _pdf_numbered_bullet_tinted(num, text, color):
    style = ParagraphStyle(
        name=f'NumBulletTint_{num}_{abs(hash(text)) % 100000}', fontName=FONT_NAME, fontSize=10.5,
        leading=13, textColor=HexColor(_PDF_DARK),
    )
    t = Table([[Paragraph(f'<font color="{color}"><b>{num}.</b></font>&nbsp;&nbsp;{text}', style)]], colWidths=[460])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), HexColor(_pdf_light_tint_hex(color))),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 10), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    return [t, Spacer(1, 2)]


def _pdf_formula_flowable(latex: str, minimal: bool = False, accent: str | None = None):
    latex = str(latex or "").strip()
    if not latex:
        return None
    try:
        return MathFlowable(latex, 15, colour=HexColor(accent or _PDF_ACCENT),
                            back_colour=None if minimal else HexColor('#F8FAFC'),
                            pad_top=7, pad_bottom=4)
    except Exception:
        return None


def _pdf_formula_card(formula: str, explanation: str, minimal: bool = False,
                      latex: str | None = None, accent: str | None = None) -> list:
    typeset = _pdf_formula_flowable(latex or _plain_to_latex(formula), minimal, accent)
    if typeset is not None:
        elements = [typeset]
        if explanation:
            elements.append(Paragraph(
                f'<i>{explanation}</i>',
                ParagraphStyle(name=f'FormulaExplV_{abs(hash(explanation)) % 100000}',
                               fontName=FONT_NAME, fontSize=9.5, leading=12,
                               textColor=HexColor('#64748B'),
                               backColor=None if minimal else HexColor('#F8FAFC'),
                               alignment=TA_CENTER, borderPadding=(1, 10, 5, 10),
                               spaceAfter=8 if minimal else 4)))
        return elements

    formula_style = ParagraphStyle(
        name=f'Formula_{abs(hash(formula)) % 100000}',
        fontName=MATH_FONT_BOLD,
        fontSize=15,
        leading=20,
        textColor=HexColor(_PDF_ACCENT),
        backColor=None if minimal else HexColor('#F8FAFC'),
        alignment=TA_CENTER,
        borderPadding=(6, 10, 4 if explanation else 6, 10),
        spaceBefore=2,
        spaceAfter=0 if explanation else 4,
    )
    elements = [Paragraph(formula, formula_style)]
    if explanation:
        expl_style = ParagraphStyle(
            name=f'FormulaExpl_{abs(hash(explanation)) % 100000}',
            fontName=FONT_NAME,
            fontSize=9.5,
            leading=12,
            textColor=HexColor('#64748B'),
            backColor=None if minimal else HexColor('#F8FAFC'),
            alignment=TA_CENTER,
            borderPadding=(1, 10, 5, 10),
            spaceAfter=8 if minimal else 4,
        )
        elements.append(Paragraph(f'<i>{explanation}</i>', expl_style))
    return elements


def _pdf_concept_card(card: dict, accent_hex: str) -> Table:
    title = card.get('title', '')
    tag = card.get('tag', '')

    header_style = ParagraphStyle(
        name=f'CardHeader_{abs(hash(title)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=11,
        textColor=HexColor(accent_hex), alignment=TA_CENTER, spaceAfter=0,
    )
    header_content = [Paragraph(title, header_style)]
    if tag:
        tag_style = ParagraphStyle(
            name=f'CardTag_{abs(hash(tag)) % 100000}', fontName=FONT_NAME, fontSize=8,
            textColor=HexColor('#64748B'), alignment=TA_CENTER, spaceBefore=1,
        )
        header_content.append(Paragraph(f'<i>{tag}</i>', tag_style))

    body_content = []
    table_rows = card.get('table')
    if table_rows:
        t = Table(table_rows, hAlign='CENTER')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), HexColor('#E2E8F0')),
            ('FONTNAME', (0, 0), (-1, 0), MATH_FONT_BOLD),
            ('FONTNAME', (0, 1), (-1, -1), MATH_FONT),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        body_content.append(t)
    formula = card.get('formula')
    if formula:
        f_style = ParagraphStyle(
            name=f'CardFormula_{abs(hash(formula)) % 100000}', fontName=MATH_FONT_BOLD, fontSize=12,
            textColor=HexColor(accent_hex), alignment=TA_CENTER, spaceBefore=3, spaceAfter=1,
        )
        body_content.append(Paragraph(formula, f_style))
    note = card.get('note')
    if note:
        n_style = ParagraphStyle(
            name=f'CardNote_{abs(hash(note)) % 100000}', fontName=FONT_NAME, fontSize=7.5,
            textColor=HexColor('#64748B'), alignment=TA_CENTER, spaceBefore=2,
        )
        body_content.append(Paragraph(f'<i>{note}</i>', n_style))
    if not body_content:
        body_content = [Spacer(1, 4)]

    card_table = Table([[header_content], [body_content]], colWidths=[228], hAlign='CENTER')
    card_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, 0), HexColor('#FFFFFF')),
        ('BACKGROUND', (0, 1), (0, 1), HexColor('#F8FAFC')),
        ('LINEABOVE', (0, 0), (0, 0), 2.5, HexColor(accent_hex)),
        ('LINEBELOW', (0, 0), (0, 0), 0.75, HexColor(accent_hex)),
        ('BOX', (0, 0), (-1, -1), 0.75, HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (0, 0), 4), ('BOTTOMPADDING', (0, 0), (0, 0), 4),
        ('TOPPADDING', (0, 1), (0, 1), 5), ('BOTTOMPADDING', (0, 1), (0, 1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8), ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    return card_table


def _card_weight(card: dict) -> int:
    w = 0
    table = card.get('table')
    if table:
        w += len(table) * 2
    if card.get('formula'):
        w += 3
    note = card.get('note')
    if note:
        w += 1 + len(note) // 60
    return w


def _pdf_concept_card_grid(cards: list, accent_hex: str, cols: int = 2) -> list:
    if not cards:
        return []
    ordered = sorted(cards, key=_card_weight)
    flowables = [_pdf_concept_card(c, accent_hex) for c in ordered]
    odd_last_row = len(flowables) % cols == 1
    while len(flowables) % cols:
        flowables.append('')
    rows = [flowables[i:i + cols] for i in range(0, len(flowables), cols)]
    grid = Table(rows, colWidths=[232] * cols, hAlign='CENTER')
    style = [
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]
    if odd_last_row:
        last_row = len(rows) - 1
        style.append(('SPAN', (0, last_row), (cols - 1, last_row)))
        style.append(('ALIGN', (0, last_row), (cols - 1, last_row), 'CENTER'))
    grid.setStyle(TableStyle(style))
    return [Spacer(1, 4), grid, Spacer(1, 6)]


_SUBSCRIPT_DIGITS = {chr(0x2080 + d): str(d) for d in range(10)}


_MATH_DIR = os.path.join(os.path.dirname(__file__), "..", "uploads", "math")
_MATH_SPAN = re.compile(r"\$([^$\n]{1,400}?)\$")
_MATH_PAD = 4
_MATH_OVERSAMPLE = 6
_MATH_CACHE_VERSION = 2



_MATH_PDF_FONTS: dict = {}


def _math_pdf_font(path: str) -> str:
    if path in _MATH_PDF_FONTS:
        return _MATH_PDF_FONTS[path]
    name = FONT_NAME
    try:
        if os.path.exists(path):
            candidate = "MathFace" + str(abs(hash(path)) % 100000)
            if candidate not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(candidate, path))
            name = candidate
    except Exception:
        name = MATH_FONT or FONT_NAME
    _MATH_PDF_FONTS[path] = name
    return name


class _PdfMathPainter:

    def __init__(self, canvas, origin_x: float, baseline_y: float,
                 scale: float, colour):
        self.c = canvas
        self.x0 = origin_x
        self.y0 = baseline_y
        self.k = scale
        self.colour = colour

    def text(self, x, y, s, font_path, size_px):
        self.c.setFont(_math_pdf_font(font_path), size_px * self.k)
        self.c.setFillColor(self.colour)
        self.c.drawString(self.x0 + x * self.k, self.y0 - y * self.k, s)

    def line(self, x1, y1, x2, y2, width):
        self.c.setStrokeColor(self.colour)
        self.c.setLineWidth(max(0.35, width * self.k))
        self.c.line(self.x0 + x1 * self.k, self.y0 - y1 * self.k,
                    self.x0 + x2 * self.k, self.y0 - y2 * self.k)


class MathFlowable(_RLFlowable):

    def __init__(self, latex: str, size: float = 14, align: str = "CENTER",
                 colour=None, avail_width: float | None = None,
                 back_colour=None, pad_top: float = 0, pad_bottom: float = 0,
                 pad_x: float = 10):
        super().__init__()
        from app.math_render import parse, _layout, SS
        self._ss = SS
        self._box = _layout(parse(latex), size)
        self._colour = colour or HexColor("#111111")
        self.align = align
        self.width = self._box.w / SS
        self.height = (self._box.asc + self._box.desc) / SS
        self._avail = avail_width
        self.back_colour = back_colour
        self.pad_top = pad_top
        self.pad_bottom = pad_bottom
        self.pad_x = pad_x

    def wrap(self, avail_w, avail_h):
        self._avail = avail_w
        usable = max(1.0, avail_w - 2 * self.pad_x)
        self._shrink = min(1.0, (usable / self.width) if self.width else 1.0)
        return (avail_w, self.height * self._shrink + self.pad_top + self.pad_bottom)

    def draw(self):
        shrink = getattr(self, "_shrink", 1.0)
        width = self.width * shrink
        avail = self._avail if self._avail is not None else width
        if self.back_colour is not None:
            self.canv.setFillColor(self.back_colour)
            self.canv.rect(0, 0, avail,
                           self.height * shrink + self.pad_top + self.pad_bottom,
                           stroke=0, fill=1)
        if self.align == "CENTER":
            x = max(0.0, (avail - width) / 2)
        elif self.align == "RIGHT":
            x = max(self.pad_x, avail - width - self.pad_x)
        else:
            x = self.pad_x
        baseline = self.pad_bottom + (self._box.desc / self._ss) * shrink
        self._box.paint(_PdfMathPainter(self.canv, x, baseline,
                                        shrink / self._ss, self._colour), 0, 0)


def _plain_to_latex(text) -> str:
    try:
        from app.math_render import plain_to_latex
        return plain_to_latex(text)
    except Exception:
        return ""


def _normalize_math(text) -> str:
    try:
        from app.math_render import normalize_math
        return normalize_math(text)
    except Exception:
        return str(text or "")


def _math_png(latex: str, size: float, bg: str = "white"):
    from app.math_render import render_png, measure
    try:
        key = hashlib.sha1(
            f"{latex}|{size:.2f}|{bg}|x{_MATH_OVERSAMPLE}|{_MATH_CACHE_VERSION}".encode("utf-8")).hexdigest()
        path = os.path.join(_MATH_DIR, key + ".png")
        if not os.path.exists(path):
            png = render_png(latex, size, _MATH_PAD, bg, oversample=_MATH_OVERSAMPLE)
            if not png:
                return None
            os.makedirs(_MATH_DIR, exist_ok=True)
            with open(path, "wb") as f:
                f.write(png)
        w, asc, desc = measure(latex, size)
        if w <= 0:
            return None
        return path, w, asc, desc
    except Exception:
        return None


def _math_inline(text, size: float = 11, bg: str = "white"):
    raw = _normalize_math(text)
    if "$" not in raw:
        return raw

    def repl(m):
        got = _math_png(m.group(1), size, bg)
        if not got:
            return m.group(1)
        path, w, asc, desc = got
        return (f'<img src="{path}" width="{w + 2 * _MATH_PAD:.1f}" '
                f'height="{asc + desc + 2 * _MATH_PAD:.1f}" '
                f'valign="{-(desc + _MATH_PAD):.1f}"/>')
    return _MATH_SPAN.sub(repl, raw)


def _pdf_math_display(latex: str, size: float = 15, width: float = 430, bg: str = "white"):
    try:
        formula = MathFlowable(latex, size)
        if formula.width <= 0:
            return None
    except Exception:
        return None
    t = Table([[formula]], colWidths=[width])
    t.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    return t


def _sub_digits(text) -> str:
    out = str(text or "")
    for ch, digit in _SUBSCRIPT_DIGITS.items():
        if ch in out:
            out = out.replace(ch, f"<sub>{digit}</sub>")
    return out


def _pdf_subject_figure(fig: dict, minimal: bool = False, max_w: float = 300,
                        max_h: float = 230) -> list:
    path = fig.get("image")
    if not path:
        return []
    full = os.path.join(os.path.dirname(__file__), "..", path.lstrip("/"))
    if not os.path.exists(full):
        return []
    try:
        with PILImage.open(full) as im:
            iw, ih = im.size
    except Exception:
        return []
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    body = [_PdfFlowableImage(full, width=w, height=h, hAlign="CENTER")]
    caption = _sub_digits(fig.get("caption"))
    if caption:
        body.append(Paragraph(
            f"<i>{caption}</i>",
            ParagraphStyle(
                name=f"FigCap_{abs(hash(str(caption))) % 100000}",
                fontName=FONT_NAME, fontSize=9.5, leading=12.5, alignment=TA_CENTER,
                textColor=HexColor("#000000" if minimal else "#64748B"),
                spaceBefore=3, spaceAfter=8,
            ),
        ))
    return [Spacer(1, 6), KeepTogether(body), Spacer(1, 6)]


class _PdfContentAnchor(Spacer):

    def __init__(self):
        super().__init__(1, 0)


def _pdf_body_splice_index(body: list, nth_section: int = 2) -> int:
    if nth_section == 1:
        for i, flowable in enumerate(body):
            if isinstance(flowable, _PdfContentAnchor):
                return i + 1
    seen = 0
    for i, flowable in enumerate(body):
        if isinstance(flowable, KeepTogether):
            seen += 1
            if seen == nth_section:
                return i
    return min(len(body), max(1, len(body) * (nth_section - 1) // 3))


class _PdfPageGate(_RLFlowable):

    def __init__(self, target_page: int, max_defers: int = 4):
        super().__init__()
        self.target_page = target_page
        self._defers_left = max_defers

    def wrap(self, avail_w, avail_h):
        current = self._doctemplateAttr("page") or 1
        if current < self.target_page and self._defers_left > 0:
            self._defers_left -= 1
            return (avail_w, avail_h + 1)
        return (0, 0)

    def draw(self):
        pass


def _pdf_lesson_image_block(image: dict, label: str, accent_hex: str) -> list:
    body = _pdf_illustration_image(
        image.get("path", ""), image.get("caption", ""),
        image.get("credit", ""), max_w=300, max_h=215,
    )
    if not body:
        return []
    key = abs(hash(image.get("path", ""))) % 100000
    tag = Paragraph(
        label.upper(),
        ParagraphStyle(name=f'LessonImgTag_{key}', fontName=FONT_NAME_BOLD, fontSize=8.5,
                       textColor=HexColor(accent_hex), alignment=TA_CENTER,
                       spaceBefore=8, spaceAfter=2),
    )
    parts = [tag] + body
    explanation = str(image.get("explanation") or "").strip()
    if explanation:
        parts.append(Paragraph(
            explanation,
            ParagraphStyle(name=f'LessonImgNote_{key}', fontName=FONT_NAME, fontSize=9.5,
                           leading=13, alignment=TA_CENTER, textColor=HexColor('#334155'),
                           leftIndent=40, rightIndent=40, spaceBefore=2, spaceAfter=2),
        ))
    return [KeepTogether(parts), Spacer(1, 8)]


def _pdf_illustration_image(image_path: str, caption: str, attribution: str, max_w: float = 460, max_h: float = 320) -> list | None:
    full_path = os.path.join(os.path.dirname(__file__), "..", image_path.lstrip("/"))
    if not os.path.exists(full_path):
        return None
    try:
        with PILImage.open(full_path) as img:
            iw, ih = img.size
    except Exception:
        return None
    w, h = max_w, max_w * ih / iw
    if h > max_h:
        h = max_h
        w = max_h * iw / ih
    elements = [
        Spacer(1, 4),
        _PdfFlowableImage(full_path, width=w, height=h, hAlign='CENTER'),
    ]
    if caption:
        cap_style = ParagraphStyle(
            name=f'ImgCaption_{abs(hash(caption)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=9.5,
            alignment=TA_CENTER, textColor=HexColor('#334155'), spaceBefore=5, spaceAfter=1,
        )
        elements.append(Paragraph(_xml_escape(caption), cap_style))
    if attribution:
        attr_style = ParagraphStyle(
            name=f'ImgAttribution_{abs(hash(attribution)) % 100000}', fontName=FONT_NAME, fontSize=7.5,
            alignment=TA_CENTER, textColor=HexColor('#94A3B8'), spaceAfter=4,
        )
        elements.append(Paragraph(_xml_escape(attribution), attr_style))
    return elements


def _pdf_answer_box(text, accent: str | None = None):
    style = ParagraphStyle(
        name='AnswerBox',
        fontName=FONT_NAME,
        fontSize=10,
        leading=14,
        textColor=HexColor('#475569'),
        backColor=HexColor(_PDF_LIGHT),
        borderPadding=(6, 8, 6, 8),
        leftIndent=25,
        rightIndent=15,
        spaceAfter=8,
    )
    note = Table([[Paragraph(str(text), style)]], colWidths=[470])
    note.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(_PDF_LIGHT)),
        ("LINEBEFORE", (0, 0), (0, -1), 2.5, HexColor(accent or _PDF_ACCENT)),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    return note


def _pdf_footer():
    elements = []
    elements.append(Spacer(1, 20))
    d = Drawing(500, 20)
    d.add(Rect(0, 0, 500, 20, fillColor=HexColor(_PDF_DARK), strokeColor=None))
    d.add(String(150, 6, 'Dastyor — AI-помощник для учителей', fontName=FONT_NAME, fontSize=8, fillColor=HexColor('#94A3B8')))
    elements.append(d)
    return elements


def _styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name='SlideTitle',
        fontName=FONT_NAME_BOLD,
        fontSize=13,
        leading=16,
        textColor=HexColor(_PDF_DARK),
        spaceAfter=2,
    ))
    styles.add(ParagraphStyle(
        name='SlideBody',
        fontName=FONT_NAME,
        fontSize=11,
        leading=15,
        textColor=HexColor('#374151'),
        leftIndent=20,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name='DocTitle',
        fontName=FONT_NAME_BOLD,
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        textColor=HexColor(_PDF_DARK),
        spaceAfter=12,
    ))
    styles.add(ParagraphStyle(
        name='DocSubtitle',
        fontName=FONT_NAME,
        fontSize=12,
        leading=15,
        alignment=TA_CENTER,
        textColor=HexColor('#6B7280'),
        spaceAfter=8,
    ))
    return styles


_COLORS = [_PDF_ACCENT, _PDF_GREEN, _PDF_ORANGE, _PDF_RED, _PDF_PURPLE, '#06B6D4', '#EC4899', '#14B8A6']


_PDF_KONSPEKT_LABELS = {
    'Русский': {
        'objective': 'Цель лекции', 'lecture_plan': 'План лекции', 'references': 'Литература',
        'competencies': 'Компетенции', 'objectives': 'Цели урока', 'key_concepts': 'Ключевые понятия',
        'key_terms': 'Словарь урока', 'fun_facts': 'Интересные факты', 'formulas': 'Формулы',
        'lesson_program': 'План урока', 'tools': 'Материалы', 'warmup': 'Разминка',
        'main_content': 'Основное содержание', 'real_life_examples': 'Примеры из жизни', 'map': 'Карта', 'illustration': 'Иллюстрация', 'diagrams': 'Схемы',
        'pair_work': 'Работа в парах', 'consolidation': 'Закрепление', 'common_mistakes': 'Частые ошибки',
        'group_work': 'Групповая работа', 'group': 'Группа',
        'homework': 'Домашнее задание', 'summary': 'Итоги урока', 'assessment': 'Оценивание',
        'code_blocks': 'Код', 'quick_check': 'Быстрая проверка', 'answer': 'Ответ',
        'worked_examples': 'Решённые примеры', 'solution': 'Решение',
        'practice_problems': 'Практические задачи', 'answers': 'Ответы',
        'plan_note': 'Внимание', 'plan_date': 'Дата', 'plan_class': 'Класс', 'plan_school': 'Школа', 'example': 'Пример', 'warmup': 'Проверка знаний', 'plan_year': 'Учебный год',
        'tag': 'КОНСПЕКТ', 'grade_suffix': 'класс',
    },
    'Таджикский': {
        'objective': 'Мақсади лексия', 'lecture_plan': 'Нақшаи лексия', 'references': 'Адабиёт',
        'competencies': 'Салоҳиятҳо', 'objectives': 'Мақсадҳои дарс', 'key_concepts': 'Мафҳумҳои асосӣ',
        'key_terms': 'Луғати дарс', 'fun_facts': 'Фактҳои ҷолиб', 'formulas': 'Формулаҳо',
        'lesson_program': 'Барномаи дарс', 'tools': 'Воситаҳои аёнӣ', 'warmup': 'Санҷиши дониш',
        'main_content': 'Шиносоӣ бо мазмуни мавзӯъ', 'real_life_examples': 'Дар ҳаёти воқеӣ', 'map': 'Харита', 'illustration': 'Тасвир', 'diagrams': 'Расмҳо',
        'pair_work': 'Кори дунафара', 'consolidation': 'Мустаҳкамкунии дарс', 'common_mistakes': 'Хатогиҳои маъмул',
        'group_work': 'Кори гурӯҳӣ', 'group': 'Гурӯҳи',
        'homework': 'Супориши хонагӣ', 'summary': 'Хулосаи дарс', 'assessment': 'Арзёбӣ',
        'code_blocks': 'Код', 'quick_check': 'Санҷиши зуд', 'answer': 'Ҷавоб',
        'worked_examples': 'Мисолҳои ҳалшуда', 'solution': 'Ҳал',
        'practice_problems': 'Машқҳои мустақил', 'answers': 'Ҷавобҳо',
        'plan_note': 'Диққат', 'plan_date': 'Сана', 'plan_class': 'Синф', 'plan_school': 'Мактаб', 'example': 'Мисоли', 'warmup': 'Санҷиши дониш', 'plan_year': 'Соли таҳсили',
        'tag': 'КОНСПЕКТ', 'grade_suffix': 'синф',
    },
    'English': {
        'objective': 'Objective', 'lecture_plan': 'Lecture plan', 'references': 'References',
        'competencies': 'Competencies', 'objectives': 'Objectives', 'key_concepts': 'Key Concepts',
        'key_terms': 'Lesson Glossary', 'fun_facts': 'Fun Facts', 'formulas': 'Formulas',
        'lesson_program': 'Lesson Program', 'tools': 'Tools', 'warmup': 'Warm-up',
        'main_content': 'Main Content', 'real_life_examples': 'Real Life Examples', 'map': 'Map', 'illustration': 'Illustration', 'diagrams': 'Diagrams',
        'pair_work': 'Pair Work', 'consolidation': 'Consolidation', 'common_mistakes': 'Common Mistakes',
        'group_work': 'Group Work', 'group': 'Group',
        'homework': 'Homework', 'summary': 'Summary', 'assessment': 'Assessment',
        'code_blocks': 'Code', 'quick_check': 'Quick Check', 'answer': 'Answer',
        'worked_examples': 'Worked Examples', 'solution': 'Solution',
        'practice_problems': 'Practice Problems', 'answers': 'Answers',
        'plan_note': 'Note', 'plan_date': 'Date', 'plan_class': 'Class', 'plan_school': 'School', 'example': 'Example', 'warmup': 'Knowledge check', 'plan_year': 'Academic year',
        'tag': 'KONSPEKT', 'grade_suffix': 'grade',
    },
}
_PDF_KONSPEKT_LABELS['Английский'] = _PDF_KONSPEKT_LABELS['English']


def build_konspekt_pdf(content: dict, language: str = 'Русский', doc_type_label: str | None = None, skip_cover: bool = False) -> io.BytesIO:
    if doc_type_label is None and not skip_cover:
        _tmpl_check = get_template(content.get('template') or 'zamonaviy')
        if not _tmpl_check.plan_layout:
            try:
                from app.konspekt_builder import build_konspekt_pdf as _build_modern
                return _build_modern(content, language)
            except Exception as e:
                logger.exception(f"modern konspekt renderer failed, falling back: {e}")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=0.9 * cm, bottomMargin=1.3 * cm,
                            leftMargin=1.4 * cm, rightMargin=1.4 * cm)
    story = []

    L = dict(_PDF_KONSPEKT_LABELS.get(language, _PDF_KONSPEKT_LABELS['Русский']))
    if doc_type_label:
        L['tag'] = doc_type_label

    title = content.get("title", "Конспект урока")
    subtitle = content.get("subtitle", "")
    subject = content.get("subject", "")
    grade = content.get("grade", "")

    from app.cover_builder import _academic_year

    cover_png = None
    if not skip_cover:
        try:
            from app.cover_builder import build_cover_image
            cover_png = build_cover_image(subject, title, grade, language, template_id=content.get("template") or "zamonaviy", doc_type_label=doc_type_label)
        except Exception as e:
            logger.warning(f"Cover generation failed for '{title[:50]}': {e}")

    _plan_sheet = get_template(content.get("template") or "zamonaviy").plan_layout

    def _draw_running_footer(canvas_obj, doc_obj):
        if _plan_sheet:
            canvas_obj.saveState()
            canvas_obj.setFont(MATH_FONT, 10)
            canvas_obj.setFillColor(HexColor('#000000'))
            canvas_obj.drawCentredString(A4[0] / 2, 1.0 * cm, str(doc_obj.page))
            canvas_obj.restoreState()
            return
        canvas_obj.saveState()
        canvas_obj.setStrokeColor(HexColor('#E2E8F0'))
        canvas_obj.setLineWidth(0.6)
        canvas_obj.line(1.6 * cm, 1.0 * cm, A4[0] - 1.6 * cm, 1.0 * cm)
        canvas_obj.setFont(FONT_NAME, 8)
        canvas_obj.setFillColor(HexColor('#94A3B8'))
        canvas_obj.drawString(1.6 * cm, 0.65 * cm, title[:70])
        canvas_obj.drawRightString(A4[0] - 1.6 * cm, 0.65 * cm, f'Dastyor  •  Страница {doc_obj.page}')
        canvas_obj.setFont(FONT_NAME, 7.5)
        canvas_obj.setFillColor(HexColor('#CBD5E1'))
        canvas_obj.drawRightString(A4[0] - 1.6 * cm, A4[1] - 0.7 * cm, _academic_year())
        canvas_obj.restoreState()

    def _draw_cover(canvas_obj, doc_obj):
        if not cover_png:
            _draw_running_footer(canvas_obj, doc_obj)
            return
        canvas_obj.drawImage(ImageReader(io.BytesIO(cover_png)), 0, 0, width=A4[0], height=A4[1])

    if cover_png:
        from reportlab.platypus import PageBreak
        story.append(PageBreak())

    tag_line = L['tag']
    if subject:
        tag_line += f"  •  {subject}"
    if grade:
        tag_line += f"  •  {_format_grade(grade, L['grade_suffix'])}"

    _tmpl = get_template(content.get("template") or "zamonaviy")
    accent_hex = get_subject_accent_hex(subject)

    if _tmpl.plan_layout:
        body = _pdf_plan_body(content, L)
    elif content.get("lecture_plan") or content.get("objective"):
        story.extend(_pdf_lecture_masthead(
            title, subtitle, subject, grade, language, L['grade_suffix']))
        body = _add_konspekt_body_pdf(content, L)
    else:
        story.extend(_pdf_header(
            title, subtitle, tag_line,
            accent=accent_hex,
            minimal=_tmpl.minimal_chrome,
        ))
        body = _add_konspekt_body_pdf(content, L)

    story.extend(body)

    doc.build(story, onFirstPage=_draw_cover, onLaterPages=_draw_running_footer)
    buf.seek(0)
    return buf


_LECTURE_TAG = {
    'Русский': 'ЛЕКЦИЯ', 'Таджикский': 'ЛЕКСИЯ', 'English': 'LECTURE',
    'Английский': 'LECTURE',
}


def build_lecture_pdf(content: dict, language: str = 'Русский') -> io.BytesIO:
    from app.lecture_builder import build_lecture_pdf as _build
    try:
        return _build(content, language)
    except Exception as e:
        logger.exception(f"lecture renderer failed, falling back: {e}")
        tag = _LECTURE_TAG.get(language, _LECTURE_TAG['Русский'])
        return build_konspekt_pdf(content, language, doc_type_label=tag, skip_cover=True)


def _pdf_important_note(text: str, accent_hex: str) -> Table:
    style = ParagraphStyle(
        name=f'ImportantNote_{abs(hash(text)) % 100000}', fontName=FONT_NAME, fontSize=9.5,
        leading=13, textColor=HexColor('#7F1D1D'),
    )
    t = Table([[Paragraph(f'<b>!</b> <i>{text}</i>', style)]], colWidths=[470])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), HexColor(_PDF_BRAND_SOFT)),
        ('LINEBEFORE', (0, 0), (0, -1), 3, HexColor(accent_hex)),
        ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10), ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    return t


def _pdf_code_card(code: str, language: str, explanation: str, accent_hex: str) -> list:
    escaped = (
        code.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('\n', '<br/>')
    )
    badge_style = ParagraphStyle(
        name=f'CodeBadge_{abs(hash(code)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=8,
        textColor=HexColor('#FFFFFF'),
    )
    code_style = ParagraphStyle(
        name=f'Code_{abs(hash(code)) % 100000}', fontName=CODE_FONT, fontSize=9, leading=13,
        textColor=HexColor('#E2E8F0'),
    )
    t = Table(
        [[Paragraph((language or 'code').upper(), badge_style)], [Paragraph(escaped, code_style)]],
        colWidths=[470],
    )
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, 0), HexColor(accent_hex)),
        ('BACKGROUND', (0, 1), (0, 1), HexColor('#1E293B')),
        ('TOPPADDING', (0, 0), (0, 0), 3), ('BOTTOMPADDING', (0, 0), (0, 0), 3),
        ('LEFTPADDING', (0, 0), (0, 0), 10), ('RIGHTPADDING', (0, 0), (0, 0), 8),
        ('TOPPADDING', (0, 1), (0, 1), 8), ('BOTTOMPADDING', (0, 1), (0, 1), 8),
        ('LEFTPADDING', (0, 1), (0, 1), 10), ('RIGHTPADDING', (0, 1), (0, 1), 8),
    ]))
    elements = [t]
    if explanation:
        expl_style = ParagraphStyle(
            name=f'CodeExpl_{abs(hash(explanation)) % 100000}', fontName=FONT_NAME, fontSize=9.5,
            leading=12, textColor=HexColor('#64748B'), spaceBefore=4, spaceAfter=6,
        )
        elements.append(Paragraph(f'<i>{explanation}</i>', expl_style))
    else:
        elements.append(Spacer(1, 6))
    return [KeepTogether(elements)]


def _pdf_quick_check(items: list, accent_hex: str, answer_label: str) -> list:
    elements = []
    for item in items:
        if not isinstance(item, dict):
            continue
        question = item.get('question', '')
        answer = item.get('answer', '')
        if not question:
            continue
        q_style = ParagraphStyle(
            name=f'QCQ_{abs(hash(question)) % 100000}', fontName=FONT_NAME, fontSize=9.5,
            leading=13, textColor=HexColor('#1F2937'),
        )
        rows = [[Paragraph(f'<b>?</b>&nbsp;&nbsp;{question}', q_style)]]
        if answer:
            a_style = ParagraphStyle(
                name=f'QCA_{abs(hash(answer)) % 100000}', fontName=FONT_NAME, fontSize=9,
                leading=12, textColor=HexColor('#166534'),
            )
            rows.append([Paragraph(f'<b>{answer_label}:</b> <i>{answer}</i>', a_style)])
        t = Table(rows, colWidths=[470])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), HexColor('#F0FDF4')),
            ('LINEBEFORE', (0, 0), (0, -1), 3, HexColor(accent_hex)),
            ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 10), ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 4))
    return elements


def _pdf_worked_examples(items: list, accent_hex: str, solution_label: str, minimal: bool = False) -> list:
    elements = []
    tint = _pdf_light_tint_hex(accent_hex, amount=0.93)
    for i, item in enumerate(items):
        if not isinstance(item, dict):
            continue
        problem = item.get('problem', '')
        solution = item.get('solution', '')
        if not problem:
            continue
        n_style = ParagraphStyle(
            name=f'WEN_{i}_{abs(hash(problem)) % 100000}', fontName=FONT_NAME_BOLD, fontSize=13,
            leading=16,
            textColor=HexColor(accent_hex if minimal else '#FFFFFF'),
            alignment=1,
        )
        p_style = ParagraphStyle(
            name=f'WEP_{abs(hash(problem)) % 100000}', fontName=FONT_NAME, fontSize=9.5,
            leading=13, textColor=HexColor('#1F2937'),
        )
        body = [Paragraph(f'<b>{problem}</b>', p_style)]
        if solution:
            s_style = ParagraphStyle(
                name=f'WES_{abs(hash(solution)) % 100000}', fontName=FONT_NAME, fontSize=9,
                leading=12.5, textColor=HexColor('#334155'), spaceBefore=3,
            )
            body.append(Paragraph(
                f'<font color="#1D4ED8"><b>{solution_label}:</b></font> {solution}', s_style,
            ))
        t = Table([[Paragraph(str(i + 1), n_style), body]], colWidths=[26, 444])
        if minimal:
            style = [
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LINEBELOW', (0, 0), (-1, -1), 0.5, HexColor('#E2E8F0')),
                ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('LEFTPADDING', (0, 0), (0, 0), 0), ('RIGHTPADDING', (0, 0), (0, 0), 6),
                ('LEFTPADDING', (1, 0), (1, 0), 0), ('RIGHTPADDING', (1, 0), (1, 0), 0),
            ]
        else:
            style = [
                ('BACKGROUND', (0, 0), (0, 0), HexColor(accent_hex)),
                ('BACKGROUND', (1, 0), (1, 0), HexColor(tint)),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('LEFTPADDING', (0, 0), (0, 0), 2), ('RIGHTPADDING', (0, 0), (0, 0), 2),
                ('LEFTPADDING', (1, 0), (1, 0), 10), ('RIGHTPADDING', (1, 0), (1, 0), 10),
            ]
        t.setStyle(TableStyle(style))
        elements.append(KeepTogether([t, Spacer(1, 5)]))
    return elements


def _pdf_sidebar_panel(content: dict, sections: tuple, accent_hex: str, L: dict) -> list:
    rows_present = [(key, content.get(key)) for key in sections if content.get(key)]
    if not rows_present:
        return []
    tint = _pdf_light_tint_hex(accent_hex, amount=0.92)
    cell_elements = []
    for key, items in rows_present:
        header_style = ParagraphStyle(
            name=f'SidebarHdr_{key}', fontName=FONT_NAME_BOLD, fontSize=8.5,
            textColor=HexColor(accent_hex), spaceBefore=6, spaceAfter=2,
        )
        cell_elements.append(Paragraph(L.get(key, key).upper(), header_style))
        item_style = ParagraphStyle(
            name=f'SidebarItem_{key}', fontName=FONT_NAME, fontSize=9.5,
            leading=12, textColor=HexColor(_PDF_DARK), leftIndent=8, spaceAfter=3,
        )
        for item in items:
            cell_elements.append(Paragraph(f'•&nbsp;&nbsp;{item}', item_style))
    t = Table([[cell_elements]], colWidths=[470])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), HexColor(tint)),
        ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 12), ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    return [Spacer(1, 4), t, Spacer(1, 8)]


_VISUAL_TABLE_WIDTH = 500.0


_LECTURE_SUMMARY_LABEL = {
    'Русский': 'Выводы', 'Таджикский': 'Хулосаҳо',
    'English': 'Conclusions',
}
_LECTURE_SUMMARY_LABEL['Английский'] = _LECTURE_SUMMARY_LABEL['English']


def _pdf_visual_block(block: dict, accent_hex: str, minimal: bool = False,
                      serif: bool = False) -> list:
    btype = block.get('type')
    data = block.get('data') or {}
    elements = []
    title = block.get('title')
    if title:
        cap_style = ParagraphStyle(
            name=f'VisualBlockCaption_{abs(hash(title)) % 100000}', fontName=FONT_NAME, fontSize=9.5,
            textColor=HexColor('#000000' if minimal else '#64748B'), spaceBefore=6, spaceAfter=3,
        )
        elements.append(Paragraph(f'<i>{title}</i>', cap_style))

    try:
        if btype in ('table', 'comparison'):
            if btype == 'table':
                headers = data.get('headers') or []
                rows = data.get('rows') or []
            else:
                criteria = data.get('criteria') or []
                items = data.get('items') or []
                headers = [''] + criteria
                rows = [[it.get('name', '')] + list(it.get('values') or []) for it in items]
            if not headers or not rows:
                return []
            col_count = max(1, len(headers))
            weights = []
            for c in range(col_count):
                longest = max([len(str(headers[c]))] +
                              [len(str(r[c])) for r in rows if c < len(r)] or [1])
                weights.append(max(6, longest))
            total_w = _VISUAL_TABLE_WIDTH
            min_w = min(52.0, total_w / col_count)
            free = total_w - min_w * col_count
            weight_sum = sum(weights) or 1
            col_widths = [min_w + free * (w / weight_sum) for w in weights]

            cell_style = ParagraphStyle(
                name='VisCell', fontName=FONT_NAME, fontSize=9, leading=12,
                alignment=TA_CENTER,
                textColor=HexColor('#000000' if minimal else '#1F2937'))
            head_style = ParagraphStyle(
                name='VisCellHead', parent=cell_style, fontName=FONT_NAME_BOLD,
                textColor=HexColor(accent_hex if minimal else '#FFFFFF'))
            headers = [Paragraph(_math_inline(str(h), 9), head_style) for h in headers]
            rows = [[Paragraph(_math_inline(str(c), 9), cell_style) for c in row]
                    for row in rows]
            t = Table([headers] + rows, colWidths=col_widths, hAlign='CENTER')
            t.setStyle(TableStyle([
            ] + ([] if minimal else [
                ('BACKGROUND', (0, 0), (-1, 0), HexColor(accent_hex)),
            ]) + [
                ('TEXTCOLOR', (0, 0), (-1, 0), HexColor(accent_hex if minimal else '#FFFFFF')),
                ('FONTNAME', (0, 0), (-1, 0), FONT_NAME_BOLD),
                ('FONTNAME', (0, 1), (-1, -1), FONT_NAME),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ] + ([
                ('LINEBELOW', (0, 0), (-1, 0), 0.8, HexColor(accent_hex)),
                ('LINEBELOW', (0, 1), (-1, -2), 0.4, HexColor('#E2E8F0')),
            ] if minimal else [
                ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#E2E8F0')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [HexColor('#FFFFFF'), HexColor('#FEF2F2')]),
            ])))
            elements.extend([Spacer(1, 2), t, Spacer(1, 4)])

        elif btype == 'timeline':
            events = data.get('events') or []
            if events:
                elements.append(Spacer(1, 2))
                for ev in events:
                    date = str(ev.get('date', '')).strip()
                    label = str(ev.get('label', '')).strip()
                    head = (f'<font color="{accent_hex}"><b>{date}</b></font>&nbsp;&nbsp;' if date else '')
                    elements.append(Paragraph(
                        f'{head}<b>{label}</b>',
                        ParagraphStyle(
                            name=f'TlT{abs(hash(label + date)) % 99999}',
                            fontName=MATH_FONT_BOLD if serif else FONT_NAME_BOLD,
                            fontSize=11 if serif else 10.5, leading=15 if serif else 14,
                            textColor=HexColor(_LECTURE_INK if serif else _PDF_DARK),
                            leftIndent=22, firstLineIndent=-22, spaceBefore=5, spaceAfter=1)))
                    if ev.get('description'):
                        elements.append(Paragraph(
                            str(ev['description']),
                            ParagraphStyle(
                                name=f'TlD{abs(hash(str(ev["description"]))) % 99999}',
                                fontName=MATH_FONT if serif else FONT_NAME,
                                fontSize=10 if serif else 9.5, leading=14 if serif else 13,
                                textColor=HexColor(_LECTURE_MUTED if serif else '#475569'),
                                leftIndent=22, spaceAfter=2)))
                elements.append(Spacer(1, 4))

        elif btype in ('flowchart', 'process'):
            steps = data.get('steps') or []
            if steps:
                elements.append(Spacer(1, 2))
                for i, step in enumerate(steps):
                    title = str(step.get('title', '')).strip()
                    desc = str(step.get('description', '')).strip()
                    if title:
                        elements.append(Paragraph(
                            f'<font color="{accent_hex}"><b>{i + 1}.</b></font>'
                            f'&nbsp;&nbsp;<b>{title}</b>',
                            ParagraphStyle(
                                name=f'StepT{abs(hash(title)) % 99999}',
                                fontName=MATH_FONT_BOLD if serif else FONT_NAME_BOLD,
                                fontSize=11 if serif else 10.5, leading=15 if serif else 14,
                                textColor=HexColor(_LECTURE_INK if serif else _PDF_DARK),
                                leftIndent=22, firstLineIndent=-22, spaceBefore=5, spaceAfter=1)))
                    if desc:
                        elements.append(Paragraph(
                            desc,
                            ParagraphStyle(
                                name=f'StepD{abs(hash(desc)) % 99999}',
                                fontName=MATH_FONT if serif else FONT_NAME,
                                fontSize=10 if serif else 9.5, leading=14 if serif else 13,
                                textColor=HexColor(_LECTURE_MUTED if serif else '#475569'),
                                leftIndent=22, spaceAfter=2)))
                elements.append(Spacer(1, 4))

        elif btype == 'figure':
            elements.extend(_pdf_subject_figure(block, minimal=minimal, max_w=340, max_h=260))

        elif btype == 'chart':
            categories = [str(c) for c in (data.get('categories') or [])]
            series = data.get('series') or []
            if categories and series:
                headers = [''] + categories
                rows = [[str(s.get('name', ''))] + [str(v) for v in (s.get('values') or [])] for s in series]
                t = Table([headers] + rows, hAlign='CENTER')
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), HexColor(accent_hex)),
                    ('TEXTCOLOR', (0, 0), (-1, 0), HexColor('#FFFFFF')),
                    ('FONTNAME', (0, 0), (-1, 0), FONT_NAME_BOLD),
                    ('FONTNAME', (0, 1), (-1, -1), FONT_NAME),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#E2E8F0')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [HexColor('#FFFFFF'), HexColor('#FEF2F2')]),
                    ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ]))
                elements.extend([Spacer(1, 2), t, Spacer(1, 4)])

        elif btype == 'concept_map':
            nodes = {n.get('id'): n.get('label', n.get('id', '')) for n in (data.get('nodes') or [])}
            for edge in (data.get('edges') or []):
                frm = nodes.get(edge.get('from'), edge.get('from', ''))
                to = nodes.get(edge.get('to'), edge.get('to', ''))
                text = f'<b>{frm}</b> &nbsp;→&nbsp; <b>{to}</b>'
                if edge.get('label'):
                    text += f" — {edge['label']}"
                elements.append(_pdf_bullet(text, accent_hex, serif=serif))
    except Exception:
        return []
    if not elements:
        return elements
    if btype in ('table', 'comparison', 'chart'):
        return elements
    return [KeepTogether(elements)]



_PLAN_RUN_IN = [
    ("tools", "prose"),
    ("main_content", "prose"),
    ("lesson_program", "prose"),
    ("pair_work", "prose"),
    ("consolidation", "prose"),
    ("summary", "prose"),
    ("assessment", "prose"),
]

_PLAN_BULLETS = ["objectives", "key_concepts", "key_terms", "real_life_examples", "group_work", "homework"]

_PLAN_FILL_IN = ["competencies", "warmup"]


_PLAN_BG = "#EAF3FB"
_PLAN_BAR = "#8FBEDC"
_PLAN_INK = "#1B4460"
_PLAN_MARK = "#D3E7F6"
_PLAN_EX_BG = "#F4F9FD"


def _plan_panel(flowables, width=470, bar=True):
    t = Table([[flowables]], colWidths=[width])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(_PLAN_BG)),
    ] + ([("LINEBEFORE", (0, 0), (0, -1), 2.5, HexColor(_PLAN_BAR))] if bar else []) + [
        ("LEFTPADDING", (0, 0), (-1, -1), 11), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return [Spacer(1, 4), t, Spacer(1, 8)]


def _plan_feature_head(text, S, badge=None):
    cells = [Paragraph(f"<b>{text}</b>", S["head"])]
    widths = [470]
    if badge:
        cells.append(Paragraph(
            f'<para alignment="right">{badge}</para>',
            ParagraphStyle(name="PlanBadge", parent=S["head"], fontSize=9.5),
        ))
        widths = [370, 100]
    t = Table([cells], colWidths=widths)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(_PLAN_BG)),
        ("LINEBEFORE", (0, 0), (0, -1), 2.5, HexColor(_PLAN_BAR)),
        ("LEFTPADDING", (0, 0), (0, -1), 11), ("RIGHTPADDING", (-1, 0), (-1, -1), 11),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return [Spacer(1, 6), t, Spacer(1, 6)]


def _plan_example_card(label, problem, solution, S):
    inner = [Paragraph(f'<b>{label}</b> {problem}', S["ex"])]
    if solution:
        inner.append(Paragraph(solution, S["exsol"]))
    t = Table([[inner]], colWidths=[470])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(_PLAN_EX_BG)),
        ("LINEBEFORE", (0, 0), (0, -1), 2, HexColor(_PLAN_BAR)),
        ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return [t, Spacer(1, 5)]


_ANSWER_WORDS = ("ҷавоб", "җавоб", "жавоб", "javob", "ответ", "answer",
                 "натиҷа", "натича", "натижа", "natija", "результат", "result")


def _plan_mark_answer(text: str) -> str:
    body = str(text or "")
    low = body.lower()
    cut = max((low.rfind(w) for w in _ANSWER_WORDS), default=-1)
    if cut < 0:
        return body
    head, tail = body[:cut], body[cut:]
    out = []
    pos = 0
    for m in _MATH_SPAN.finditer(tail):
        text = tail[pos:m.start()]
        if text:
            body, sep = text.rstrip(), text[len(text.rstrip()):] or " "
            if body:
                out.append(f'<font backColor="{_PLAN_MARK}"><b>{body}</b></font>')
            out.append(sep)
        got = _math_png(m.group(1), 10.5, _PLAN_MARK)
        if got:
            path, w, asc, desc = got
            out.append(f'<img src="{path}" width="{w + 2 * _MATH_PAD:.1f}" '
                       f'height="{asc + desc + 2 * _MATH_PAD:.1f}" '
                       f'valign="{-(desc + _MATH_PAD):.1f}"/>')
        else:
            out.append(f'<font backColor="{_PLAN_MARK}"><b>{m.group(1)}</b></font>')
        pos = m.end()
    rest = tail[pos:]
    if rest:
        out.append(f'<font backColor="{_PLAN_MARK}"><b>{rest}</b></font>')
    return head + "".join(out)


def _plan_styles():
    body = ParagraphStyle(
        name="PlanBody", fontName=MATH_FONT, fontSize=11, leading=15,
        alignment=TA_JUSTIFY, textColor=HexColor("#000000"), spaceAfter=6,
    )
    return {
        "body": body,
        "bullet": ParagraphStyle(
            name="PlanBullet", parent=body, alignment=TA_LEFT,
            leftIndent=22, firstLineIndent=-11, spaceAfter=2,
        ),
        "head": ParagraphStyle(
            name="PlanHead", fontName=MATH_FONT_BOLD, fontSize=11, leading=14,
            textColor=HexColor(_PLAN_INK), spaceBefore=0, spaceAfter=0,
        ),
        "title": ParagraphStyle(
            name="PlanTitle", fontName=MATH_FONT_BOLD, fontSize=16, leading=20,
            alignment=TA_CENTER, textColor=HexColor("#000000"),
            spaceBefore=0, spaceAfter=8,
        ),
        "ex": ParagraphStyle(
            name="PlanEx", parent=body, alignment=TA_LEFT, leading=22, spaceAfter=2,
        ),
        "exsol": ParagraphStyle(
            name="PlanExSol", parent=body, alignment=TA_LEFT,
            leftIndent=13, fontSize=10.5, leading=21, spaceAfter=0,
        ),
        "meta": ParagraphStyle(
            name="PlanMeta", fontName=MATH_FONT, fontSize=10.5, leading=14,
            textColor=HexColor("#000000"),
        ),
    }


def _plan_rule(width=470, count=2):
    out = []
    for _ in range(count):
        t = Table([[""]], colWidths=[width], rowHeights=[15])
        t.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 0.6, HexColor("#000000")),
            ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        out.append(t)
    out.append(Spacer(1, 4))
    return out


def _plan_flatten(value):
    if isinstance(value, list):
        parts = [str(v).strip().rstrip(".") for v in value if v]
        return "; ".join(parts) + ("." if parts else "")
    return str(value or "")


def _pdf_plan_body(content: dict, L: dict) -> list:
    S = _plan_styles()
    story = []

    from app.cover_builder import _academic_year
    grade = str(content.get("grade") or "")
    subject = str(content.get("subject") or "")
    language = str(content.get("language") or "Русский")

    story.append(Paragraph(
        f'<para alignment="right">{L.get("plan_year", "Соли таҳсили")} {_academic_year()}</para>',
        ParagraphStyle(name="PlanYear", parent=S["meta"], alignment=TA_RIGHT),
    ))
    fill = Table(
        [[Paragraph(f'{L.get("plan_date", "Сана")} ____________', S["meta"]),
          Paragraph(
              f'{L.get("plan_class", "Синф")} {_bare_grade(grade) or grade or "________"}',
              S["meta"],
          ),
          Paragraph(f'{L.get("plan_school", "Мактаб")} ____________', S["meta"])]],
        colWidths=[160, 150, 160],
    )
    fill.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
    ]))
    story.append(fill)
    story.append(Spacer(1, 14))

    if subject:
        story.append(Paragraph(
            "&nbsp;".join(_display_subject(subject, language).upper()),
            ParagraphStyle(name="PlanSubject", parent=S["meta"], fontName=MATH_FONT_BOLD,
                           fontSize=8.5, alignment=TA_CENTER, spaceAfter=4),
        ))
    story.append(Paragraph(str(content.get("title") or "").upper(), S["title"]))
    rule = Table([[""]], colWidths=[120], rowHeights=[1.6])
    rule.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(_PLAN_BAR)),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(Table([[rule]], colWidths=[470], style=TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
    ])))
    story.append(_PdfContentAnchor())

    def heading(text):
        t = Table([[Paragraph(f"<b>{text}</b>", S["head"])]], colWidths=[470])
        t.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 0.7, HexColor(_PLAN_BAR)),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    def bullets(value, bg="white"):
        out = []
        for item in (value if isinstance(value, list) else [value]):
            out.append(Paragraph(
                f"&nbsp;&nbsp;•&nbsp;&nbsp;{_math_inline(item, 11, bg)}", S["bullet"]))
        out.append(Spacer(1, 4))
        return out

    heavy: list[list] = []

    formulas = content.get("formulas") or []
    if formulas:
        head = heading(L.get("formulas", "Формулаҳо"))
        blk = [head]
        for f in formulas:
            latex = f.get("latex") if isinstance(f, dict) else None
            plain = (f.get("formula", "") if isinstance(f, dict) else str(f))
            if not latex and plain:
                latex = _plain_to_latex(plain)
            drawn = _pdf_math_display(latex, 14, bg=_PLAN_BG) if latex else None
            if drawn is not None:
                blk.append(drawn)
                if isinstance(f, dict) and f.get("explanation"):
                    blk.append(Paragraph(
                        f'<para alignment="center">'
                        f'{_math_inline(f["explanation"], 9.5, _PLAN_BG)}</para>',
                        ParagraphStyle(name=f"Fx{abs(hash(plain)) % 99999}",
                                       parent=S["body"], fontSize=9.5, leading=12.5,
                                       spaceBefore=1, spaceAfter=7),
                    ))
            else:
                line = f"<b>{plain}</b>"
                if isinstance(f, dict) and f.get("explanation"):
                    line += f" — {_math_inline(f['explanation'], 9.5)}"
                blk.append(Paragraph(f"&nbsp;&nbsp;•&nbsp;&nbsp;{line}", S["bullet"]))
        heavy.append([head] + _plan_panel(blk[1:]))

    examples = [e for e in (content.get("worked_examples") or [])
                if isinstance(e, dict) and e.get("problem")]
    if examples:
        blk = list(_plan_feature_head(
            L.get("worked_examples", "Мисолҳо"), S,
            badge=str(len(examples)),
        ))
        for i, ex in enumerate(examples, 1):
            solution = ""
            if ex.get("solution"):
                solution = (f'<i>{L.get("solution", "Ҳал")}:</i> '
                            f'{_math_inline(_plan_mark_answer(ex["solution"]), 10.5, _PLAN_EX_BG)}')
            blk.extend(_plan_example_card(
                f'{L.get("example", "Мисол")} {i}.',
                _math_inline(ex["problem"], 11, _PLAN_EX_BG), solution, S
            ))
        blk.append(Spacer(1, 4))
        heavy.append(blk)

    drills = [d for d in (content.get("practice_problems") or [])
              if isinstance(d, dict) and d.get("problem")]
    if drills:
        blk = list(_plan_feature_head(
            L.get("practice_problems", "Машқҳои мустақил"), S, badge=str(len(drills)),
        ))
        for i, d in enumerate(drills, 1):
            blk.extend(_plan_example_card(
                f'{L.get("example", "Мисол")} {i}.',
                _math_inline(d["problem"], 11, _PLAN_EX_BG), "", S
            ))
        blk.append(Spacer(1, 4))
        heavy.append(blk)

    quick_check = [q for q in (content.get("quick_check") or []) if isinstance(q, dict) and q.get("question")]
    if quick_check:
        blk = list(_plan_feature_head(L.get("quick_check", "Санҷиши зуд"), S, badge=str(len(quick_check))))
        blk.extend(_pdf_quick_check(quick_check, _PLAN_BAR, L.get("answer", "Ҷавоб")))
        heavy.append(blk)

    for block in (content.get("visual_blocks") or []):
        if not isinstance(block, dict):
            continue
        btype = block.get("type")
        data = block.get("data") or {}
        if btype in ("table", "comparison"):
            blk = list(_pdf_visual_block(block, _PLAN_INK, minimal=True))
        elif btype in ("process", "flowchart", "timeline"):
            blk = []
            if block.get("title"):
                blk.append(heading(block["title"]))
            if btype == "timeline":
                for ev in (data.get("events") or []):
                    bits = [b for b in [ev.get("date"), ev.get("label")] if b]
                    line = f"<b>{' — '.join(bits)}</b>"
                    if ev.get("description"):
                        line += f" — {ev['description']}"
                    blk.append(Paragraph(
                        f"&nbsp;&nbsp;•&nbsp;&nbsp;{_math_inline(line, 11)}", S["bullet"]))
            else:
                for i, st in enumerate(data.get("steps") or [], 1):
                    line = f"<b>{st.get('title', '')}</b>"
                    if st.get("description"):
                        line += f" — {st['description']}"
                    blk.append(Paragraph(
                        f"&nbsp;&nbsp;{i}.&nbsp;&nbsp;{_math_inline(line, 11)}", S["bullet"]))
            if len(blk) <= (1 if block.get("title") else 0):
                continue
            blk.append(Spacer(1, 4))
        else:
            continue
        note = block.get("description") or block.get("note")
        if note:
            blk.append(Paragraph(
                f"<i>{_math_inline(note, 10)}</i>",
                ParagraphStyle(name=f"PlanVisNote_{abs(hash(str(note))) % 100000}",
                               parent=S["body"], fontSize=10, leading=13,
                               alignment=TA_LEFT, spaceBefore=1, spaceAfter=8),
            ))
        if blk:
            heavy.append(blk)

    for cb in (content.get("code_blocks") or []):
        if isinstance(cb, dict) and cb.get("code"):
            heavy.append(list(_pdf_code_card(
                cb["code"], cb.get("language", ""), cb.get("explanation", ""), _PLAN_INK)))

    for fig in (content.get("figures") or []):
        if isinstance(fig, dict):
            blk = _pdf_subject_figure(fig, minimal=True, max_w=270, max_h=205)
            if blk:
                heavy.append(blk)

    lesson_by_anchor: dict[str, list] = {}
    for image in (content.get("lesson_images") or []):
        if isinstance(image, dict) and image.get("path"):
            lesson_by_anchor.setdefault(image.get("position_after") or "main_content", []).append(image)

    PLAN_ORDER = [
        ("competencies", "fill"),
        ("warmup", "fill"),
        ("objectives", "bullets"),
        ("key_concepts", "bullets"),
        ("key_terms", "bullets"),
        ("tools", "runin"),
        ("main_content", "runin"),
        ("important_notes", "notes"),
        ("real_life_examples", "bullets"),
        ("lesson_program", "runin"),
        ("pair_work", "runin"),
        ("group_work", "bullets"),
        ("consolidation", "runin"),
        ("homework", "bullets"),
        ("summary", "runin"),
        ("assessment", "runin"),
    ]

    printable = [
        key for key, kind in PLAN_ORDER
        if kind == "fill" or content.get(key)
    ]
    slots = printable[1:]
    drop_after: dict[str, list] = {}
    if slots and heavy:
        for i, blk in enumerate(heavy):
            pos = round((i + 1) * len(slots) / (len(heavy) + 1))
            pos = max(0, min(pos, len(slots) - 1))
            drop_after.setdefault(slots[pos], []).append(blk)

    PANELLED = {"key_concepts", "key_terms"}

    for key, kind in PLAN_ORDER:
        value = content.get(key)
        label = L.get(key, key)
        if kind == "fill":
            story.append(heading(f"{label}:"))
            story.extend(bullets(value) if value else _plan_rule())
        elif value and kind == "notes":
            for note in (value if isinstance(value, list) else [value]):
                if not str(note).strip():
                    continue
                story.extend(_plan_panel([Paragraph(
                    f'<b>{L.get("plan_note", "Диққат")}.</b> '
                    f'{_math_inline(note, 11, _PLAN_BG)}', S["body"]
                )]))
        elif value and kind == "bullets":
            story.append(heading(label))
            if key in PANELLED:
                story.extend(_plan_panel(bullets(value, _PLAN_BG)[:-1]))
            else:
                story.extend(bullets(value))
        elif value and kind == "runin":
            story.append(Paragraph(
                f"<b><i>{label}.</i></b> {_math_inline(_plan_flatten(value), 11)}",
                S["body"]))
        for image in lesson_by_anchor.pop(key, []):
            story.extend(_pdf_lesson_image_block(
                image, L.get("illustration", "Тасвир"), _PLAN_INK))
        for blk in drop_after.pop(key, []):
            story.extend(blk)

    for images in lesson_by_anchor.values():
        for image in images:
            story.extend(_pdf_lesson_image_block(
                image, L.get("illustration", "Тасвир"), _PLAN_INK))
    for leftover in drop_after.values():
        for blk in leftover:
            story.extend(blk)

    drill_answers = [d.get("answer", "") for d in drills if str(d.get("answer") or "").strip()]
    if drill_answers:
        story.extend(_plan_feature_head(L.get("answers", "Ҷавобҳо"), S))
        for i, ans in enumerate(drill_answers, 1):
            story.append(Paragraph(
                f'<b>{i}.</b>&nbsp;&nbsp;{_math_inline(ans, 11, _PLAN_BG)}', S["bullet"]))
        story.append(Spacer(1, 4))

    return story


def _add_konspekt_body_pdf(content: dict, L: dict, academic: bool = False,
                           numbered: bool = True) -> list:
    ACCENT = get_subject_accent_hex(content.get("subject"))
    if content.get("group_work"):
        content = {
            **content,
            "group_work": [_strip_group_label(t) for t in content["group_work"]],
        }
    tmpl = get_template(content.get("template") or "zamonaviy")
    header_font = MATH_FONT if tmpl.font_family == 'serif' else FONT_NAME

    def _header(label, color, number):
        if is_lecture:
            return _pdf_lecture_section_header(label, color, number if numbered else None)
        if tmpl.header_style == "numbered":
            return _pdf_section_header(label, color, number)
        renderer = _PDF_SECTION_RENDERERS.get(tmpl.header_style, _pdf_section_header)
        return renderer(label, color, number, header_font)

    def _num_bullet(j, item, color):
        if tmpl.alt_row_tint:
            return _pdf_numbered_bullet_tinted(j, item, color)
        return [_pdf_numbered_bullet(j, item, color, serif=is_lecture)]

    _visuals_by_position: dict[str, list] = {}
    for _block in (content.get("visual_blocks") or []):
        if isinstance(_block, dict):
            _visuals_by_position.setdefault(_block.get("position_after") or "", []).append(_block)

    _figures = [f for f in (content.get("figures") or [])
                if isinstance(f, dict) and f.get("image")]
    _figure_anchors = ["key_concepts", "main_content", "real_life_examples", "consolidation"]
    _figures_by_anchor: dict[str, list] = {}
    for _i, _fig in enumerate(_figures):
        _figures_by_anchor.setdefault(_figure_anchors[_i % len(_figure_anchors)], []).append(_fig)

    def _render_figures(position_key):
        out = []
        for fig in _figures_by_anchor.pop(position_key, []):
            out.extend(_pdf_subject_figure(fig, tmpl.minimal_chrome))
        return out

    def _render_visuals(position_key):
        rendered = []
        for block in _visuals_by_position.pop(position_key, []):
            rendered.extend(_pdf_visual_block(block, ACCENT, tmpl.minimal_chrome,
                                              serif=is_lecture))
        return rendered

    _lesson_by_anchor: dict[str, list] = {}
    for _img in (content.get("lesson_images") or []):
        if isinstance(_img, dict) and _img.get("path"):
            _lesson_by_anchor.setdefault(_img.get("position_after") or "main_content", []).append(_img)

    def _render_lesson_images(position_key):
        out = []
        for image in _lesson_by_anchor.pop(position_key, []):
            out.extend(_pdf_lesson_image_block(image, L.get('illustration', 'Illustration'), ACCENT))
        return out

    story = []
    is_lecture = academic or bool(content.get("lecture_plan") or content.get("objective"))
    sidebar_sections = () if is_lecture else tmpl.sidebar_sections
    if is_lecture:
        L = dict(L)
        L['summary'] = _LECTURE_SUMMARY_LABEL.get(
            str(content.get('language') or 'Русский'), L['summary'])
    if sidebar_sections:
        story.extend(_pdf_sidebar_panel(content, sidebar_sections, ACCENT, L))
    sections_map = [
        ("objective", L.get('objective', 'Цель'), ACCENT),
        ("lecture_plan", L.get('lecture_plan', 'План'), ACCENT),
        ("objectives", L['objectives'], ACCENT),
        ("competencies", L['competencies'], ACCENT),
        ("key_concepts", L['key_concepts'], ACCENT),
        ("key_terms", L['key_terms'], ACCENT),
        ("tools", L['tools'], ACCENT),
        ("lesson_program", L['lesson_program'], ACCENT),
        ("main_content", L['main_content'], ACCENT),
        ("real_life_examples", L['real_life_examples'], ACCENT),
        ("pair_work", L['pair_work'], ACCENT),
        ("group_work", L['group_work'], ACCENT),
        ("consolidation", L['consolidation'], ACCENT),
    ]

    num = 0
    for key, label, color in sections_map:
        val = content.get(key)
        if val is None:
            continue
        if key not in sidebar_sections:
            if isinstance(val, list):
                if val:
                    num += 1
                    header_flowables = _header(label, color, num)
                    first_item = _num_bullet(1, val[0], color)
                    story.append(KeepTogether(header_flowables + first_item))
                    for j, item in enumerate(val[1:], start=2):
                        story.extend(_num_bullet(j, item, color))
            elif isinstance(val, str) and val.strip():
                num += 1
                header_flowables = _header(label, color, num)
                if key == "main_content":
                    paragraphs = _pdf_prose_paragraph(val, serif=is_lecture)
                    if paragraphs:
                        story.append(KeepTogether(header_flowables + [paragraphs[0]]))
                        story.extend(paragraphs[1:])
                    else:
                        story.extend(header_flowables)
                else:
                    story.append(KeepTogether(header_flowables + [_pdf_bullet(val, color, serif=is_lecture)]))
        story.extend(_render_lesson_images(key))
        story.extend(_render_figures(key))
        story.extend(_render_visuals(key))

        if key == "key_concepts":
            formulas = content.get("formulas")
            if formulas:
                story.append(Spacer(1, 4))
                for f in formulas:
                    if isinstance(f, dict):
                        story.extend(_pdf_formula_card(
                            f.get('formula', ''), f.get('explanation', ''),
                            tmpl.minimal_chrome or is_lecture, f.get('latex'), ACCENT))
                    else:
                        story.extend(_pdf_formula_card(str(f), '', tmpl.minimal_chrome or is_lecture, accent=ACCENT))
            concept_cards = content.get("concept_cards")
            if concept_cards:
                story.extend(_pdf_concept_card_grid(concept_cards, ACCENT))
        if key == "key_terms":
            for cb in (content.get("code_blocks") or []):
                if isinstance(cb, dict) and cb.get("code"):
                    story.extend(_pdf_code_card(cb["code"], cb.get("language", ""), cb.get("explanation", ""), ACCENT))
        if key == "main_content":
            worked_examples = content.get("worked_examples")
            if worked_examples:
                num += 1
                story.extend(_header(L['worked_examples'], ACCENT, num))
                story.extend(_pdf_worked_examples(worked_examples, ACCENT, L['solution'], tmpl.minimal_chrome))
                story.extend(_render_lesson_images("worked_examples"))
        if key == "real_life_examples":
            map_image = content.get("map_image")
            if map_image:
                locations = content.get("map_locations", [])
                caption = ', '.join(locations) if locations else ''
                img = _pdf_illustration_image(map_image, caption, '© OpenStreetMap contributors', max_w=400, max_h=262)
                if img:
                    story.extend(img)
            real_image = content.get("real_image")
            if isinstance(real_image, dict) and real_image.get("path"):
                box = _REAL_IMAGE_BOX.get(real_image.get("style"), _REAL_IMAGE_BOX_DEFAULT)
                img = _pdf_illustration_image(real_image["path"], real_image.get("caption", ""), "Wikipedia", max_w=box, max_h=box)
                if img:
                    story.extend(img)
    for _left in list(_lesson_by_anchor):
        story.extend(_render_lesson_images(_left))
    for _left in list(_figures_by_anchor):
        story.extend(_render_figures(_left))
    story.extend(_render_visuals("visual_aid"))
    for _leftover_key in list(_visuals_by_position):
        story.extend(_render_visuals(_leftover_key))

    important_notes = content.get("important_notes") or []
    for note in important_notes:
        num += 1
        story.append(Spacer(1, 4))
        story.append(_pdf_important_note(str(note), ACCENT))

    quick_check = content.get("quick_check") or []
    if quick_check:
        num += 1
        story.extend(_header(L['quick_check'], ACCENT, num))
        story.extend(_pdf_quick_check(quick_check, ACCENT, L['answer']))

    homework = content.get("homework")
    if homework:
        num += 1
        story.extend(_header(L['homework'], ACCENT, num))
        if isinstance(homework, list):
            for i, h in enumerate(homework):
                story.extend(_num_bullet(i + 1, h, ACCENT))
        else:
            story.append(_pdf_bullet(str(homework), ACCENT, serif=is_lecture))

    summary = content.get("summary", "")
    if summary:
        num += 1
        story.extend(_header(L['summary'], ACCENT, num))
        story.append(_pdf_bullet(summary, ACCENT, serif=is_lecture))

    assessment = content.get("assessment", "")
    if assessment:
        num += 1
        story.extend(_header(L['assessment'], ACCENT, num))
        story.append(_pdf_bullet(assessment, ACCENT, serif=is_lecture))

    references = [r for r in (content.get("references") or []) if str(r).strip()]
    if references:
        num += 1
        story.extend(_header(L.get('references', 'Литература'), ACCENT, num))
        for i, ref in enumerate(references, start=1):
            story.extend(_num_bullet(i, str(ref), ACCENT))

    return story



_MONO_INK = '#111111'
_MONO_TEXT = '#222222'

_CURRICULUM_LABELS = {
    'Русский': {
        'duration': 'Время урока', 'competencies': 'Компетенции', 'objectives': 'Цели урока',
        'key_concepts': 'Ключевые понятия', 'key_terms': 'Словарь урока', 'lesson_program': 'Программа урока',
        'tools': 'Инструменты', 'warmup': 'Разминка', 'main_content': 'Основное содержание',
        'real_life_examples': 'Примеры из жизни', 'visual_aid': 'Наглядное пособие',
        'pair_work': 'Парная работа', 'consolidation': 'Закрепление', 'fun_facts': 'Интересные факты',
        'common_mistakes': 'Частые ошибки', 'summary': 'Итоги урока', 'homework': 'Домашнее задание',
        'assessment': 'Оценка', 'group_work': 'Групповая работа', 'group': 'Группа',
        'day': 'День', 'date': 'Дата', 'topic': 'Тема', 'type': 'Тип', 'lesson': 'Урок', 'exam': 'Экзамен',
        'roadmap': 'Дорожная карта курса', 'answer': 'Примерный ответ', 'explanation': 'Объяснение',
    },
    'Таджикский': {
        'duration': 'Вақти дарс', 'competencies': 'Салоҳиятҳо', 'objectives': 'Мақсадҳои дарс',
        'key_concepts': 'Мафҳумҳои асосӣ', 'key_terms': 'Луғати дарс', 'lesson_program': 'Барномаи дарс',
        'tools': 'Воситаҳои аёнӣ', 'warmup': 'Санҷиши дониш', 'main_content': 'Шиносоӣ бо мазмуни мавзӯъ',
        'real_life_examples': 'Дар ҳаёти воқеӣ', 'visual_aid': 'Ёрии визуалӣ',
        'pair_work': 'Кори дунафара', 'consolidation': 'Мустаҳкамкунии дарс', 'fun_facts': 'Фактҳои ҷолиб',
        'common_mistakes': 'Хатогиҳои маъмул', 'summary': 'Хулосаи дарс', 'homework': 'Супориши хонагӣ',
        'assessment': 'Арзёбӣ', 'group_work': 'Кори гурӯҳӣ', 'group': 'Гурӯҳи',
        'day': 'Рӯз', 'date': 'Сана', 'topic': 'Мавзӯъ', 'type': 'Навъ', 'lesson': 'Дарс', 'exam': 'Имтиҳон',
        'roadmap': 'Харитаи курс', 'answer': 'Ҷавоби намунавӣ', 'explanation': 'Шарҳ',
    },
    'English': {
        'duration': 'Duration', 'competencies': 'Competencies', 'objectives': 'Objectives',
        'key_concepts': 'Key Concepts', 'key_terms': 'Lesson Glossary', 'lesson_program': 'Lesson Program',
        'tools': 'Teaching Tools', 'warmup': 'Warm-up', 'main_content': 'Main Content',
        'real_life_examples': 'Real Life', 'visual_aid': 'Visual Aid',
        'pair_work': 'Pair Work', 'consolidation': 'Consolidation', 'fun_facts': 'Fun Facts',
        'common_mistakes': 'Common Mistakes', 'summary': 'Summary', 'homework': 'Homework',
        'assessment': 'Assessment', 'group_work': 'Group Work', 'group': 'Group',
        'day': 'Day', 'date': 'Date', 'topic': 'Topic', 'type': 'Type', 'lesson': 'Lesson', 'exam': 'Exam',
        'roadmap': 'Course Roadmap', 'answer': 'Model Answer', 'explanation': 'Explanation',
    },
}
_CURRICULUM_LABELS['Английский'] = _CURRICULUM_LABELS['English']


def _pdf_mono_para(text, size=10.5, italic=False):
    tag = 'i' if italic else 'span'
    style = ParagraphStyle(
        name=f'MonoPara_{hash(text) % 10000}', fontName=FONT_NAME, fontSize=size,
        leading=size + 4, textColor=HexColor(_MONO_TEXT), spaceAfter=5,
    )
    return Paragraph(f'<{tag}>{text}</{tag}>' if italic else text, style)


def _pdf_curriculum_exam_content(content: dict, R: dict, accent: str | None = None) -> list:
    accent = accent or _PDF_ACCENT
    elements = []
    description = content.get('description', '')
    if description:
        elements.append(Paragraph(
            str(description),
            ParagraphStyle(name='ExamDesc', fontName=MATH_FONT, fontSize=10, leading=14,
                           textColor=HexColor(_LECTURE_MUTED), spaceAfter=6)))
    for i, q in enumerate(content.get('questions', [])):
        q_type = q.get('type', 'multiple_choice')
        q_style = ParagraphStyle(
            name=f'ExamQ_{i}', fontName=MATH_FONT_BOLD, fontSize=11.5, leading=16,
            textColor=HexColor(_LECTURE_INK), leftIndent=20, firstLineIndent=-20,
            spaceBefore=9, spaceAfter=4,
        )
        elements.append(Paragraph(
            f'<font color="{accent}"><b>{i + 1}.</b></font>&nbsp;&nbsp;{q.get("question", "")}',
            q_style))
        if q_type == 'open_ended':
            if q.get('model_answer'):
                elements.append(_pdf_answer_box(f'{R["answer"]}: {q["model_answer"]}', accent))
        else:
            options = q.get('options', [])
            correct_indices = set(q.get('correct_indices') or [])
            correct = q.get('correct_index', 0)
            for j, opt in enumerate(options):
                is_correct = (j in correct_indices) if q_type == 'multiple_select' else (j == correct)
                letter = chr(65 + j)
                opt_style = ParagraphStyle(
                    name=f'ExamOpt_{i}_{j}',
                    fontName=MATH_FONT_BOLD if is_correct else MATH_FONT,
                    fontSize=10.5, leading=15, textColor=HexColor(_LECTURE_INK),
                    backColor=HexColor(_pdf_light_tint_hex(accent, amount=0.90)) if is_correct else None,
                    borderPadding=(2, 5, 2, 5) if is_correct else 0,
                    leftIndent=36 if is_correct else 40, spaceAfter=2,
                )
                text = f'<b><u>{letter}) {opt}</u></b>' if is_correct else f'{letter}) {opt}'
                elements.append(Paragraph(text, opt_style))
        if q.get('explanation'):
            elements.append(_pdf_answer_box(f'{R["explanation"]}: {q["explanation"]}', accent))
    return elements


_TEST_COUNT_LABEL = {
    'Русский': 'вопросов', 'Таджикский': 'савол',
    'English': 'questions',
}
_TEST_COUNT_LABEL['Английский'] = _TEST_COUNT_LABEL['English']


_TEST_UI = {
    'Русский': {
        'name': 'Ф.И.О.', 'klass': 'Класс', 'date': 'Дата', 'mark': 'Оценка',
        'hint': 'Внимательно прочитайте каждый вопрос. В заданиях с одним ответом '
                'обведите букву верного варианта, в заданиях с несколькими ответами — '
                'отметьте все верные, на открытый вопрос ответьте письменно.',
        'key_title': 'КЛЮЧ К ТЕСТУ', 'key_hint': 'Лист для учителя — не выдавайте ученикам.',
        'col_no': '№', 'col_answer': 'Ответ', 'col_why': 'Пояснение',
        'open': 'развёрнутый ответ', 'model': 'Примерный ответ',
        'total': 'Всего вопросов', 'max_score': 'Максимальный балл',
        'types': {'true_false': 'верно / неверно', 'multiple_select': 'несколько ответов',
                  'open_ended': 'открытый вопрос'},
        'select_hint': 'выберите все верные', 'tf': ('Верно', 'Неверно'),
    },
    'Таджикский': {
        'name': 'Ном ва насаб', 'klass': 'Синф', 'date': 'Сана', 'mark': 'Баҳо',
        'hint': 'Ҳар саволро бодиққат хонед. Дар саволҳои якҷавоба ҳарфи ҷавоби дурустро '
                'давра кунед, дар саволҳои бисёрҷавоба ҳамаи ҷавобҳои дурустро қайд кунед, '
                'ба саволи кушод хаттӣ ҷавоб диҳед.',
        'key_title': 'КАЛИДИ ТЕСТ', 'key_hint': 'Варақаи муаллим — ба хонандагон надиҳед.',
        'col_no': '№', 'col_answer': 'Ҷавоб', 'col_why': 'Шарҳ',
        'open': 'ҷавоби муфассал', 'model': 'Ҷавоби намунавӣ',
        'total': 'Ҳамагӣ саволҳо', 'max_score': 'Холи ниҳоӣ',
        'types': {'true_false': 'дуруст / нодуруст', 'multiple_select': 'якчанд ҷавоб',
                  'open_ended': 'саволи кушод'},
        'select_hint': 'ҳамаи ҷавобҳои дурустро интихоб кунед', 'tf': ('Дуруст', 'Нодуруст'),
    },
    'English': {
        'name': 'Name', 'klass': 'Class', 'date': 'Date', 'mark': 'Mark',
        'hint': 'Read each question carefully. For single-answer questions circle the letter '
                'of the correct option, for multiple-answer questions mark every correct '
                'option, and answer the open question in writing.',
        'key_title': 'ANSWER KEY', 'key_hint': 'Teacher\'s sheet — do not hand out to students.',
        'col_no': '#', 'col_answer': 'Answer', 'col_why': 'Explanation',
        'open': 'written answer', 'model': 'Model answer',
        'total': 'Questions', 'max_score': 'Maximum score',
        'types': {'true_false': 'true / false', 'multiple_select': 'multiple answers',
                  'open_ended': 'open question'},
        'select_hint': 'choose all that apply', 'tf': ('True', 'False'),
    },
}
_TEST_UI['Английский'] = _TEST_UI['English']


def _test_ui(language: str) -> dict:
    return _TEST_UI.get(str(language or 'Русский'), _TEST_UI['Русский'])


def _pdf_student_header(ui: dict, accent: str):
    cell = ParagraphStyle(name='TestFieldLbl', fontName=FONT_NAME, fontSize=9,
                          leading=13, textColor=HexColor(_LECTURE_MUTED))
    def field(label, width):
        return Paragraph(f'{label}&nbsp;<font color="{_LECTURE_RULE}">'
                         f'{"_" * width}</font>', cell)
    row = [[field(ui['name'], 26), field(ui['klass'], 7),
            field(ui['date'], 9), field(ui['mark'], 7)]]
    t = Table(row, colWidths=[228, 74, 86, 72])
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LINEBELOW', (0, 0), (-1, -1), 0.7, HexColor(_LECTURE_RULE)),
    ]))
    return [t, Spacer(1, 8)]


def _pdf_answer_lines(count: int = 3, width: str = '100%'):
    out = []
    for _ in range(count):
        out.append(Spacer(1, 15))
        out.append(HRFlowable(width=width, thickness=0.5,
                              color=HexColor(_LECTURE_RULE), hAlign='RIGHT'))
    out.append(Spacer(1, 6))
    return out


def build_test_pdf(content: dict, include_key: bool = True,
                   student_only: bool = False) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    story = []

    title = content.get("title", "Тест")
    desc = content.get("description", "")
    subject = content.get("subject", "")
    grade = content.get("grade", "")
    language = str(content.get("language") or "Русский")
    ui = _test_ui(language)
    questions = content.get("questions", [])
    accent = get_subject_accent_hex(subject)

    story.extend(_pdf_lecture_masthead(
        title, desc, subject, grade, language,
        _PDF_KONSPEKT_LABELS.get(language,
                                 _PDF_KONSPEKT_LABELS['Русский']).get('grade_suffix', 'класс'),
        accent=accent))
    story.extend(_pdf_student_header(ui, accent))
    story.append(Paragraph(
        ui['hint'],
        ParagraphStyle(name="TestHint", fontName=FONT_NAME, fontSize=8.5, leading=12,
                       alignment=TA_JUSTIFY, textColor=HexColor(_LECTURE_MUTED),
                       spaceAfter=2)))

    _LETTERS = ['A', 'B', 'C', 'D', 'E', 'F']

    def letters_for(q):
        q_type = q.get('type', 'multiple_choice')
        options = q.get('options') or []
        if q_type == 'open_ended':
            return ui['open']
        if q_type == 'multiple_select':
            idx = sorted(int(i) for i in (q.get('correct_indices') or [])
                         if isinstance(i, (int, float)) and 0 <= int(i) < len(options))
            return ', '.join(_LETTERS[i] for i in idx) if idx else '—'
        i = q.get('correct_index')
        if not isinstance(i, int) or not (0 <= i < len(options)):
            return '—'
        return (f'{_LETTERS[i]} · {options[i]}' if q_type == 'true_false'
                else _LETTERS[i])

    for i, q in enumerate(questions):
        q_type = q.get('type', 'multiple_choice')
        question_style = ParagraphStyle(
            name=f'Q_{i}', fontName=MATH_FONT_BOLD, fontSize=12.5, leading=17,
            textColor=HexColor(_LECTURE_INK), leftIndent=22, firstLineIndent=-22,
            spaceAfter=7)
        opt_style = ParagraphStyle(
            name=f'Opt_{i}', fontName=MATH_FONT, fontSize=11, leading=15.5,
            textColor=HexColor(_LECTURE_INK), leftIndent=40, spaceAfter=3)

        story.append(Spacer(1, 12))
        story.append(HRFlowable(width="100%", thickness=0.7, color=HexColor(_LECTURE_RULE)))
        story.append(Spacer(1, 7))

        hint = (f' <font color="{_LECTURE_MUTED}" size="9">({ui["select_hint"]})</font>'
                if q_type == 'multiple_select' else '')
        story.append(Paragraph(
            f'<font color="{accent}"><b>{i + 1}.</b></font>&nbsp;&nbsp;'
            f'{_math_inline(q.get("question", ""), 12.5)}{hint}', question_style))

        img = q.get("image") if isinstance(q.get("image"), dict) else None
        if img and img.get("path"):
            body = _pdf_illustration_image(img.get("path", ""), "", img.get("credit", ""),
                                           max_w=220, max_h=160)
            if body:
                story.append(KeepTogether([Spacer(1, 4)] + body))

        if q_type == 'open_ended':
            story.extend(_pdf_answer_lines(3))
        else:
            for j, opt in enumerate(q.get("options", [])):
                letter = _LETTERS[j] if j < len(_LETTERS) else str(j)
                box = '<font color="%s">[&nbsp;&nbsp;]</font>&nbsp;' % _LECTURE_RULE \
                    if q_type == 'multiple_select' else ''
                story.append(Paragraph(f'{box}{letter})&nbsp;{_math_inline(opt, 11)}', opt_style))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.7, color=HexColor(_LECTURE_RULE)))
    story.append(Paragraph(
        f"{ui['total']}: {len(questions)}&nbsp;&nbsp;·&nbsp;&nbsp;"
        f"{ui['max_score']}: {len(questions)}",
        ParagraphStyle(name="TestFoot", fontName=FONT_NAME, fontSize=8.5, leading=12,
                       alignment=TA_CENTER, textColor=HexColor(_LECTURE_MUTED),
                       spaceBefore=5)))

    if include_key and not student_only and questions:
        from reportlab.platypus import PageBreak
        story.append(PageBreak())
        story.append(Spacer(1, 4))
        story.append(Paragraph(
            ui['key_title'],
            ParagraphStyle(name="KeyTitle", fontName=MATH_FONT_BOLD, fontSize=15, leading=19,
                           alignment=TA_CENTER, textColor=HexColor(_LECTURE_INK))))
        story.append(Paragraph(
            ui['key_hint'],
            ParagraphStyle(name="KeyHint", fontName=FONT_NAME, fontSize=8.5, leading=12,
                           alignment=TA_CENTER, textColor=HexColor(_LECTURE_MUTED),
                           spaceAfter=12)))

        head = ParagraphStyle(name="KeyHead", fontName=FONT_NAME_BOLD, fontSize=9, leading=12,
                              textColor=HexColor("#FFFFFF"))
        num = ParagraphStyle(name="KeyNum", fontName=FONT_NAME_BOLD, fontSize=10, leading=14,
                             alignment=TA_CENTER, textColor=HexColor(_LECTURE_INK))
        ans = ParagraphStyle(name="KeyAns", fontName=FONT_NAME_BOLD, fontSize=10, leading=14,
                             textColor=HexColor(accent))
        why = ParagraphStyle(name="KeyWhy", fontName=FONT_NAME, fontSize=9, leading=12.5,
                             textColor=HexColor(_LECTURE_MUTED))

        rows = [[Paragraph(ui['col_no'], head), Paragraph(ui['col_answer'], head),
                 Paragraph(ui['col_why'], head)]]
        for i, q in enumerate(questions):
            note = _math_inline(str(q.get('explanation') or '').strip(), 9)
            if q.get('type') == 'open_ended' and q.get('model_answer'):
                model = _math_inline(str(q['model_answer']).strip(), 9)
                note = (f'<b>{ui["model"]}:</b> {model}<br/>{note}' if note
                        else f'<b>{ui["model"]}:</b> {model}')
            rows.append([Paragraph(str(i + 1), num),
                         Paragraph(letters_for(q), ans),
                         Paragraph(note or '&nbsp;', why)])

        table = Table(rows, colWidths=[30, 95, 345], repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), HexColor(accent)),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1),
             [HexColor("#FFFFFF"), HexColor(_pdf_light_tint_hex(accent, amount=0.955))]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('GRID', (0, 0), (-1, -1), 0.5, HexColor(_LECTURE_RULE)),
            ('LEFTPADDING', (0, 0), (-1, -1), 7),
            ('RIGHTPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(table)

    doc.build(story)
    buf.seek(0)
    return buf



_PRACTICAL_UI = {
    'Русский': {'individual': 'Индивидуальные задания', 'group': 'Групповые задания',
               'difficulty': {'easy': 'лёгкое', 'medium': 'среднее', 'hard': 'сложное'},
               'roles': 'Роли в группе', 'outcome': 'Результат', 'group_size': 'Состав группы'},
    'Таджикский': {'individual': 'Супоришҳои инфиродӣ', 'group': 'Супоришҳои гурӯҳӣ',
                  'difficulty': {'easy': 'сабук', 'medium': 'миёна', 'hard': 'душвор'},
                  'roles': 'Нақшҳо дар гурӯҳ', 'outcome': 'Натиҷа', 'group_size': 'Таркиби гурӯҳ'},
    'English': {'individual': 'Individual tasks', 'group': 'Group tasks',
               'difficulty': {'easy': 'easy', 'medium': 'medium', 'hard': 'hard'},
               'roles': 'Roles in the group', 'outcome': 'Expected outcome', 'group_size': 'Group size'},
}


def build_practical_pdf(content: dict) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    story = []

    title = content.get("title", "Амалӣ супоришҳо")
    desc = content.get("description", "")
    subject = content.get("subject", "")
    grade = content.get("grade", "")
    language = str(content.get("language") or "Русский")
    ui = _test_ui(language)
    pui = _PRACTICAL_UI.get(language, _PRACTICAL_UI['Русский'])
    accent = get_subject_accent_hex(subject)

    story.extend(_pdf_lecture_masthead(
        title, desc, subject, grade, language,
        _PDF_KONSPEKT_LABELS.get(language,
                                 _PDF_KONSPEKT_LABELS['Русский']).get('grade_suffix', 'класс'),
        accent=accent))
    story.extend(_pdf_student_header(ui, accent))

    section_style = ParagraphStyle(name="PracticalSection", fontName=FONT_NAME_BOLD, fontSize=13.5,
                                   leading=17, textColor=HexColor(accent), spaceBefore=16, spaceAfter=8)
    task_title_style = ParagraphStyle(name="PracticalTaskTitle", fontName=MATH_FONT_BOLD, fontSize=12,
                                      leading=16, textColor=HexColor(_LECTURE_INK), spaceAfter=3)
    body_style = ParagraphStyle(name="PracticalBody", fontName=FONT_NAME, fontSize=10.5, leading=15,
                                alignment=TA_JUSTIFY, textColor=HexColor(_LECTURE_INK), spaceAfter=4)
    label_style = ParagraphStyle(name="PracticalLabel", fontName=FONT_NAME_BOLD, fontSize=9,
                                 leading=13, textColor=HexColor(_LECTURE_MUTED))
    diff_style = ParagraphStyle(name="PracticalDiff", fontName=FONT_NAME_BOLD, fontSize=9,
                                leading=13, textColor=HexColor("#FFFFFF"))

    def diff_badge(level: str):
        label = pui['difficulty'].get(str(level or 'medium').lower(), level or '')
        badge_color = {'easy': '#16A34A', 'medium': '#D97706', 'hard': '#DC2626'}.get(
            str(level or 'medium').lower(), '#6B7280')
        t = Table([[Paragraph(label.upper(), diff_style)]], colWidths=[60])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), HexColor(badge_color)),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        return t

    individual = content.get("individual_tasks") or []
    if individual:
        story.append(Paragraph(pui['individual'], section_style))
        for i, task in enumerate(individual):
            if not isinstance(task, dict):
                continue
            head = Table([[Paragraph(f"{i + 1}. {task.get('title', '')}", task_title_style),
                          diff_badge(task.get('difficulty'))]], colWidths=[400, 65])
            head.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                                      ('LEFTPADDING', (0, 0), (-1, -1), 0),
                                      ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
            body = [head, Spacer(1, 3), Paragraph(task.get("instructions", ""), body_style)]
            if task.get("expected_outcome"):
                body.append(Paragraph(f"{pui['outcome']}: {task['expected_outcome']}", label_style))
            story.append(KeepTogether(body))
            story.append(Spacer(1, 10))

    group = content.get("group_tasks") or []
    if group:
        story.append(Paragraph(pui['group'], section_style))
        for i, task in enumerate(group):
            if not isinstance(task, dict):
                continue
            body = [Paragraph(f"{i + 1}. {task.get('title', '')}", task_title_style)]
            if task.get("group_size"):
                body.append(Paragraph(f"{pui['group_size']}: {task['group_size']}", label_style))
            body.append(Spacer(1, 3))
            body.append(Paragraph(task.get("instructions", ""), body_style))
            roles = task.get("roles") or []
            if roles:
                body.append(Paragraph(pui['roles'], label_style))
                for role in roles:
                    body.append(Paragraph(f"• {role}", body_style))
            if task.get("expected_outcome"):
                body.append(Paragraph(f"{pui['outcome']}: {task['expected_outcome']}", label_style))
            story.append(KeepTogether(body))
            story.append(Spacer(1, 10))

    doc.build(story)
    buf.seek(0)
    return buf


def _pdf_notebook_page(canvas_obj, doc_obj):
    w, h = A4
    canvas_obj.saveState()

    canvas_obj.setFillColor(HexColor('#FFFDF7'))
    canvas_obj.rect(0, 0, w, h, stroke=0, fill=1)

    canvas_obj.setStrokeColor(HexColor('#EEE4C8'))
    canvas_obj.setLineWidth(0.9)
    y = 40
    while y < h - 30:
        canvas_obj.line(0, y, w, y)
        y += 23

    canvas_obj.setStrokeColor(HexColor('#F3C6C6'))
    canvas_obj.setLineWidth(1.4)
    canvas_obj.line(38, 20, 38, h - 20)

    canvas_obj.setStrokeColor(HexColor('#D8D2C4'))
    canvas_obj.setFillColor(HexColor('#FFFDF7'))
    canvas_obj.setLineWidth(1.1)
    y = h - 44
    while y > 30:
        canvas_obj.circle(19, y, 6, stroke=1, fill=1)
        y -= 39

    canvas_obj.restoreState()


def build_presentation_pdf(content: dict) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    story = []

    title = content.get("title", "Презентация")
    desc = content.get("description", "")
    subject = content.get("subject", "")
    grade = content.get("grade", "")
    language = content.get("language", "Русский")

    slides = content.get("slides", [])

    tmpl = get_template(content.get("template"))
    ACCENT = get_subject_accent_hex(content.get("subject"))
    ACCENT_DARK = _pdf_dark_shade_hex(ACCENT)
    ACCENT_SOFT = _pdf_light_tint_hex(ACCENT)
    title_font = MATH_FONT_BOLD if tmpl.font_family == 'serif' else FONT_NAME_BOLD
    body_font = MATH_FONT if tmpl.font_family == 'serif' else FONT_NAME

    cover_png = None
    try:
        from app.cover_builder import build_cover_image
        cover_png = build_cover_image(
            subject, title, grade, language,
            template_id=content.get("template") or "zamonaviy",
            doc_type_label="ПРЕЗЕНТАЦИЯ",
        )
    except Exception as e:
        logger.warning(f"Presentation cover generation failed for '{title[:50]}': {e}")

    def _draw_cover(canvas_obj, doc_obj):
        if cover_png:
            canvas_obj.drawImage(ImageReader(io.BytesIO(cover_png)), 0, 0, width=A4[0], height=A4[1])

    if cover_png:
        from reportlab.platypus import PageBreak
        story.append(PageBreak())
    elif desc:
        story.append(Paragraph(f'<font color="#64748B">{desc}</font>', _styles()['DocSubtitle']))
        story.append(Spacer(1, 14))

    minimal_header = tmpl.header_style in ("underline", "smallcaps")

    for i, s in enumerate(slides):
        slide_title = s.get("title", f"Слайд {i + 1}")
        bullets = s.get("bullet_points", [])
        notes = s.get("speaker_notes", "")

        story.append(Spacer(1, 8))

        if minimal_header:
            d = Drawing(460, 22)
            num_str = f'{i + 1:02d}'
            d.add(String(0, 4, num_str, fontName=title_font, fontSize=11, fillColor=HexColor(ACCENT)))
            d.add(String(24, 4, slide_title.upper() if tmpl.header_style == "smallcaps" else slide_title,
                          fontName=title_font, fontSize=13, fillColor=HexColor(_PDF_DARK)))
            story.append(d)
            story.append(HRFlowable(width="18%", thickness=1.5, color=HexColor(ACCENT), spaceAfter=6))
        elif tmpl.header_style == "bar":
            d = Drawing(460, 30)
            d.add(Rect(0, 0, 460, 30, fillColor=HexColor(ACCENT), strokeColor=None))
            d.add(String(10, 10, f'{i + 1}. {slide_title}', fontName=title_font, fontSize=13, fillColor=HexColor('#FFFFFF')))
            story.append(d)
        else:
            d = Drawing(460, 32)
            d.add(Rect(0, 0, 460, 30, fillColor=HexColor(ACCENT_SOFT), strokeColor=None, rx=7, ry=7))
            d.add(Rect(0, 0, 5, 30, fillColor=HexColor(ACCENT), strokeColor=None))
            d.add(Circle(27, 15, 11, fillColor=HexColor(ACCENT), strokeColor=None))
            num_str = str(i + 1)
            d.add(String(27 - (3.5 * len(num_str)), 11, num_str, fontName=FONT_NAME_BOLD, fontSize=11, fillColor=HexColor('#FFFFFF')))
            d.add(String(48, 10, slide_title, fontName=title_font, fontSize=13, fillColor=HexColor(_PDF_DARK)))
            story.append(d)

        for j, bp in enumerate(bullets):
            story.append(_pdf_numbered_bullet(j + 1, bp, ACCENT))

        body_text = str(s.get("body") or "").strip()
        if body_text:
            story.append(Paragraph(body_text, ParagraphStyle(
                name=f'DeckBody_{i}', fontName=body_font, fontSize=10,
                leading=14.5, textColor=HexColor('#334155'),
                alignment=TA_JUSTIFY, leftIndent=20, rightIndent=6,
                spaceBefore=3, spaceAfter=6)))

        image = s.get("image")
        if isinstance(image, dict) and image.get("path"):
            figure = _pdf_illustration_image(
                image.get("path", ""), "", image.get("credit", ""),
                max_w=300, max_h=200)
            if figure:
                story.append(Spacer(1, 4))
                story.extend(figure)

        visual = s.get("visual")
        if visual:
            story.extend(_pdf_visual_block(visual, ACCENT))

        if notes:
            notes_style = ParagraphStyle(
                name=f'Notes_{i}',
                fontName=body_font,
                fontSize=9,
                leading=13,
                textColor=HexColor('#64748B'),
                backColor=HexColor('#F8FAFC'),
                borderPadding=(4, 6, 4, 6),
                leftIndent=20,
                rightIndent=10,
                spaceAfter=8,
            )
            story.append(Spacer(1, 4))
            story.append(Paragraph(f'<font color="{ACCENT}"><b>Заметки:</b></font> {notes}', notes_style))

        story.append(HRFlowable(width="100%", thickness=0.5, color=HexColor('#E2E8F0')))

    story.extend(_pdf_footer())
    notebook = (content.get("template") or "") == "playful"
    doc.build(
        story,
        onFirstPage=_draw_cover,
        onLaterPages=_pdf_notebook_page if notebook else (lambda c, d: None),
    )
    buf.seek(0)
    return buf
